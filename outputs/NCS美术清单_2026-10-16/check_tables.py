from pathlib import Path
import json,re,collections,hashlib,csv
from openpyxl import load_workbook
out=Path(__file__).parent; root=out.parents[1]
wb=load_workbook(root/'Assets/ConfigTables/Data/Narrative/Narrative.xlsx',read_only=True,data_only=False)
fields={'Object':['portraitPath','portraitIndoorPath','iconPath','avtarPath'],'ConversationPage':['bgOverride'],'ConversationLine':['speakerKey','side','emote','accentFxKey','accentTargetKey','stageLeftKey','stageRightKey','stageCenterKey'],'ReadingPage':['imageKey'],'Notice':['imageKey'],'MapPosition':['imgPath'],'OpeningPresentation':['visualsResource'],'CafeVisit':['seatedVisualPath','characterId']}
rowsout={}; diffs=[]
for name,cols in fields.items():
 ws=wb[name]; rows=list(ws.iter_rows(values_only=True)); headers=rows[0]; parsed=[]
 for rn,vals in enumerate(rows[4:],5):
  if not vals[1] or (vals[0] and str(vals[0]).startswith('##')): continue
  row={k:v for k,v in zip(headers[1:],vals[1:]) if k}; row['_excel_row']=rn; parsed.append(row)
 rowsout[name]=parsed
 table='tb'+('' if name.startswith('Opening') else 'ncs')+name.lower()
 data=json.loads((root/'Assets/Settings/Resources/Luban'/f'{table}.json').read_text(encoding='utf-8-sig')); idx={str(next(iter(r.values()))):r for r in data}
 for row in parsed:
  id=str(row[headers[1]]); other=idx.get(id)
  if other is None: diffs.append({'sheet':name,'id':id,'row':row['_excel_row'],'reason':'Excel有、JSON无'}); continue
  for k in cols:
   if (row.get(k) or '') != (other.get(k) or ''): diffs.append({'sheet':name,'id':id,'field':k,'excel':row.get(k),'json':other.get(k)})
(out/'excel_art_rows.json').write_text(json.dumps(rowsout,ensure_ascii=False,indent=2),encoding='utf-8')
(out/'excel_json_differences.json').write_text(json.dumps(diffs,ensure_ascii=False,indent=2),encoding='utf-8')
print('Counts', {k:len(v) for k,v in rowsout.items()}); print('DIFFS',json.dumps(diffs,ensure_ascii=False))
# Consolidated source expression tags, preserving exact labels and locations.
source=json.loads((out/'source_records.json').read_text(encoding='utf-8')); emotes=collections.defaultdict(lambda:collections.defaultdict(list))
for r in source:
 if not r['file'].endswith('.docx'): continue
 t=r['text']; names=re.match(r'^([^：:「⌈\[【\n]{1,35})[：:]',t)
 if names:
  for e in re.findall(r'[#＃]([^）)\s]+)',t): emotes[names[1].strip()][e].append(r['file']+'/'+r['loc'])
(out/'source_emotes.json').write_text(json.dumps(emotes,ensure_ascii=False,indent=2),encoding='utf-8')
print('SOURCE EMOTES')
for k,v in emotes.items(): print(k,','.join(v))
