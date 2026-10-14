import shutil
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from validate_player import validate


class PlayerValidationTests(unittest.TestCase):
    def setUp(self):
        root = Path(__file__).resolve().parent
        scratch = root.parents[1] / 'Library/Locus/tmp'
        scratch.mkdir(parents=True, exist_ok=True)
        self.tmp = tempfile.TemporaryDirectory(dir=scratch)
        self.root = Path(self.tmp.name)
        for relative in ['Data/Player/Player.xlsx', 'Data/Narrative/Narrative.xlsx']:
            target = self.root / relative
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(root / relative, target)

    def tearDown(self):
        self.tmp.cleanup()

    def mutate(self, sheet, column, value, row=5):
        path = self.root / 'Data/Player/Player.xlsx'
        book = load_workbook(path)
        fields = [c.value for c in book[sheet][1]]
        book[sheet].cell(row, fields.index(column) + 1, value)
        book.save(path)

    def test_real_tables(self):
        validate(self.root)

    def test_money_outside_origin_rejected(self):
        self.mutate('CreationOption', 'startingMoney', 100, 8)
        with self.assertRaisesRegex(AssertionError, 'ambiguous origin'):
            validate(self.root)

    def test_negative_action_max_rejected(self):
        self.mutate('PlayerConfig', 'actionPointMax', -1)
        with self.assertRaisesRegex(AssertionError, 'negative'):
            validate(self.root)

    def test_deleted_required_group_rejected(self):
        self.mutate('CreationGroup', 'groupId', 'unknown')
        with self.assertRaisesRegex(AssertionError, 'missing required'):
            validate(self.root)

    def test_duplicate_reward_rejected(self):
        path = self.root / 'Data/Player/Player.xlsx'
        book = load_workbook(path)
        book['CreationBonus'].append([cell.value for cell in book['CreationBonus'][5]])
        book.save(path)
        with self.assertRaisesRegex(AssertionError, 'duplicate reward'):
            validate(self.root)


if __name__ == '__main__':
    unittest.main()
