"""Import the cafe manuscript, then apply player-facing tutorial corrections.
Paragraph indices are zero-based document body paragraphs, including blank paragraphs.
Original manuscript directions remain under cafe.tutorial.source.* for traceability;
only OpeningLine-linked text is presented to players. Run GenerateAll.bat afterwards.
"""
from pathlib import Path
import re
import zipfile
import xml.etree.ElementTree as ET
import openpyxl

ROOT = Path(__file__).resolve().parent
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

def apply_player_facing_content(narrative, text):
    """Idempotent source-table corrections, also applied after a manuscript re-import."""
    corrections = {
        # 顾客B 已按《文案-新手教程.docx》P014/P019/P025 统一为女性（老太太／她／这位太太）；
        # 旧修订表里的男性版本会把这些行冲回去，不要恢复。
        'customer.name': 'B太太',
        'customer_arrives.01': '进来的是一个衣着普通的老太太，记不住脸。',
        'customer_arrives.06': '她看了我一眼，我回以微笑。',
        'perception_success.01': '老太太说这家店是乱套，而不用「落魄」做形容。',
        'perception_success.02': '我想她可能是叔叔的「老客人」，或者说他为数不多的朋友。',
        'perception_success.03': '我听过许多他朋友的故事，结果无一例外。没想到还有活着的人。',
        'perception_success.06': '阿黛尔，你和这位太太说说这两天你做了什么。',
        'perception_failure.01': '不知怎么，总感觉她身上的气息突然变了。',
        'perception_failure.02': '我没心思琢磨这位客人在想些什么。',
        'perception_failure.03': '还好她把视线收了回去。',
        'order_confirmed.01': '可能是看出老太太想点餐，阿黛尔递来一张自己无聊时手绘的餐单。',
        'order_confirmed.02': '大杯黑咖啡，不加糖，就这样。',
        'order_confirmed.03': 'B太太是新沸点咖啡的第一位客人。',
        'workstation_open.01': '接下来去工作台制作客人的咖啡。点击大厅右侧吧台处的箭头或吧台，进入工作台。',
        'make_coffee_base.01': '别紧张，店长，我会陪着你的。客人要的是大杯无糖黑咖啡，我们就从磨豆开始吧。',
        'make_coffee_base.02': '制作目标：研磨咖啡豆，用热水煮成黑咖啡，再装入杯中。不需要加糖。',
        'make_coffee_base.03': '工作台会一步步提示你。先学会取放原料，再试着启动设备。',
        'pour_into_cup.01': '咖啡做好了。准备一个空杯，把壶里的咖啡倒进去，就可以出餐啦。',
        'adjust_strength.01': '加清水会让咖啡变淡。这次先按当前配方完成黑咖啡，不必额外调浓度。',
        'adjust_sweetness.01': '客人要无糖黑咖啡，这次不需要加糖。',
    }
    for key, title in {
        'customer_arrives': '第一位客人', 'perception_success': '熟悉的老客人',
        'perception_failure': '难以捉摸的客人', 'order_confirmed': '确认订单',
        'workstation_open': '前往工作台', 'make_coffee_base': '制作黑咖啡',
        'pour_into_cup': '装杯', 'adjust_strength': '认识浓度',
        'adjust_sweetness': '确认甜度', 'serve_customer': '交付咖啡', 'ending': '新的开始',
    }.items():
        corrections[key + '.title'] = title
    rows = {row[1].value: row for row in text['Text'].iter_rows(min_row=5)}
    for suffix, value in corrections.items():
        key = 'cafe.tutorial.' + suffix
        if key not in rows:
            raise ValueError('Missing tutorial text key: ' + key)
        rows[key][2].value = value
        rows[key][3].value = '玩家教程修订；原稿保留在 cafe.tutorial.source.*；按键细节由工作台指引负责'
    customer_segments = {'customer_arrives', 'perception_success', 'perception_failure', 'order_confirmed'}
    for row in narrative['OpeningPresentation'].iter_rows(min_row=5):
        pid = str(row[1].value or '')
        if pid != 'first_day_cafe' and not pid.startswith('cafe_tutorial.'):
            continue
        row[5].value = 'Dialogue/CafeSceneVisuals' if pid == 'first_day_cafe' else 'Dialogue/CafeTutorialVisuals'
        segment = pid.removeprefix('cafe_tutorial.')
        row[6].value = 'cafe.tutorial.customer.name' if segment in customer_segments or segment.startswith('result_') else None


def configure():
    with zipfile.ZipFile(ROOT.parent / '文案-新手教程.docx') as archive:
        tree = ET.fromstring(archive.read('word/document.xml'))
    paragraphs = [''.join(t.text or '' for t in p.findall('.//w:t', NS))
                  for p in tree.findall('.//w:body/w:p', NS)]
    narrative = openpyxl.load_workbook(ROOT / 'Data/Narrative/Narrative.xlsx')
    text = openpyxl.load_workbook(ROOT / 'Data/Shared/Text.xlsx')
    tw = text['Text']
    def put(key, value, comment='文案-新手教程.docx；原文保留'):
        for row in tw.iter_rows(min_row=5):
            if row[1].value == key:
                row[2].value = value
                row[3].value = comment
                return
        tw.append([None, key, value, comment])
    def clear_rows(sheet, predicate):
        for i in range(sheet.max_row, 4, -1):
            if predicate([c.value for c in sheet[i]]): sheet.delete_rows(i)
    presentations = narrative['OpeningPresentation']
    lines = narrative['OpeningLine']
    clear_rows(presentations, lambda r: str(r[1]).startswith('cafe_tutorial.'))
    clear_rows(lines, lambda r: r[2] == 'first_day_cafe' or str(r[2]).startswith('cafe_tutorial.'))
    clear_rows(tw, lambda r: str(r[1]).startswith('opening.line.') or str(r[1]).startswith('cafe.tutorial.'))
    for i, value in enumerate(['centerNameKey', 'string', 'c', '第三角色显示名'], 1): presentations.cell(i, 7, value)
    put('cafe.tutorial.customer.name', '顾客B')
    put('opening.title', '重新启航')
    steps = []
    opening_line_seq = 0
    tutorial_line_seq = 0
    def segment(step, indices, next_step='', focus='', optional=False, title=None):
        nonlocal opening_line_seq, tutorial_line_seq
        pid = 'first_day_cafe' if step == 'introduction' else 'cafe_tutorial.' + step
        title_key = 'opening.title' if step == 'introduction' else 'cafe.tutorial.' + step + '.title'
        if step != 'introduction': put(title_key, title or step)
        existing = any(row[1] == pid for row in presentations.iter_rows(min_row=5, values_only=True))
        if not existing:
            presentations.append([None, pid, title_key, 'opening.player.name', 'opening.adele.name',
                                  'Dialogue/FirstDayIntroductionVisuals', 'cafe.tutorial.customer.name'])
        seq = 0
        for index in indices:
            raw = paragraphs[index] if isinstance(index, int) else index
            speaker = 'narrator'
            if raw.startswith('阿黛尔·米尼耶：'): speaker, raw = 'right', raw.split('：', 1)[1]
            elif raw.startswith('{playername}：'): speaker, raw = 'left', raw.split('：', 1)[1]
            elif raw.startswith('顾客B：'): speaker, raw = 'center', raw.split('：', 1)[1]
            # Stage directions remain in the source text registry, not spoken dialogue.
            spoken = re.sub(r'^（#?[^）]+）', '', raw)
            # Only real speaker paragraphs use 「」 as quotation marks. In narration the
            # corner brackets are emphasis, so splitting would reduce the line to a bare
            # word (e.g. 落魄 / 老客人) and drop the rest of the sentence.
            parts = re.findall('「([^」]*)」', spoken) if speaker != 'narrator' else []
            if not parts: parts = [spoken]
            for part in parts:
                seq += 1
                if step == 'introduction':
                    opening_line_seq += 1
                    line_id = f'opening_{opening_line_seq:03}'
                    key = f'opening.line.{opening_line_seq:02}'
                else:
                    tutorial_line_seq += 1
                    line_id = f'cafe_tutorial_{tutorial_line_seq:03}'
                    key = f'cafe.tutorial.{step}.{seq:02}'
                put(key, part, f'原稿段落 {index}' if isinstance(index, int) else '原稿枚举展开；未添加玩法')
                lines.append([None, line_id, pid, seq, speaker, key, 'none', 0])
        steps.append([None, step, 'first_cafe_tutorial', pid, focus,
                      'workstation_opened' if step == 'workstation_open' else '', next_step, optional])
    segment('introduction', list(range(9,14)), 'customer_arrives')
    segment('customer_arrives', list(range(15,21)), 'perception_check')
    segment('perception_success', [22,23,25,26,27,28,29], 'order_confirmed')
    segment('perception_failure', [31,32,33], 'order_confirmed')
    segment('order_confirmed', [35,36,37,38], 'workstation_open')
    segment('workstation_open', [40], 'make_coffee_base', 'cafe.workstation')
    segment('make_coffee_base', [45,46,51], 'pour_into_cup', 'cafe.coffee_base')
    segment('pour_into_cup', [51], 'adjust_strength', 'cafe.cup')
    segment('adjust_strength', [52], 'adjust_sweetness', 'cafe.strength')
    segment('adjust_sweetness', [52,53], 'serve_customer', 'cafe.sweetness')
    for key, label in [('grind','研磨'),('boil','烧水'),('brew','冲泡'),('cup','装杯'),('base','基底'),('strength','调浓度'),('sweetness','甜度'),('serve','出餐')]:
        segment('demo_' + key, [label], optional=True, title=label)
    segment('serve_customer', [67], 'result', 'cafe.delivery')
    for key,label,reward in [('excellent','独一无二','金钱+200'),('good','滴滴香浓','金钱+150'),('average','中庸之道','金钱+50'),('bad','食不知味','金钱+10'),('refund','好想退钱','金钱-200，打回重做')]:
        segment('result_' + key, ['顾客B：「这个咖啡嘛……简直就是' + label + '！」', reward], 'make_coffee_base' if key == 'refund' else 'ending', title=label)
    segment('ending', [70,71,73,74,75,76])
    # Preserve all author directions, duplicate order copy and ambiguous sentences for traceability.
    for i, paragraph in enumerate(paragraphs):
        if paragraph: put(f'cafe.tutorial.source.p{i:03}', paragraph, '原稿完整归档（含舞台说明）；不自动播放')
    if 'CafeTutorialStep' in narrative: del narrative['CafeTutorialStep']
    ws = narrative.create_sheet('CafeTutorialStep')
    ws.append(['##var','stepId','tutorialId','presentationId','focusTargetKey','completionSignalId','nextStepId','optionalDemo'])
    ws.append(['##type','string','string','string','string','string','string','bool'])
    ws.append(['##group'] + ['c'] * 7)
    ws.append(['##','步骤','教程','NCS段落','焦点','完成信号','后续步骤或分支标记','可选演示'])
    for row in steps: ws.append(row)
    # Existing NCS ability check owns selection; branch events become silent routing events.
    for row in narrative['Event'].iter_rows(min_row=5):
        if row[1].value in ('EV_FIRST_DAY_PERCEPTION_SUCCESS','EV_FIRST_DAY_PERCEPTION_FAILURE'):
            row[6].value = None
    for row in narrative['SetupParams'].iter_rows(min_row=5):
        if row[1].value in ('first_day_perception_success','first_day_perception_failure'): row[3].value = None
    apply_player_facing_content(narrative, text)
    narrative.save(ROOT / 'Data/Narrative/Narrative.xlsx')
    text.save(ROOT / 'Data/Shared/Text.xlsx')
    print(f'Configured {len(steps)} tutorial segments; manuscript archive preserved.')

if __name__ == '__main__': configure()
