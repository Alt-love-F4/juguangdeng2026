"""Export original PSD pixels and synchronize map art coordinates in the existing table."""
from pathlib import Path
import csv
import json
import numpy as np
from scipy import ndimage
from psd_tools import PSDImage
from PIL import Image
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
PSD = PSDImage.open(ROOT / 'Assets/Art/地图.psd')
LAYERS = {layer.name: layer for layer in PSD}
OUT = ROOT / 'Assets/Resources/Art/Map'
OUT.mkdir(parents=True, exist_ok=True)
for source, name in [('地图底图', 'background'), ('指北', 'compass'), ('圣诺特城', 'title'), ('其他位置图钉', 'pin')]:
    LAYERS[source].topil().save(OUT / (name + '.png'))

def regions(layer):
    labels, _ = ndimage.label(np.array(layer.topil())[:, :, 3] > 20)
    result = []
    for i, s in enumerate(ndimage.find_objects(labels), 1):
        if s and (labels[s] == i).sum() > 500:
            result.append((s[1].start + layer.left, s[0].start + layer.top,
                           s[1].stop + layer.left, s[0].stop + layer.top))
    return sorted(result, key=lambda b: (b[1], b[0]))

buildings = regions(LAYERS['地点图标'])
signs = regions(LAYERS['标牌'])
assert len(buildings) == len(signs) == 9
names = ['旧城区', '圣西尔万郊区', '圣徒街', '大堂区', '圣地大教堂', '首都法院', '市政厅', '灰墙学院', '花园区']
keys = ['temporary', 'saint_sylvan', 'saint_street', 'cafe', 'market', 'court', 'city_hall', 'library', 'garden']
records = []
for i, (building, sign, name, key) in enumerate(zip(buildings, signs, names, keys), 1):
    # Include the full sign and disconnected lettering within its rectangle.
    x0 = max(0, min(building[0], sign[0]) - 20)
    y0 = max(0, min(building[1], sign[1]) - 20)
    x1 = min(PSD.width, max(building[2], sign[2]) + 20)
    y1 = min(PSD.height, max(building[3], sign[3]) + 20)
    bounds = (x0, y0, x1, y1)
    image = Image.new('RGBA', (x1-x0, y1-y0))
    for layer_name in ['地点图标', '标牌', '地点名称']:
        layer = LAYERS[layer_name]
        crop = layer.topil().crop((x0-layer.left, y0-layer.top, x1-layer.left, y1-layer.top))
        image.alpha_composite(crop)
    image.save(OUT / f'location_{i:02d}.png')
    records.append({'key': key, 'name': name, 'x': ((x0+x1)/2-1920)*.625,
                    'y': (1080-(y0+y1)/2)*.625, 'path': f'Art/Map/location_{i:02d}', 'bounds': bounds})

book_path = ROOT / 'Assets/ConfigTables/Data/Narrative/Narrative.xlsx'
book = openpyxl.load_workbook(book_path)
sheet = book['MapPosition']
columns = {cell.value: cell.column for cell in sheet[1] if cell.value}
existing = {sheet.cell(r, columns['positionKey']).value: r for r in range(5, sheet.max_row+1)}
for record in records:
    row = existing.get(record['key'])
    if row is None:
        row = sheet.max_row + 1
        for field, value in {'positionKey': record['key'], 'nameText': record['name'], 'isImportant': False,
                             'defaultAction': '', 'extendedActions': '', 'enabled': True}.items():
            sheet.cell(row, columns[field], value)
    for field, value in {'posX': record['x'], 'posY': record['y'], 'imgPath': record['path']}.items():
        sheet.cell(row, columns[field], value)
sheet.cell(4, columns['posX'], '地图中心原点，向右为正；PSD像素转换比例0.625')
sheet.cell(4, columns['posY'], '地图中心原点，向上为正；PSD像素转换比例0.625')
sheet.cell(4, columns['imgPath'], 'Resources下Sprite路径，不含扩展名；图片原尺寸×0.625')
book.save(book_path)
# Export only the changed table. Preserve all other generated data/code.
types = {sheet.cell(1,c).value: sheet.cell(2,c).value for c in range(2,sheet.max_column+1)}
data = []
for r in range(5, sheet.max_row+1):
    if not sheet.cell(r, columns['positionKey']).value: continue
    entry = {}
    for field, c in columns.items():
        if field == '##var': continue
        value = sheet.cell(r,c).value
        entry[field] = (bool(value) if types[field]=='bool' else float(value or 0) if types[field]=='float' else str(value or ''))
    data.append(entry)
(ROOT/'Assets/Settings/Resources/Luban/tbncsmapposition.json').write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
manifest = ROOT/'Assets/Art/Gameplay/DayFlow/MapLocations/location_manifest.csv'
with manifest.open('w',encoding='utf-8',newline='') as f:
    writer=csv.writer(f); writer.writerow(['positionKey','display_name','resource_path','map_x','map_y','psd_x','psd_y','psd_width','psd_height'])
    for r in records:
        x0,y0,x1,y1=r['bounds']; writer.writerow([r['key'],r['name'],r['path'],r['x'],r['y'],x0,y0,x1-x0,y1-y0])
print('Exported 9 building/sign/name sprites and 4 map layers; synchronized MapPosition table.')
for r in records: print(r)
