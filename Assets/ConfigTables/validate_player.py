"""Player semantic checks in addition to Luban strict references and duplicate keys."""
from pathlib import Path
from openpyxl import load_workbook


def rows(book, sheet):
    values = list(book[sheet].values)
    return [dict(zip(values[0][1:], row[1:])) for row in values[4:] if row[1] is not None]


def validate(root):
    player = load_workbook(root / 'Data/Player/Player.xlsx', data_only=True)
    narrative = load_workbook(root / 'Data/Narrative/Narrative.xlsx', data_only=True)
    config = rows(player, 'PlayerConfig')
    assert len(config) == 1 and config[0]['playerId'] == 'player', 'PlayerConfig: exactly player required'
    assert config[0]['baseStartingMoney'] >= 0 and config[0]['actionPointMax'] >= 0, 'PlayerConfig: negative money/max'
    groups = rows(player, 'CreationGroup')
    assert {g['groupId'] for g in groups} == {'origin', 'temperament', 'memory', 'education', 'trait'}, 'CreationGroup: missing required group'
    assert all(g['required'] for g in groups), 'CreationGroup: all five groups required'
    assert len({g['seq'] for g in groups}) == len(groups), 'CreationGroup: duplicate sequence'
    options = rows(player, 'CreationOption')
    assert len({o['optionId'] for o in options}) == len(options), 'CreationOption: duplicate ID'
    for group in groups:
        choices = [o for o in options if o['groupId'] == group['groupId']]
        assert choices and len({o['seq'] for o in choices}) == len(choices), 'CreationOption: empty group/duplicate sequence'
    for option in options:
        assert option['startingMoney'] >= -1, 'CreationOption: illegal money'
        assert option['groupId'] == 'origin' or (not option['originId'] and option['startingMoney'] == -1), 'CreationOption: ambiguous origin/money'
        assert not option['specializationConditionKey'] or option['groupId'] == 'education', 'CreationOption: specialization outside education'
    origins = {o['originId']: o['startingMoney'] for o in options if o['groupId'] == 'origin'}
    assert set(origins) == {'waif', 'commoner', 'noble'} and 0 <= origins['waif'] < origins['commoner'] < origins['noble'], 'CreationOption: invalid origin money order'
    attributes = rows(player, 'PlayerAttribute')
    assert {a['attributeId'] for a in attributes} == {'language', 'courage', 'wisdom', 'perception'}, 'PlayerAttribute: missing formal attribute'
    bars = {b['ownerId'] + '.' + b['barId']: b for b in rows(narrative, 'Bar')}
    for attr in attributes:
        bar = bars[attr['barKey']]
        assert bar['sourceMode'] in (None, '', 'ncs') and bar['advanceMode'] != 'derived', 'PlayerAttribute: must reference internal Bar'
    bonuses = rows(player, 'CreationBonus')
    assert len({(b['optionId'], b['attributeId']) for b in bonuses}) == len(bonuses), 'CreationBonus: duplicate reward'


if __name__ == '__main__':
    validate(Path(__file__).resolve().parent)
    print('Player semantic validation passed.')
