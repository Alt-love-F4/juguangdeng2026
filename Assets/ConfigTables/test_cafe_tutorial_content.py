"""Source/generated parity and manuscript re-import regression checks (no Unity required)."""
import json
from pathlib import Path
import unittest

import openpyxl
from configure_cafe_tutorial import apply_player_facing_content

ROOT = Path(__file__).resolve().parent
GENERATED = ROOT.parent / 'Settings/Resources/Luban'


def records(sheet):
    headers = [cell.value for cell in sheet[1]][1:]
    return [dict(zip(headers, row[1:])) for row in sheet.iter_rows(min_row=5, values_only=True)
            if row[1] is not None]


class CafeTutorialContentTests(unittest.TestCase):
    def setUp(self):
        self.narrative = openpyxl.load_workbook(ROOT / 'Data/Narrative/Narrative.xlsx')
        self.text = openpyxl.load_workbook(ROOT / 'Data/Shared/Text.xlsx')

    def tearDown(self):
        self.narrative.close()
        self.text.close()

    def generated(self, table):
        return json.loads((GENERATED / (table + '.json')).read_text(encoding='utf-8-sig'))

    def test_tutorial_source_and_generated_data_match(self):
        for sheet, table, key in [('OpeningPresentation', 'tbopeningpresentation', 'presentationId'),
                                  ('OpeningLine', 'tbopeningline', 'lineId')]:
            source = records(self.narrative[sheet])
            generated = {row[key]: row for row in self.generated(table)}
            for row in source:
                if row.get('presentationId') != 'first_day_cafe' and not str(row.get('presentationId', '')).startswith('cafe_tutorial.'):
                    continue
                for field, value in row.items():
                    self.assertEqual(value if value is not None else '', generated[row[key]][field], (row[key], field))
        generated_text = {row['key']: row for row in self.generated('tbtext')}
        for row in records(self.text['Text']):
            if row['key'].startswith('cafe.tutorial.'):
                self.assertEqual(row['text'] or '', generated_text[row['key']]['text'], row['key'])

    def test_reimport_corrections_are_idempotent_and_preserve_archive_and_headers(self):
        headers = [[tuple(cell.value for cell in row) for row in book[sheet].iter_rows(max_row=4)]
                   for book, sheet in [(self.narrative, 'OpeningPresentation'), (self.text, 'Text')]]
        archive = [row for row in records(self.text['Text']) if row['key'].startswith('cafe.tutorial.source.')]
        apply_player_facing_content(self.narrative, self.text)
        first = (records(self.narrative['OpeningPresentation']), records(self.text['Text']))
        apply_player_facing_content(self.narrative, self.text)
        self.assertEqual(first, (records(self.narrative['OpeningPresentation']), records(self.text['Text'])))
        self.assertEqual(archive, [row for row in records(self.text['Text']) if row['key'].startswith('cafe.tutorial.source.')])
        self.assertEqual(headers, [[tuple(cell.value for cell in row) for row in book[sheet].iter_rows(max_row=4)]
                                  for book, sheet in [(self.narrative, 'OpeningPresentation'), (self.text, 'Text')]])

    def test_player_text_has_no_internal_directions_or_gender_mismatch(self):
        texts = {row['key']: row['text'] or '' for row in records(self.text['Text'])}
        for row in records(self.narrative['OpeningLine']):
            if not row['presentationId'].startswith('cafe_tutorial.'):
                continue
            text = texts[row['textKey']]
            self.assertFalse(row['textKey'].startswith('cafe.tutorial.source.'))
            for marker in ['演出说明', '配料表', '太太', '→']:
                self.assertNotIn(marker, text, row['textKey'])
            if row['speaker'] == 'right':
                for marker in ['[', ']', '#', '鼠标', '左键', '右键', '高亮']:
                    self.assertNotIn(marker, text, row['textKey'])
                self.assertLess(len(text), 110, row['textKey'])
        entry = texts['cafe.tutorial.workstation_open.01']
        self.assertNotIn('方框', entry)
        self.assertIn('吧台', entry)
        self.assertNotIn('挂牌', entry)
        self.assertIn('箭头或吧台', entry)
        self.assertIn('男人', texts['cafe.tutorial.customer_arrives.01'])
        self.assertIn('大杯', texts['cafe.tutorial.order_confirmed.02'])
        self.assertIn('不加糖', texts['cafe.tutorial.order_confirmed.02'])


if __name__ == '__main__':
    unittest.main()
