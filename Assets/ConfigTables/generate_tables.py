#!/usr/bin/env python3
"""
等效 Luban 生成器：读取 ConfigTables/Data/Simulation 下的玩法 xlsx，输出 JSON 到 Assets/Settings/Resources/Luban/
规则：
  - 表头遵循 Luban 约定：row1=##var, row2=##type, row3=##group, row4=##, 数据从 row5 起
  - 输出文件名：tb{table_name_lower}.json，JSON 数组，字段名与 ##var 一致
  - 支持类型：string / int / float / bool
  - bool 在 xlsx 中填 TRUE/FALSE 或 true/false
"""
import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent
DATA_DIR = ROOT / "Data" / "Simulation"
OUT_DIR = ROOT.parent / "Settings" / "Resources" / "Luban"

def _to_value(type_name: str, value):
    if value is None:
        if type_name == "string":
            return ""
        if type_name == "int":
            return 0
        if type_name == "float":
            return 0.0
        if type_name == "bool":
            return False
        return value
    if type_name == "string":
        return str(value)
    if type_name == "int":
        return int(float(value))
    if type_name == "float":
        return float(value)
    if type_name == "bool":
        if isinstance(value, bool):
            return value
        return str(value).strip().lower() in ("true", "1", "yes")
    return value


def read_table(xlsx_path: Path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(min_row=1, max_row=5, values_only=True))
    if len(rows) < 4:
        raise RuntimeError(f"{xlsx_path.name}: sheet too short")
    var_row = rows[0]
    type_row = rows[1]
    if not var_row or str(var_row[0]).strip() != "##var":
        raise RuntimeError(f"{xlsx_path.name}: row1 col1 must be ##var")

    fields = []
    for i in range(1, len(var_row)):
        name = var_row[i]
        if name is None:
            continue
        type_name = type_row[i] if i < len(type_row) else "string"
        fields.append((i, str(name), str(type_name).strip().lower()))

    table_name = xlsx_path.stem.lower()
    out_name = f"tb{table_name}.json"
    items = []
    for row in ws.iter_rows(min_row=5, values_only=True):
        # skip entirely empty rows
        if not any(v is not None for v in row):
            continue
        item = {}
        for idx, name, type_name in fields:
            value = row[idx] if idx < len(row) else None
            item[name] = _to_value(type_name, value)
        items.append(item)

    wb.close()
    return out_name, items


def validate(tables):
    for name, rows in tables.items():
        ids = [row.get("id", row.get("key")) for row in rows]
        if any(not key for key in ids) or len(ids) != len(set(ids)):
            raise ValueError(f"{name}: empty or duplicate id/key")
    entities = {r["id"]: r for r in tables.get("tbworldobject.json", [])}
    ingredients = {r["id"] for r in tables["tbingredient.json"]}
    devices = {r["id"] for r in tables["tbdevice.json"]}
    recipes = {r["id"] for r in tables["tbrecipe.json"]}
    operations = {"SelectCup", "AddIngredient", "BeanSelection", "GrindClick", "Process", "CocoaMinigame", "LatteArt", "Serve"}
    for entity in entities.values():
        if entity["role"] not in {"item", "source", "slot", "machine", "npc"} or entity["capacity"] < 1 or entity["height"] <= 0:
            raise ValueError(f"Invalid entity: {entity['id']}")
        if entity["role"] == "source" and (entity["item_id"] not in entities or entities[entity["item_id"]]["role"] != "item"):
            raise ValueError(f"Unknown source item: {entity['id']}")
        if entity["role"] == "machine" and entity["behavior_id"] not in devices:
            raise ValueError(f"Unknown machine: {entity['id']}")
        if "ingredient" in entity["tags"].split("|") and entity["behavior_id"] not in ingredients:
            raise ValueError(f"Unknown ingredient: {entity['id']}")
        if entity["sprite_path"] and not (ROOT.parent.parent / entity["sprite_path"]).is_file():
            raise ValueError(f"Missing sprite: {entity['sprite_path']}")
    grouped = {}
    for step in tables.get("tbrecipestep.json", []):
        if step["recipe_id"] not in recipes or step["operation"] not in operations or step["amount"] <= 0:
            raise ValueError(f"Invalid recipe step: {step['id']}")
        if step["operation"] == "AddIngredient" and step["target_id"] not in ingredients:
            raise ValueError(f"Unknown ingredient in step: {step['id']}")
        if step["device_id"] and step["device_id"] not in devices:
            raise ValueError(f"Unknown device in step: {step['id']}")
        grouped.setdefault(step["recipe_id"], []).append(step)
    for recipe in recipes:
        steps = sorted(grouped.get(recipe, []), key=lambda r: r["sequence"])
        if not steps or [s["sequence"] for s in steps] != list(range(1, len(steps) + 1)) or steps[0]["operation"] != "SelectCup" or steps[-1]["operation"] != "Serve":
            raise ValueError(f"Missing, duplicated or unordered steps: {recipe}")


def main():
    xlsx_files = sorted(DATA_DIR.glob("*.xlsx"))
    if not xlsx_files:
        print("No xlsx found in Data/Simulation/")
        return 1
    tables = dict(read_table(xlsx) for xlsx in xlsx_files)
    validate(tables)
    if "--validate" in sys.argv:
        print("[OK] All table IDs, references and recipe steps are valid")
        return 0
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for name, items in tables.items():
        out_path = OUT_DIR / name
        temporary = out_path.with_suffix(".json.tmp")
        temporary.write_text(json.dumps(items, ensure_ascii=False, indent=2), encoding="utf-8")
        temporary.replace(out_path)
        print(f"[OK] {name}: {len(items)} rows")
    return 0


if __name__ == "__main__":
    sys.exit(main())
