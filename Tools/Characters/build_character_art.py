"""Build NCS portrait/avatar/seated placeholders from untouched PSD PNG exports.
Run from project root: python Tools/Characters/build_character_art.py
Requires Pillow and openpyxl. Does not edit Unity YAML or source workbooks.
"""
from pathlib import Path
import hashlib
import json
import zipfile
import xml.etree.ElementTree as ET
from PIL import Image, ImageDraw, ImageFont, ImageOps
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'Assets/Resources/Characters'
SCRATCH = ROOT / 'Library/Locus/tmp/character-art'
BATCH1 = 'Assets/Art/Gameplay/人物-常规色'
BATCH2 = 'Assets/Art/Gameplay/人物2-常规色 (2)'
IDS = 'laurent clovis renaud elion alienor queen bertrand sade mamion lucien lesou sylvain antoine lefebvre duval adele emmanuel sparrow thibaut octave eugene colette pascal marie guest_a guest_b guest_c guest_d soldier_a soldier_b soldier_c soldier_d'.split()
SOURCES = {
 'laurent': (BATCH2, '洛朗.png'), 'clovis': (BATCH2, '克洛维·康多尔 皇帝.png'),
 'elion': (BATCH1, '埃利昂.康多尔.png'), 'alienor': (BATCH1, '阿莉耶诺尔(公主.png'),
 'bertrand': (BATCH1, '贝特朗（剪影.png'), 'sade': (BATCH1, '萨德神父.png'),
 'mamion': (BATCH2, '马米翁 大主教.png'), 'lucien': (BATCH2, '吕西安 神父.png'),
 'lesou': (BATCH1, '勒苏.png'), 'sylvain': (BATCH2, '西尔万.瓦隆 教授.png'),
 'antoine': (BATCH1, '安托万富商.png'), 'lefebvre': (BATCH2, '勒费弗尔.png'),
 'adele': (BATCH1, '阿黛尔.png'), 'emmanuel': (BATCH2, '埃马纽埃尔·索雷尔医生.png'),
 'sparrow': (BATCH2, '麻雀.png'), 'thibaut': (BATCH1, '蒂博·卡隆.png'),
 'eugene': (BATCH1, '欧仁.png'), 'colette': (BATCH2, '柯莱特·布吕内.png'),
 'pascal': (BATCH2, '帕斯卡尔·勒鲁.png'), 'marie': (BATCH2, '玛丽 女扮男.png'),
}
# Source-space square crop (left, top, side), fractional coordinates in original PNG.
# Kept explicit/reviewable: no guessed runtime cropping or face-detection dependency.
CROPS = {
 'laurent': (.24,0,.50), 'clovis': (.20,0,.60), 'elion': (.20,0,.60),
 'alienor': (.27,.10,.60), 'bertrand': (.39,0,.49), 'sade': (.18,0,.65),
 'mamion': (.24,0,.52), 'lucien': (.03,0,.97), 'lesou': (.12,0,.83),
 'sylvain': (.18,0,.64), 'antoine': (.27,0,.60), 'lefebvre': (.18,0,.53),
 'adele': (.12,0,.78), 'emmanuel': (.20,0,.65), 'sparrow': (.05,0,.90),
 'thibaut': (.10,0,.80), 'eugene': (.05,0,.90), 'colette': (0,0,1),
 'pascal': (.20,0,.55), 'marie': (.12,0,.78),
}
FONT = 'C:/Windows/Fonts/msyh.ttc'

def font(size):
 return ImageFont.truetype(FONT, size)

def label(draw, xy, text, size=32, fill='#f5ead5'):
 draw.text(xy, text, font=font(size), fill=fill, anchor='mm')

def save_image(im, path):
 path.parent.mkdir(parents=True, exist_ok=True)
 im.save(path)

def copy_image(src, dest):
 dest.parent.mkdir(parents=True, exist_ok=True)
 data = src.read_bytes()
 if not dest.exists() or dest.read_bytes() != data:
  dest.write_bytes(data)

def silhouette(source, size):
 # Reuse approved silhouette shapes without labels, frames or invented character art.
 art = Image.open(ROOT / source).convert('RGBA')
 art = art.crop(art.getchannel('A').getbbox())
 art = ImageOps.contain(art, (size[0] - 40, size[1] - 40), Image.Resampling.LANCZOS)
 black = Image.new('RGBA', art.size, (0, 0, 0, 0))
 black.putalpha(art.getchannel('A'))
 canvas = Image.new('RGBA', size)
 canvas.alpha_composite(black, ((size[0] - art.width) // 2, size[1] - art.height - 20))
 return canvas

def placeholder(cid, name):
 return silhouette('Assets/Art/NarrativeGenerated/Characters/CH35_tutorial_elderly_woman_silhouette.png', (768,1536))

def avatar(im,cid):
 if cid not in SOURCES:
  shape = im.crop(im.getchannel('A').getbbox())
  return ImageOps.fit(shape.crop((0,0,shape.width,max(1,int(shape.height*.40)))), (512,512), Image.Resampling.LANCZOS)
 if cid in SOURCES:
  w,h=im.size
  x,y,s=CROPS.get(cid,(0.0,0.0,1.0))
  side=int(w*s)
  crop=im.crop((int(w*x),int(h*y),int(w*x)+side,int(h*y)+side))
 else:
  crop=im.crop((120,160,648,688))
 crop=crop.resize((512,512),Image.Resampling.LANCZOS)
 bg=Image.new('RGBA',(512,512),'#e3d4bb'); bg.alpha_composite(crop)
 d=ImageDraw.Draw(bg); d.rounded_rectangle((5,5,506,506),radius=18,outline='#806c50',width=8)
 return bg

def seated(cid,name,av):
 return silhouette('Assets/Art/FreeCafe/Seating/Characters/silhouette_seated.png', (768,1024))

def rebuild_placeholders():
 # Preserve formal art, variant art, GUIDs and all existing importer settings.
 manifest = json.loads((ROOT/'Tools/Characters/character-art-manifest.json').read_text(encoding='utf-8'))
 count = 0
 for record in manifest:
  folder = OUT / record['id']
  if record['portraitPlaceholder']:
   im = placeholder(record['id'], record['name'])
   save_image(im, folder/'neutral.png')
   save_image(avatar(im,record['id']), folder/'avatar.png')
   count += 2
  if record['seatedPlaceholder']:
   save_image(seated(record['id'],record['name'],None), folder/'seated_placeholder.png')
   count += 1
 print(f'Rebuilt {count} NPC placeholder images as black silhouettes; formal/variant art and .meta files untouched.')

def contact_sources():
 SCRATCH.mkdir(parents=True,exist_ok=True)
 sheet=Image.new('RGB',(5*300,4*460),'#d2d0c9')
 for i,(cid,(directory,filename)) in enumerate(SOURCES.items()):
  im=Image.open(ROOT/directory/filename).convert('RGBA')
  thumb=ImageOps.contain(im,(280,410))
  x=(i%5)*300;y=(i//5)*460
  sheet.paste(thumb,(x+(300-thumb.width)//2,y),thumb)
  label(ImageDraw.Draw(sheet),(x+150,y+435),cid,24,fill='#252525')
 save_image(sheet,SCRATCH/'sources.jpg')

def build():
 ws=load_workbook(ROOT/'策划文件/策划/文案/人物.xlsx',read_only=True,data_only=False)['人物']
 names=[ws.cell(row,1).value for row in range(3,35)]
 names[23]='雅克·马尚'
 assert len(names)==len(IDS)==32
 manifest=[]
 preview=Image.new('RGB',(8*420,4*580),'#c9c5ba')
 for i,(cid,name) in enumerate(zip(IDS,names)):
  folder=OUT/cid; folder.mkdir(parents=True,exist_ok=True)
  if cid in SOURCES:
   directory,filename=SOURCES[cid]; src=ROOT/directory/filename
   copy_image(src,folder/'neutral.png'); im=Image.open(src).convert('RGBA')
  else:
   src=None; im=placeholder(cid,name);save_image(im,folder/'neutral.png')
  av=avatar(im,cid);save_image(av,folder/'avatar.png')
  seat=seated(cid,name,av);save_image(seat,folder/'seated_placeholder.png')
  x=(i%8)*420;y=(i//8)*580
  for art,rect in [(im,(x,y,190,450)),(av,(x+205,y+10,200,200)),(seat,(x+205,y+230,200,270))]:
   thumb=ImageOps.contain(art,(rect[2],rect[3]));preview.paste(thumb,(rect[0]+(rect[2]-thumb.width)//2,rect[1]),thumb)
  label(ImageDraw.Draw(preview),(x+210,y+535),cid,25,fill='#252525')
  manifest.append({'id':cid,'name':name,'source':src.relative_to(ROOT).as_posix() if src else None,'portraitPlaceholder':src is None,'avatarCrop':CROPS.get(cid,[0,0,1]) if src else None,'portraitPath':f'Characters/{cid}/neutral','avtarPath':f'Characters/{cid}/avatar','portraitIndoorPath':f'Characters/{cid}/seated_placeholder','seatedPlaceholder':True})
 for cid,source,dest in [('alienor','阿莉耶诺尔(日常.png','indoor.png'),('bertrand','贝特朗.png','revealed.png')]:
  copy_image(ROOT/BATCH1/source,OUT/cid/dest)
 save_image(preview,SCRATCH/'generated-contact.jpg')
 (ROOT/'Tools/Characters/character-art-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
 print(f'Built {len(manifest)} character triples: {len(SOURCES)} authored, {len(manifest)-len(SOURCES)} placeholder portraits. Original PSD/Prefab exports untouched.')

if __name__=='__main__':
 import argparse
 parser=argparse.ArgumentParser();parser.add_argument('--contact-sources',action='store_true');parser.add_argument('--placeholders-only',action='store_true');args=parser.parse_args()
 if args.contact_sources: contact_sources()
 elif args.placeholders_only: rebuild_placeholders()
 else: build()
