"""Stage B-rose source-mapped Text rows; never activates an event.
Run from project root with --event 1009|1011|1012|1016.
Generation must subsequently use Assets/ConfigTables/GenerateAll.bat.
"""
from __future__ import annotations
import argparse
from copy import copy
from datetime import datetime
import hashlib
import json
from pathlib import Path
import re
import shutil
import xml.etree.ElementTree as ET
import zipfile

import openpyxl
from openpyxl.comments import Comment

ROOT = Path(__file__).resolve().parents[2]
MAP = ROOT / 'Docs/superpowers/specs/2026-10-16-content-batches/B-rose/source-map.md'
TEXT = ROOT / 'Assets/ConfigTables/Data/Shared/Text.xlsx'
SOURCES = {
    '1009': ('MA10.09-S1.docx', 76, '5ab8f53a3139abd6b0598d0dc5a78687095931b37aab500a3e74600c38f217b2'),
    '1011': ('MA70.10.11-S1.docx', 140, '97634d4d7ce507c87760d2def53bfd6bbf0d69e72b8597c0ec4253cda32920a0'),
    '1012': ('MA70.10.12-S1.docx', 86, 'f0c61782db5e85c876b716bcb038ec619e1492f117aaab6862de038888f01d4e'),
    '1016': ('MA70.10.16-S1.docx', 17, '38beea960ed414579b1b094a617f3ecd9f0b8238d01b732840ed0b39b78176a1'),
}
MIXED = {'1009': {36, 38, 43, 51}, '1011': {52, 85, 87, 110, 116, 117, 122, 125}, '1012': {4, 9, 67}, '1016': set()}
PREFIX = re.compile(r'^(?:\{playername\}|男人|布塔克|老爷|阿黛尔|麻雀|西尔万|埃莉诺|萨德)\s*[：:]\s*')
EMOTE = re.compile(r'[（(]\s*#.*?[）)]')


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def source_rows(event):
    rows = []
    for line in MAP.read_text(encoding='utf-8').splitlines():
        if not re.match(rf'^\| B-{event}-P\d{{3}} \|', line):
            continue
        cells = [s.strip() for s in line.strip('|').split('|')]
        if len(cells) != 6:
            continue
        number = int(re.search(r'P(\d+)', cells[0]).group(1))
        # Remove exactly the Markdown table's single padding space; preserve source indentation.
        source_cell = line.split('|')[6]
        if source_cell.startswith(' '):
            source_cell = source_cell[1:]
        if source_cell.endswith(' '):
            source_cell = source_cell[:-1]
        original = source_cell.replace('&#9;', '\t').replace('<br>', '\n')
        rows.append((number, cells[1], cells[2].strip('`'), cells[4].strip('`'), original))
    assert len(rows) == SOURCES[event][1]
    assert [r[0] for r in rows] == list(range(1, len(rows) + 1))
    return rows


def verify_docx(event, rows):
    filename, count, expected = SOURCES[event]
    path = Path('D:/桌面/文案/新玫瑰') / filename
    if not path.is_file():
        raise RuntimeError(f'Original source unavailable: {path}; refusing to stage unverified text')
    assert digest(path) == expected, f'Original changed: {path}'
    with zipfile.ZipFile(path) as archive:
        xml = ET.fromstring(archive.read('word/document.xml'))
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    paragraphs = []
    for p in xml.findall('.//w:body//w:p', ns):
        text = ''.join(n.text or '' if n.tag.endswith('}t') else '\t' if n.tag.endswith('}tab') else '\n' if n.tag.endswith(('}br', '}cr')) else '' for n in p.iter())
        if text.strip():
            paragraphs.append(text)
    assert len(paragraphs) == count
    # Markdown cell padding cannot distinguish source trailing layout spaces.
    assert [p.rstrip() for p in paragraphs] == [r[4].rstrip() for r in rows], 'DOCX/source-map mismatch'
    return str(path)


def clean(text, choice=False):
    text = text.strip()
    if choice:
        text = re.sub(r'^[①②③]\s*', '', text)
    text = PREFIX.sub('', text)
    text = EMOTE.sub('', text).strip()
    return text


def split_text(event, number, raw):
    if event == '1009' and number == 36:
        return [raw.split('（注：', 1)[0]]
    # Every specified mixed paragraph is an alternating sequence of quoted
    # speech and narration. Keep the original quotation marks and punctuation.
    parts = [s for s in re.split(r'(「[^」]*」)', raw) if s.strip()]
    result = []
    for part in parts:
        stripped = clean(part)
        if stripped:
            result.append(stripped)
    return result


def expected_rows(event):
    rows = source_rows(event)
    source = verify_docx(event, rows)
    entries = []
    for number, kind, page, key, original in rows:
        if key == '—':
            continue
        raw = clean(original, kind == '选项')
        parts = split_text(event, number, raw) if number in MIXED[event] else [raw]
        for index, value in enumerate(parts):
            suffix = '.' + chr(ord('a') + index) if number in MIXED[event] else ''
            entry_key = key + suffix
            assert value and not PREFIX.match(value) and not EMOTE.search(value)
            entries.append({'key': entry_key, 'text': value, 'source': f'B-{event}-P{number:03}', 'page': page,
                            'kind': kind, 'original': original, 'path': source})
    assert len({e['key'] for e in entries}) == len(entries)
    return entries


def signature(cell):
    return (cell.value, cell.data_type, copy(cell._style), cell.comment.text if cell.comment else None)


def stage(event):
    entries = expected_rows(event)
    initial_hash = digest(TEXT)
    wb = openpyxl.load_workbook(TEXT)
    ws = wb['Text']
    before = {(s.title, c.coordinate): signature(c) for s in wb for row in s for c in row}
    keys = {}
    for row in ws.iter_rows(min_row=5):
        key = row[1].value
        if key:
            assert key not in keys, f'Duplicate Text key: {key}'
            keys[key] = row[1].row
    prefix = f'rose.{event}.s1.text.'
    expected_keys = {e['key'] for e in entries}
    assert not {k for k in keys if str(k).startswith(prefix)} - expected_keys, 'Unexpected existing B rows; merge required'
    added = []
    for entry in entries:
        if entry['key'] in keys:
            assert ws.cell(keys[entry['key']], 3).value == entry['text'], f'Concurrent/existing text differs: {entry["key"]}'
            continue
        row = ws.max_row + 1
        for col in range(1, 5):
            ws.cell(row, col)._style = copy(ws.cell(599, col)._style)
        ws.cell(row, 2, entry['key'])
        ws.cell(row, 3, entry['text'])
        ws.cell(row, 4, f'{entry["source"]}; {entry["kind"]}; {entry["page"]}; text staging only, event disabled')
        ws.cell(row, 3).comment = Comment(f'{entry["path"]}\n{entry["source"]}\n{entry["original"]}', 'B-rose source')
        added.append(row)
    stamp = datetime.now().strftime('%Y%m%d-%H%M%S-%f')
    out = ROOT / f'Library/Locus/tmp/rose-content/{stamp}-{event}'
    out.mkdir(parents=True)
    shutil.copy2(TEXT, out / 'Text.before.xlsx')
    candidate = out / 'Text.candidate.xlsx'
    wb.save(candidate)
    saved = openpyxl.load_workbook(candidate)
    for (sheet, coordinate), value in before.items():
        assert signature(saved[sheet][coordinate]) == value, f'Existing cell changed: {sheet}!{coordinate}'
    assert digest(TEXT) == initial_hash, 'Concurrent Text.xlsx change; candidate preserved, not applied'
    if added:
        # Atomic replacement on the same project volume, after compare-and-swap check.
        candidate.replace(TEXT)
    manifest = {'time': datetime.now().astimezone().isoformat(), 'event': f'rose.{event}.s1', 'source_sha256': SOURCES[event][2],
                'source_paragraphs': SOURCES[event][1], 'text_rows': len(entries), 'added_rows': added,
                'before_sha256': initial_hash, 'after_sha256': digest(TEXT), 'preserved_existing_cells': len(before),
                'entries': entries, 'scope': 'Text only. No Event, Date, Action, Result, or completed fact created.'}
    (out / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({k: v for k, v in manifest.items() if k != 'entries'}, ensure_ascii=False))
    print('Evidence:', out)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--event', choices=SOURCES, required=True)
    from rose_write_guard import source_write_guard
    event = parser.parse_args().event
    with source_write_guard('import Text ' + event):
        stage(event)
