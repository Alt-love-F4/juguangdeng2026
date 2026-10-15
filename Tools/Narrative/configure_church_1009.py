"""Append church.1009.s2 NCS configuration; formal release remains blocked by B completion.
Uses existing choice: session conditions. Does not register dates or simulate upstream facts.
"""
from copy import copy
from datetime import datetime
import hashlib
import json
from pathlib import Path
import shutil
import openpyxl
ROOT = Path(__file__).resolve().parents[2]
P = ROOT / 'Assets/ConfigTables/Data/Narrative/Narrative.xlsx'
T = ROOT / 'Assets/ConfigTables/Data/Shared/Text.xlsx'
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def sig(c): return (c.value, c.data_type, copy(c._style), c.comment.text if c.comment else None)

def main():
    initial = {p: sha(p) for p in [P,T]}
    books = {p: openpyxl.load_workbook(p) for p in [P,T]}
    before = {p: {(s.title,c.coordinate):sig(c) for s in w for row in s for c in row} for p,w in books.items()}
    def add(p,sheet,values):
        ws=books[p][sheet]; heads={c.value:c.column for c in ws[1]}; key=list(values.values())[0]
        existing=[r for r in range(5,ws.max_row+1) if ws.cell(r,2).value==key]
        if existing:
            for k,v in values.items(): assert ws.cell(existing[0],heads[k]).value==v, (sheet,key,k)
            return
        r=ws.max_row+1
        for col in range(1,ws.max_column+1): ws.cell(r,col)._style=copy(ws.cell(r-1,col)._style)
        for k,v in values.items(): ws.cell(r,heads[k],v)
    root='church.1009.s2'
    for key,text in [('church.broadcast.name','广播'),('church.crowd.name','人群'),('church.duke.name','埃利昂公爵'),('church.man.name','男人'),('condition.story.church.1009.s2.completed','观礼完成')]:
        add(T,'Text',dict(key=key,text=text,comment='D 1009 identity/completion label; no narrative addition'))
    for obj in ['broadcast','crowd','duke','man']:
        add(P,'Object',dict(objId='church.'+obj,objType='npc',nameText='church.'+obj+'.name'))
    add(P,'Event',dict(eventId=root,eventKind='world',actorId='world',setupKey=root,onceOnly=True,logTag='D1009',presentationClass='Conversation',presentationKey=root+'.conversation',presentationLogTag='D1009',eventProperty='entry=Immediate;blocksSave=true;blockedBy=D-B03-rose-real-completion'))
    add(P,'PresentationIndex',dict(presentationKey=root+'.conversation',presentationClass='Conversation',logTag='D1009'))
    add(P,'SetupParams',dict(setupKey=root,tplName='Tpl_ArrivalConversation',slotValues='opts=',skipAble=False,autoAdvance=False))
    add(P,'ConditionKey',dict(key='story.'+root+'.completed',nameText='condition.story.'+root+'.completed',initialValue=False,onlySetTrue=True))
    add(P,'Condition',dict(conditionId=root+'.day',ownerType='event',ownerId=root,groupIndex=1,condType='BAR_BETWEEN',target='world.calendar',value='4,4'))
    add(P,'Condition',dict(conditionId=root+'.incomplete',ownerType='event',ownerId=root,groupIndex=1,condType='BOOL_EQ',target='story.'+root+'.completed',value='false'))
    # B has no established producer yet. blockedBy prevents every formal start; do not invent a B key.
    add(P,'Result',dict(resultId=root+'.complete',eventId=root,seq=1,resultType='SET_FLAG',target='story.'+root+'.completed',value='true',note='P050 root completion only; no hat fact/reward'))
    pages=[('ceremony',list(range(4,8))+[9,10,11,12,13,15,17,18,19,20,21,23,24,25,26],None),('pick',list(range(32,40)),'choice:'+root+'.p029'),('ignore',list(range(41,47)),'choice:'+root+'.p030'),('return',[48,49,50],None)]
    for seq,(short,nums,guard) in enumerate(pages,1):
        page=root+'.'+short
        add(P,'ConversationPage',dict(pageId=page,presentationKey=root+'.conversation',seq=seq,pageCondKeys=guard,optionGroupKey=root+'.hat.choices' if seq==1 else None,bgOverride='church.cathedral.day',stepType='Dialogue'))
        for n,num in enumerate(nums,1):
            speaker='church.broadcast' if num==6 else 'church.crowd' if num in [9,10,11] else 'church.duke' if num in [15,19] else 'church.man' if num==34 else None
            line=root+f'.p{num:03}'
            add(P,'ConversationLine',dict(lineId=line,pageId=page,seq=n,speakerKey=speaker,side='none',textKey=line+'.text'))
    for seq,num in enumerate([29,30],1):
        key=root+f'.p{num:03}'
        add(P,'Choice',dict(choiceId=key,groupKey=root+'.hat.choices',textKey=key+'.text',seq=seq))
    out=ROOT/('Library/Locus/tmp/church-content/'+datetime.now().strftime('%Y%m%d-%H%M%S-%f')+'-1009-config');out.mkdir(parents=True)
    candidates={}
    for p,w in books.items():
        shutil.copy2(p,out/(p.stem+'.before.xlsx'))
        cand=out/(p.stem+'.candidate.xlsx'); w.save(cand); candidates[p]=cand
        saved=openpyxl.load_workbook(cand)
        for (s,c),v in before[p].items(): assert sig(saved[s][c])==v,(p,s,c)
    for p in books: assert sha(p)==initial[p],f'Concurrent change: {p}'
    for p,cand in candidates.items(): cand.replace(p)
    manifest=dict(time=datetime.now().astimezone().isoformat(),event=root,before={str(p):h for p,h in initial.items()},after={str(p):sha(p) for p in books},preserved={str(p):len(v) for p,v in before.items()},formal_blocker='D-B03-rose-real-completion',pages=4,lines=36,choices=2)
    (out/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(manifest,ensure_ascii=False));print(out)
if __name__=='__main__': main()
