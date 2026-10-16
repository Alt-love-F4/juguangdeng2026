"""Complete three source-backed neutral entrance emotes and Antoine's existing portrait.

Read-only by default. --apply requires --expected-sha256; never generates or
imports assets. Metadata and every non-target cell are checked before CAS.
"""
from __future__ import annotations
import argparse
from datetime import datetime
import json
import shutil
from zipfile import ZipFile, ZIP_DEFLATED
import openpyxl
from openpyxl.writer.excel import ExcelWriter
from configure_neutral_remaining import MAP, NARRATIVE, ROOT, digest, require
from repair_neutral_stage import snapshot, metadata

TARGETS = (
    ('ConversationLine', 'lineId', 'neutral.1008.s1.line.p023', 'emote', '慌张', 'colette'),
    ('ConversationLine', 'lineId', 'neutral.1010.s1.line.p024', 'emote', '开心', 'adele'),
    ('ConversationLine', 'lineId', 'neutral.1010.s2.line.p011', 'emote', '平静', 'antoine'),
    ('Object', 'objId', 'antoine', 'portraitPath', 'Characters/antoine/neutral', None),
)


def repair(apply=False, expected_sha256=None):
    initial, source_sha = digest(NARRATIVE), digest(MAP)
    require(not apply or expected_sha256, '--apply requires --expected-sha256')
    require(expected_sha256 is None or expected_sha256 == initial, 'SHA changed; re-plan')
    text = MAP.read_text(encoding='utf-8')
    for cue in ('【立绘出场：柯莱特·布吕内（#慌张）】', '【立绘出场：阿黛尔·米尼耶（#开心）】',
                '【立绘出场：安托万·马尚（#平静）】'):
        require(cue in text, 'Source cue missing: ' + cue)
    portrait = ROOT / 'Assets/Resources/Characters/antoine/neutral.png'
    require(portrait.is_file(), 'Existing portrait missing')
    portrait_sha = digest(portrait)
    wb = openpyxl.load_workbook(NARRATIVE)
    try:
        before, meta = snapshot(wb), metadata(wb)
        edits = []
        for sheet, key_field, key, field, value, speaker in TARGETS:
            ws = wb[sheet]
            headers = {c.value: c.column for c in ws[1] if c.value}
            matches = [r for r in range(5, ws.max_row + 1) if ws.cell(r, headers[key_field]).value == key]
            require(len(matches) == 1, 'Missing/duplicate: ' + key)
            row = matches[0]
            if speaker:
                require(ws.cell(row, headers['speakerKey']).value == speaker, 'Unexpected speaker: ' + key)
                require(ws.cell(row, headers['textKey']).value == key.replace('.line.', '.text.'), 'Unexpected text key')
                require(key in text, 'Source body missing')
            cell = ws.cell(row, headers[field])
            require(cell.value in (None, value), 'Authored value differs: ' + key)
            if cell.value != value:
                edits.append(dict(sheet=sheet, cell=cell.coordinate, key=key, field=field, before=cell.value, after=value))
                cell.value = value
        report = dict(tool='repair_neutral_visual_bindings', before_sha256=initial,
                      after_sha256=initial, source_map_sha256=source_sha, portrait_sha256=portrait_sha,
                      changed=False, idempotent=not edits, edits=edits)
        def cas():
            require(digest(NARRATIVE) == initial and digest(MAP) == source_sha and digest(portrait) == portrait_sha,
                    'Concurrent source change; CAS aborted')
        cas()
        if not apply or not edits:
            return report
        out = ROOT / 'Library/Locus/tmp/neutral-content' / (datetime.now().strftime('%Y%m%d-%H%M%S-%f') + '-visual-bindings')
        out.mkdir(parents=True, exist_ok=False)
        backup, candidate = out / 'Narrative.before.xlsx', out / 'Narrative.candidate.xlsx'
        shutil.copy2(NARRATIVE, backup)
        require(digest(backup) == initial, 'Backup differs')
        with ZipFile(candidate, 'w', ZIP_DEFLATED, allowZip64=True) as archive:
            ExcelWriter(wb, archive).save()
        saved = openpyxl.load_workbook(candidate)
        try:
            actual, desired = snapshot(saved), snapshot(wb)
            require(actual.keys() == before.keys(), 'Cell coordinates changed')
            require(metadata(saved) == meta, 'Workbook/sheet metadata changed')
            allowed = {(e['sheet'], e['cell']) for e in edits}
            for key, old in before.items():
                require(actual[key] == (desired[key] if key in allowed else old), 'Cell readback differs: ' + str(key))
                if key in allowed:
                    require(actual[key][2:] == old[2:], 'Target metadata changed')
        finally:
            saved.close()
        replacement = out / 'Narrative.replace.xlsx'
        shutil.copy2(candidate, replacement)
        after_sha = digest(candidate)
        require(digest(replacement) == after_sha, 'Replacement differs')
        cas()
        replacement.replace(NARRATIVE)
        require(digest(NARRATIVE) == after_sha, 'Final SHA differs')
        report.update(changed=True, after_sha256=after_sha, evidence=str(out),
                      cells_checked=len(before), preserved_cells=len(before)-len(edits), metadata_preserved=True)
        (out / 'repair-evidence.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        return report
    finally:
        wb.close()


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--apply', action='store_true')
    parser.add_argument('--expected-sha256')
    args = parser.parse_args()
    print(json.dumps(repair(args.apply, args.expected_sha256), ensure_ascii=False, indent=2))
