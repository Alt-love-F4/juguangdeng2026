"""Repair exactly five source-backed neutral exits; no generation or Unity calls.

Default: read-only plan. --apply requires the SHA from that plan. Only the five
stage values may change; existing metadata and all other cell signatures must
survive an openpyxl round trip. A second apply is a byte-preserving no-op.
"""
from __future__ import annotations

import argparse
from copy import deepcopy
from datetime import datetime
import hashlib
import json
from pathlib import Path
import shutil
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
from openpyxl.writer.excel import ExcelWriter
from openpyxl.xml.functions import tostring

from configure_neutral_remaining import MAP, NARRATIVE, ROOT, digest, require, signature

# lineId, field, page suffix, batch, source-map paragraph evidence.
# A01 predates the batch-comment convention; never add metadata to that row.
TARGETS = (
    ("neutral.1008.s1.line.p083", "stageRightKey", "outro", "A01", (80, 83)),
    ("neutral.1010.s1.line.p074", "stageRightKey", "join", "A02", (73, 74)),
    ("neutral.1010.s2.line.p107", "stageRightKey", "refuse_tail", "A03", (38, 107)),
    ("neutral.1010.s2.line.p112", "stageRightKey", "common_end", "A03", (103, 112)),
    ("neutral.1013.s2.line.p135", "stageLeftKey", "visit_tail", "A04", (133, 135)),
)


def snapshot(wb):
    return {(ws.title, cell.coordinate): signature(cell)
            for ws in wb for row in ws for cell in row}


def signature_sha256(cells):
    payload = [(sheet, coord, value) for (sheet, coord), value in sorted(cells.items())]
    return hashlib.sha256(json.dumps(payload, ensure_ascii=False, default=str,
                                     separators=(",", ":")).encode("utf-8")).hexdigest()


def metadata(wb):
    """Workbook and sheet metadata in addition to signature's cell metadata."""
    def dimensions(items):
        # Dimension equality can depend on in-memory parent/style objects.
        return [(key, tuple(sorted(dict(value).items())),
                 tuple(value._style) if value.has_style else None)
                for key, value in items.items()]

    return deepcopy((wb.properties, tostring(wb.custom_doc_props.to_tree()), wb.defined_names,
                     wb.calculation, wb.security, wb.views,
                     [(ws.title, ws.sheet_state, ws.sheet_properties,
                       ws.sheet_format, ws.views, ws.freeze_panes,
                       str(ws.merged_cells), dimensions(ws.row_dimensions),
                       dimensions(ws.column_dimensions), ws.auto_filter,
                       ws.data_validations, ws.print_options,
                       ws.page_margins, ws.page_setup, ws.protection)
                      for ws in wb]))


def source_evidence():
    lines = MAP.read_text(encoding="utf-8").splitlines()
    result = {}
    for key, _, _, batch, paragraphs in TARGETS:
        root = key.split(".line.")[0]
        active = False
        found = []
        for number, line in enumerate(lines, 1):
            if line.startswith("## A"):
                active = root in line and line.startswith("## " + batch + "：")
            if active and any(line.startswith(f"| P{p:03} |") for p in paragraphs):
                found.append({"line": number, "text": line})
        require(len(found) == len(paragraphs), f"Missing source evidence: {key}")
        require(any(key in item["text"] for item in found), f"Missing source line: {key}")
        result[key] = found
    return result


def repair(apply=False, expected_sha256=None):
    initial = digest(NARRATIVE)
    if expected_sha256 is not None:
        require(initial == expected_sha256.lower(), "Narrative SHA differs from expected; re-plan")
    require(not apply or expected_sha256 is not None, "--apply requires --expected-sha256")
    source_sha = digest(MAP)
    evidence = source_evidence()
    wb = openpyxl.load_workbook(NARRATIVE)
    try:
        before = snapshot(wb)
        old_metadata = metadata(wb)
        dimensions = {ws.title: (ws.max_row, ws.max_column) for ws in wb}
        ws = wb["ConversationLine"]
        headers = {cell.value: cell.column for cell in ws[1] if cell.value}
        require({"lineId", "pageId", "textKey", "stageLeftKey", "stageRightKey"} <= headers.keys(),
                "ConversationLine schema mismatch")
        edits, targets = [], []
        for key, field, page, batch, _ in TARGETS:
            matches = [r for r in range(5, ws.max_row + 1)
                       if ws.cell(r, headers["lineId"]).value == key]
            require(len(matches) == 1, f"Missing or duplicate line: {key}")
            row = matches[0]
            root = key.split(".line.")[0]
            require(ws.cell(row, headers["pageId"]).value == root + "." + page,
                    f"Unexpected page: {key}")
            require(ws.cell(row, headers["textKey"]).value == key.replace(".line.", ".text."),
                    f"Unexpected Text key: {key}")
            comment = ws.cell(row, headers["lineId"]).comment
            if batch != "A01" or comment is not None:
                require(comment is not None, f"Missing ownership metadata: {key}")
                require(json.loads(comment.text).get("batch") == batch, f"Wrong row owner: {key}")
            cell = ws.cell(row, headers[field])
            require(cell.value in (None, "-"), f"Unexpected authored stage: {key}/{field}")
            item = {"sheet": ws.title, "cell": cell.coordinate, "lineId": key,
                    "field": field, "before": cell.value, "after": "-"}
            targets.append(item)
            if cell.value is None:
                edits.append(item)
                cell.value = "-"
        require(len(edits) in (0, 5), "Partial repair state; expected all five blank or all five repaired")
        allowed = {(item["sheet"], item["cell"]) for item in edits}
        unchanged = {key: value for key, value in before.items() if key not in allowed}
        report = {"tool": "repair_neutral_stage", "changed": False,
                  "idempotent": not edits, "writes": bool(apply and edits),
                  "before_sha256": initial, "after_sha256": initial,
                  "source_map_sha256": source_sha, "targets": targets, "edits": edits,
                  "source_evidence": evidence, "cells_checked": len(before),
                  "preserved_cells": len(unchanged),
                  "preserved_cells_signature_sha256": signature_sha256(unchanged)}

        def cas_check():
            require(digest(NARRATIVE) == initial, "Concurrent Narrative change; CAS aborted")
            require(digest(MAP) == source_sha, "Concurrent source-map change; CAS aborted")

        cas_check()
        if not apply or not edits:
            return report
        out = ROOT / "Library/Locus/tmp/neutral-content" / (
            datetime.now().strftime("%Y%m%d-%H%M%S-%f") + "-stage-repair")
        out.mkdir(parents=True, exist_ok=False)
        backup = out / "Narrative.before.xlsx"
        candidate = out / "Narrative.candidate.xlsx"
        shutil.copy2(NARRATIVE, backup)
        require(digest(backup) == initial, "Backup SHA mismatch")
        # Workbook.save() overwrites properties.modified in openpyxl 3.1.5.
        # Use the same openpyxl writer while retaining the authored core metadata.
        with ZipFile(candidate, "w", ZIP_DEFLATED, allowZip64=True) as archive:
            ExcelWriter(wb, archive).save()
        saved = openpyxl.load_workbook(candidate)
        try:
            require(saved.sheetnames == wb.sheetnames, "Sheet names/order changed")
            require({s.title: (s.max_row, s.max_column) for s in saved} == dimensions,
                    "Sheet dimensions changed")
            actual = snapshot(saved)
            require(actual.keys() == before.keys(), "Cell coordinate set changed")
            require(metadata(saved) == old_metadata, "Workbook/sheet metadata changed")
            for key, old in before.items():
                if key in allowed:
                    require(actual[key] == signature(wb[key[0]][key[1]]), f"Readback differs: {key}")
                    require(actual[key][2:] == old[2:], f"Target metadata changed: {key}")
                else:
                    require(actual[key] == old, f"Non-target signature changed: {key}")
            actual_unchanged = {key: value for key, value in actual.items() if key not in allowed}
            require(signature_sha256(actual_unchanged) == report["preserved_cells_signature_sha256"],
                    "Preserved-cell signature SHA mismatch")
        finally:
            saved.close()
        candidate_sha = digest(candidate)
        report.update({"evidence": str(out), "backup_sha256": digest(backup),
                       "candidate_sha256": candidate_sha, "after_sha256": candidate_sha,
                       "changed_cells": len(allowed), "metadata_preserved": True,
                       "status": "verified-candidate"})
        report_path = out / "repair-evidence.json"
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        replacement = out / "Narrative.replace.xlsx"
        shutil.copy2(candidate, replacement)
        require(digest(replacement) == candidate_sha, "Replacement SHA mismatch")
        require(digest(candidate) == candidate_sha and digest(backup) == initial,
                "Evidence workbook SHA changed")
        cas_check()
        replacement.replace(NARRATIVE)
        require(digest(NARRATIVE) == candidate_sha, "Post-CAS SHA mismatch")
        require(digest(MAP) == source_sha, "Source-map changed during CAS")
        report.update({"changed": True, "status": "committed"})
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        return report
    finally:
        wb.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--expected-sha256", help="Narrative SHA from the read-only plan")
    args = parser.parse_args()
    print(json.dumps(repair(args.apply, args.expected_sha256), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
