"""Stage one D-church event's verified source Text without enabling gameplay.
Run --event 1009.s2|1014.s2|1020.s1|1023.s1|1023.s2, then official GenerateAll.bat.
No runtime reads this script or its manifests; Text.xlsx remains the only source.
"""
from __future__ import annotations
import argparse
from copy import copy
from datetime import datetime
import hashlib
import json
from pathlib import Path
import re
import shutil
import xml.etree.ElementTree as ET
import zipfile
import openpyxl
from openpyxl.comments import Comment

ROOT = Path(__file__).resolve().parents[2]
MAP = ROOT / 'Docs/superpowers/specs/2026-10-16-content-batches/D-church/source-map.md'
TEXT = ROOT / 'Assets/ConfigTables/Data/Shared/Text.xlsx'
SOURCES = {
    '1009.s2': (50, 38, '15c3101b85329f082efa98d9c938791089d9ac5e5cb35c50d8788fa4a3d9c58d'),
    '1014.s2': (75, 57, 'daeb634d81a175821892fb71a1a13159c4ec7c77467ca96a419d91dc03f03ad6'),
    '1020.s1': (153, 124, '250ba8e39061a0ed436aaab452500067918810ad3ec78985a24687c2dc840f6e'),
    '1023.s1': (153, 116, '34cec3f45c025050c0013758dd37f70f9f3cab70457ac5d7856e83c74c3ca29e'),
    '1023.s2': (34, 25, '632b3c47f06bdaae5440c0d45ba89ba9acaa430aba6eb5778bd6230c04fe63a1'),
}
PREFIX = re.compile(r'^(?:\{playername\}|埃利昂公爵|男人|客人|阿黛尔·米尼耶|阿黛尔|疑似领头的人|监察官|众监察员|看守)\s*[：:；;]\s*')
EMOTE = re.compile(r'[（(]\s*#.*?[）)]')


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def expected_rows(event):
    count, expected_count, sha = SOURCES[event]
    root = 'church.' + event
    section = MAP.read_text(encoding='utf-8').split('## ' + root + '\n', 1)[1].split('\n## ', 1)[0]
    mapped = []
    page = ''
    for line in section.splitlines():
        if line.startswith('### page '):
            page = line.split('`')[1]
        if not re.match(r'^\|P\d{3}\|', line):
            continue
        cells = line.split('|')[1:-1]
        assert len(cells) == 4
        number = int(cells[0][1:])
        mapped.append((number, cells[1], cells[2], cells[3], page))
    assert [x[0] for x in mapped] == list(range(1, count + 1))
    date, suffix = event.split('.')
    source = Path('D:/桌面/文案/王庭教会') / f'MA70.{date[:2]}.{date[2:]}-{suffix.upper()}.docx'
    assert digest(source) == sha, f'Source changed: {source}'
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    with zipfile.ZipFile(source) as z:
        xml = ET.fromstring(z.read('word/document.xml'))
    originals = [''.join(n.text or '' for n in p.findall('.//w:t', ns)) for p in xml.findall('.//w:body//w:p', ns)]
    originals = [p for p in originals if p.strip()]
    assert len(originals) == count
    assert [p.rstrip() for p in originals] == [r[3].rstrip() for r in mapped], 'DOCX/source-map mismatch'
    entries = []
    for number, kind, mapping, original, page in mapped:
        if not kind.startswith(('正文', '选项')):
            continue
        key = re.search(r'Text=`([^`]+)`', mapping).group(1)
        text = original.strip()
        for typo in ['{playname}', '{playernanem}', '{playernmae}']:
            text = text.replace(typo, '{playername}')
        text = re.sub(r'^[①②③④]\s*', '', text)
        text = PREFIX.sub('', text)
        text = EMOTE.sub('', text)
        if event == '1020.s1' and number == 65:
            assert text.startswith('）#平静')
            text = text[len('）#平静'):]
        if kind.startswith('选项'):
            text = re.sub(r'（(?:语言|勇气|智慧|知觉)检定）|（资金条件）|\[王庭-教会倾向[+-]\d+\]', '', text)
        if event == '1009.s2' and number in [9, 10, 11]:
            text = re.sub(r'（[近中远]景）$', '', text)
        if event == '1023.s1' and number == 45:
            text = text.replace('（数值待定）', '')
        parts = [text.strip()]
        if event == '1014.s2' and number == 31:
            assert parts[0] == '「不，不，没有任何意见。」放下咖啡「店长，我有一个问题想问你。」'
            parts = ['「不，不，没有任何意见。」', '「店长，我有一个问题想问你。」']
        for index, value in enumerate(parts):
            entry_key = key if len(parts) == 1 else key.removesuffix('.text') + '.' + chr(97 + index) + '.text'
            assert value and not PREFIX.match(value) and '#' not in value
            entries.append(dict(key=entry_key, text=value, source=f'{root}.P{number:03}', page=page,
                                kind=kind, original=original, path=str(source)))
    assert len(entries) == expected_count
    assert len({e['key'] for e in entries}) == len(entries)
    return entries


def signature(cell):
    return (cell.value, cell.data_type, copy(cell._style), cell.comment.text if cell.comment else None)


def stage(event):
    entries = expected_rows(event)
    initial_hash = digest(TEXT)
    wb = openpyxl.load_workbook(TEXT)
    ws = wb['Text']
    assert [ws.cell(1, c).value for c in range(1, 5)] == ['##var', 'key', 'text', 'comment']
    assert ws.cell(2, 3).value == 'string' and ws.cell(3, 3).value == 'c'
    before = {(s.title, c.coordinate): signature(c) for s in wb for row in s for c in row}
    keys = {}
    for row in ws.iter_rows(min_row=5):
        if row[1].value:
            assert row[1].value not in keys
            keys[row[1].value] = row[1].row
    prefix = 'church.' + event + '.'
    assert not {k for k in keys if str(k).startswith(prefix)} - {e['key'] for e in entries}
    added = []
    template_row = ws.max_row
    for entry in entries:
        if entry['key'] in keys:
            assert ws.cell(keys[entry['key']], 3).value == entry['text']
            continue
        row = ws.max_row + 1
        for col in range(1, 5):
            ws.cell(row, col)._style = copy(ws.cell(template_row, col)._style)
        ws.cell(row, 2, entry['key'])
        ws.cell(row, 3, entry['text'])
        ws.cell(row, 4, f'{entry["source"]}; {entry["page"]}; Text only, no event/route/result')
        ws.cell(row, 3).comment = Comment(f'{entry["path"]}\n{entry["source"]}\n{entry["original"]}', 'D-church source')
        added.append(row)
    out = ROOT / ('Library/Locus/tmp/church-content/' + datetime.now().strftime('%Y%m%d-%H%M%S-%f') + '-' + event)
    out.mkdir(parents=True)
    shutil.copy2(TEXT, out / 'Text.before.xlsx')
    candidate = out / 'Text.candidate.xlsx'
    wb.save(candidate)
    saved = openpyxl.load_workbook(candidate)
    for (sheet, coordinate), value in before.items():
        assert signature(saved[sheet][coordinate]) == value, f'Existing cell changed: {sheet}!{coordinate}'
    assert digest(TEXT) == initial_hash, 'Concurrent change: candidate not applied'
    if added:
        candidate.replace(TEXT)
    manifest = dict(time=datetime.now().astimezone().isoformat(), event='church.' + event,
                    source_sha256=SOURCES[event][2], source_paragraphs=SOURCES[event][0], text_rows=len(entries),
                    added_rows=added, before_sha256=initial_hash, after_sha256=digest(TEXT),
                    preserved_existing_cells=len(before), entries=entries,
                    scope='Text only. No Event/Date/Action/ConditionKey/Result or formal completion.')
    (out / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({k: v for k, v in manifest.items() if k != 'entries'}, ensure_ascii=False))
    print('Evidence:', out)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--event', choices=SOURCES, required=True)
    stage(parser.parse_args().event)
