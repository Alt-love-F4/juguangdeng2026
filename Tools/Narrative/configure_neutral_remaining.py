"""Prepare blocked, result-free A02-A05 isolated narrative previews.

This file is intentionally inert until invoked with exactly one ``--event``.
It never creates Text rows: all 311 A02-A05 body/choice keys must already exist
in the authoritative Text workbook, staged by import_neutral_text.py.

The optional --apply path appends/repairs explicitly owned Narrative fields
under SHA-256 CAS, preserving metadata and every other cell. Text is read-only.
It creates no Result, EventDate, Action, or resource data.
It does not run generation or Unity.
"""
from __future__ import annotations

import argparse
from collections import Counter
from copy import copy
from datetime import datetime
import hashlib
import json
from pathlib import Path
import re
import shutil

import openpyxl
from openpyxl.comments import Comment

ROOT = Path(__file__).resolve().parents[2]
MAP = ROOT / "Docs/superpowers/specs/2026-10-16-content-batches/A-neutral/source-map.md"
NARRATIVE = ROOT / "Assets/ConfigTables/Data/Narrative/Narrative.xlsx"
TEXT = ROOT / "Assets/ConfigTables/Data/Shared/Text.xlsx"

EVENTS = {
    "1010.s1": ("neutral.1010.s1", "A02", "5", "A-B002,A-B003,A-B004,A-B005,A-B010"),
    "1010.s2": ("neutral.1010.s2", "A03", "5", "A-B003,A-B004,A-B010"),
    "1013.s2": ("neutral.1013.s2", "A04", "8", "A-B002,A-B004,A-B005,A-B006,A-B010,A-CAP01,A-CAP04"),
    "1018.s1": ("neutral.1018.s1", "A05", "13", "A-B002,A-B007,A-B008,A-B009,A-CAP02,A-CAP05"),
}

# These are the only new display names permitted by the A-neutral source.
NAMES = {
    # Existing Text keys are reused; only the reporter display key is new.
    "adele": ("character.adele.name", "阿黛尔·米尼耶"),
    "antoine": ("character.antoine.name", "安托万·马尚"),
    "colette": ("colette.name", "柯莱特·布吕内"),
    "unknown": ("neutral.unknown.name", "？？？"),
    "reporter": ("neutral.reporter.name", "记者"),
    "shopkeeper": ("neutral.shopkeeper.name", "店主"),
}

# Source-map display rows are deliberately parsed from the authoritative map.
ROW_RE = re.compile(r"^\| P(\d{3}) \| ([^|]+) \| ([^|]+) \| (.*) \|$")
PAGE_RE = re.compile(r"^### A\d+ / ([A-Za-z0-9_]+)")

HEADERS = {
    "Event": "eventId eventKind actorId targetId titleText contentKey setupType setupKey onceOnly logTag presentationClass presentationKey presentationLogTag eventProperty",
    "PresentationIndex": "presentationKey presentationClass logTag",
    "SetupParams": "setupKey tplName slotValues skipAble autoAdvance checkId",
    "ConversationPage": "pageId presentationKey seq pageCondKeys pageBarConds optionGroupKey bgOverride bgmOverride stepType readingKey checkId",
    "ConversationLine": "lineId pageId seq speakerKey side textKey voiceKey emote sfxKey accentFxKey accentTargetKey speakerNameKey",
    "Choice": "choiceId groupKey textKey resultOwnerId seq checkId",
    "AbilityCheck": "checkId successEventId failureEventId successResultOwnerId failureResultOwnerId comment mode",
    "AbilityCheckCondition": "conditionId checkId attributeId threshold groupIndex seq",
    "Object": "objId objType nameText iconPath portraitPath portraitIndoorPath avtarPath descText tags",
    "Text": "key text comment",
}


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def signature(cell):
    # openpyxl may normalize an empty inline string to an empty numeric cell on save;
    # compare authored content and metadata while treating those two empty forms alike.
    value = cell.value
    if value is None and cell.data_type in ("inlineStr", "n"):
        data_type = "empty"
    else:
        data_type = cell.data_type
    style = tuple(cell._style) if cell.has_style else None
    hyperlink = tuple(sorted(dict(cell.hyperlink).items())) if cell.hyperlink else None
    comment = ((cell.comment.text, cell.comment.author, cell.comment.width, cell.comment.height)
               if cell.comment else None)
    return (value, data_type, style, comment, hyperlink)


def require(ok: bool, message: str):
    if not ok:
        raise RuntimeError(message)


def source_rows(root: str):
    """Return source pages and display rows for exactly one root."""
    active = False
    page = None
    pages = []
    rows = []
    for line_number, line in enumerate(MAP.read_text(encoding="utf-8").splitlines(), 1):
        if line.startswith("## A") and "：" in line:
            active = root in line
            if not active and pages:
                break
        if not active:
            continue
        page_match = PAGE_RE.match(line)
        if page_match:
            page = page_match.group(1)
            pages.append({"page": page, "map_line": line_number})
            continue
        match = ROW_RE.match(line)
        if not match:
            continue
        number, kind, location, original = match.groups()
        number = int(number)
        cells = [x.strip() for x in location.split("/")]
        key = cells[-1].strip("`")
        if (kind.strip() in ("正文", "选项") or kind.strip().startswith("选项标签")) and ".text." in key:
            rows.append({"number": number, "kind": kind.strip(), "key": key,
                         "original": original.replace("<br>", "\n"), "page": page,
                         "map_line": line_number})
    require(pages, f"No source-map pages found for {root}")
    require(rows, f"No display rows found for {root}")
    require(len({r['key'] for r in rows}) == len(rows), f"Duplicate source Text key in {root}")
    return pages, rows


def speaker_for(raw: str):
    # Require the complete name, optional expression, and colon: A05 P089 is narration.
    match = re.match(r"^([^：:（(]+)(?:[（(]#[^）)]+[）)])?[：:]", raw)
    if not match:
        return None, None
    speakers = {
        "{playername}": ("player", None),
        "阿黛尔·米尼耶": ("adele", None),
        "安托万·马尚": ("antoine", None),
        "柯莱特·布吕内": ("colette", None),
        "店主": ("neutral.shopkeeper", None),
        "记者": ("neutral.reporter", "neutral.reporter.name"),
        "？？？": ("neutral.unknown", "neutral.unknown.name"),
    }
    for suffix in "abcd":
        key = "character.guest_" + suffix
        speakers["顾客" + suffix.upper()] = (key, key + ".name")
    return speakers.get(match.group(1), (None, None))


def page_guard(root: str, short: str):
    """Gate choices before checks so skipped branches never request snapshots."""
    guards = []
    routes = {
        "neutral.1010.s1": {
            "agree": ("q1.a", None),
            "cost_pass": ("q1.b", "wisdom"), "cost_fail": ("q1.b", "wisdom"),
            "praise_pass": ("q1.c", "language"), "praise_fail": ("q1.c", "language"),
        },
        "neutral.1010.s2": {
            "accept": ("q1.a", None),
            "invite_pass": ("q1.b", "q1_language"), "invite_fail": ("q1.b", "q1_language"),
            "refuse": ("q1.c", None), "refuse_tail": ("q1.c", None),
            "plan_pass": ("q2.a", "q2_language"), "plan_fail": ("q2.a", "q2_language"),
            "price": ("q2.b", None), "honest": ("q2.c", None),
        },
        "neutral.1013.s2": {
            "cooperate": ("q1.a", None),
            "organize_pass": ("q1.b", "q1_wisdom"), "organize_fail": ("q1.b", "q1_wisdom"),
            "promote_pass": ("q1.c", "q1_language"), "promote_fail": ("q1.c", "q1_language"),
            "longterm_pass": ("q2.a", "q2_language"), "longterm_fail": ("q2.a", "q2_language"),
            "accounts_pass": ("q2.b", None), "accounts_fail": ("q2.b", None),
            "present": ("q2.c", None),
        },
        "neutral.1018.s1": {
            "advert_yes": ("q_ad.yes", None), "advert_no": ("q_ad.no", None),
        },
    }
    if root == "neutral.1010.s2" and short in {
        "serve", "q2", "plan_pass", "plan_fail", "price", "honest", "serve_tail"
    }:
        guards.append("anyChoice:" + ";".join(root + ".choice.q1." + x for x in ("a", "b")))
    if root == "neutral.1013.s2" and short in {
        "visit", "longterm_pass", "longterm_fail", "accounts_pass", "accounts_fail", "present", "visit_tail"
    }:
        guards.append("dependency:A04.antoine_approval")
    route = routes.get(root, {}).get(short)
    if route:
        choice, check = route
        guards.append("choice:" + root + ".choice." + choice)
        if check:
            guards.append(("not:" if short.endswith("fail") else "") + "check:" + root + "." + check)
    if root == "neutral.1013.s2" and short in {"accounts_pass", "accounts_fail"}:
        guards.append(("not:" if short.endswith("fail") else "") + "dependency:A04.q2B.wisdom-pass")
    if root == "neutral.1018.s1":
        if short in {"taste_pass", "recommend"}:
            guards.append("dependency:A05.quality.pass")
        if short in {"taste_fail", "advert_offer", "advert_yes", "advert_no"}:
            guards.append("dependency:A05.quality.fail")
        if short == "advert_yes":
            guards.append("dependency:A05.advertisement.fee-and-balance")
    return "|".join(guards)


def build(event: str):
    root, batch, day, blockers = EVENTS[event]
    pages, rows = source_rows(root)
    text_wb = openpyxl.load_workbook(TEXT, read_only=True)
    text_keys = {row[1].value for row in text_wb["Text"].iter_rows(min_row=5) if row[1].value}
    text_wb.close()
    missing = sorted({r["key"] for r in rows} - text_keys)
    require(not missing, f"Missing pre-staged Text keys ({len(missing)}): {missing[:5]}")

    result = []
    def add(sheet, **values):
        result.append({"sheet": sheet, "values": values})

    presentation = root + ".conversation"
    add("Event", eventId=root, eventKind="world", actorId="world", setupKey=root,
        onceOnly=True, logTag=batch, presentationClass="Conversation", presentationKey=presentation,
        presentationLogTag=batch,
        eventProperty=f"entry={'MorningBeforeMap' if event == '1010.s2' else 'MapAction' if event != '1018.s1' else 'Immediate'};blocksSave=true;blockedBy={blockers}")
    add("PresentationIndex", presentationKey=presentation, presentationClass="Conversation", logTag=batch)
    add("SetupParams", setupKey=root, tplName="Tpl_ArrivalConversation", skipAble=False, autoAdvance=False)

    # Names are only source names; no invented identities or numeric metadata.
    needed_speakers = {speaker_for(r["original"])[0] for r in rows} - {None, "player"}
    for speaker in sorted(needed_speakers):
        if speaker.startswith("character.guest_"):
            # Existing shared names are reused; the narrative schema still needs Object rows.
            add("Object", objId=speaker, objType="npc", nameText=speaker + ".name")
        elif speaker == "neutral.unknown":
            add("Object", objId=speaker, objType="npc", nameText="neutral.unknown.name")
        elif speaker == "neutral.reporter":
            add("Text", key="neutral.reporter.name", text="记者", comment=f"{batch} source identity/display")
            add("Object", objId=speaker, objType="npc", nameText="neutral.reporter.name")
        else:
            name_key, name = next((pair for key, pair in NAMES.items() if (key == speaker or key == speaker.removeprefix("neutral."))), None), None
            # Existing named objects may be reused; only add missing source-name rows.
            canonical = speaker.removeprefix("neutral.")
            if canonical in NAMES:
                name_key, name = NAMES[canonical]
                add("Text", key=name_key, text=name, comment=f"{batch} source identity/display")
                add("Object", objId=canonical, objType="npc", nameText=name_key)

    checks = {
        "1010.s1": [("wisdom", "wisdom", 3), ("language", "language", 2)],
        "1010.s2": [("q1_language", "language", 2), ("q2_language", "language", 3)],
        "1013.s2": [("q1_wisdom", "wisdom", 2), ("q1_language", "language", 3), ("q2_language", "language", 3)],
    }.get(event, [])
    for suffix, attr, threshold in checks:
        check = root + "." + suffix
        add("AbilityCheck", checkId=check, mode="Inline", comment=f"{batch} source cached Inline snapshot")
        add("AbilityCheckCondition", conditionId=check + ".1", checkId=check,
            attributeId=attr, threshold=threshold, groupIndex=1, seq=1)

    page_seq = 0
    for page_info in pages:
        short = page_info["page"]
        page_rows = [r for r in rows if r["page"] == short and r["kind"] == "正文"]
        if not page_rows:
            continue
        page_seq += 1
        option_group = None
        group = {
            "1010.s1": {"open": "q1"},
            "1010.s2": {"arrival": "q1", "serve": "q2"},
            "1013.s2": {"morning": "q1", "visit": "q2"},
            "1018.s1": {"advert_offer": "q_ad"},
        }[event].get(short)
        if group:
            option_group = root + "." + group
        add("ConversationPage", pageId=root + "." + short, presentationKey=presentation, seq=page_seq,
            pageCondKeys=page_guard(root, short), optionGroupKey=option_group,
            bgOverride="oct07.cafe.day" if event != "1013.s2" or short != "closing" else "neutral.cafe.dusk")
        for seq, row in enumerate(page_rows, 1):
            speaker, name_key = speaker_for(row["original"])
            if name_key is None and speaker == "adele":
                name_key = "character.adele.name"
            elif name_key is None and speaker == "antoine":
                name_key = "character.antoine.name"
            elif name_key is None and speaker == "colette":
                name_key = "colette.name"
            elif name_key is None and speaker == "neutral.shopkeeper":
                name_key = "neutral.shopkeeper.name"
            elif name_key is None and speaker == "player":
                name_key = "player.name"
            emote = re.search(r"[（(]#([^）)]+)[）)]", row["original"])
            values = dict(lineId=row["key"].replace(".text.", ".line."), pageId=root + "." + short,
                          seq=seq, speakerKey=speaker,
                          side="left" if event == "1013.s2" and speaker == "antoine" else "right" if speaker in {"adele", "antoine", "colette"} else "none",
                          textKey=row["key"], emote=emote.group(1) if emote else None,
                          speakerNameKey=name_key)
            add("ConversationLine", **values)

    choice_numbers = {
        "1010.s1": [("q1", "a", 34), ("q1", "b", 35), ("q1", "c", 36)],
        "1010.s2": [("q1", "a", 16), ("q1", "b", 17), ("q1", "c", 18), ("q2", "a", 59), ("q2", "b", 60), ("q2", "c", 61)],
        "1013.s2": [("q1", "a", 37), ("q1", "b", 38), ("q1", "c", 39), ("q2", "a", 100), ("q2", "b", 101), ("q2", "c", 102)],
        "1018.s1": [("q_ad", "yes", 69), ("q_ad", "no", 78)],
    }.get(event, [])
    row_by_number = {r["number"]: r for r in rows}
    for group, label, number in choice_numbers:
        require(number in row_by_number, f"Choice source row missing: {root} P{number:03}")
        add("Choice", choiceId=f"{root}.choice.{group}.{label}", groupKey=f"{root}.{group}",
            textKey=row_by_number[number]["key"], seq=1 + [x[1] for x in choice_numbers if x[0] == group].index(label))
    return result, {"root": root, "batch": batch, "pages": pages, "rows": rows, "source_map_sha256": digest(MAP)}


# Existing rows may change only these values. Stage columns and all metadata
# remain outside the repair set. Obsolete option-lines are the sole exception:
# clear every configuration column, including the ID, retaining the physical row.
REPAIR_FIELDS = {
    "ConversationPage": {"pageCondKeys", "optionGroupKey"},
    "ConversationLine": {"seq", "speakerKey", "side", "speakerNameKey"},
}
EXPECTED_COUNTS = {"1010.s1": (52, 3), "1010.s2": (67, 6),
                   "1013.s2": (100, 6), "1018.s1": (75, 2)}


def apply(event: str, rows, manifest):
    root, batch = manifest["root"], manifest["batch"]
    require((root, batch) == EVENTS[event][:2], "Manifest/event mismatch")
    counts = Counter(item["sheet"] for item in rows)
    require((counts["ConversationLine"], counts["Choice"]) == EXPECTED_COUNTS[event],
            f"Unexpected body/choice counts: {counts}")
    hashes = {p: digest(p) for p in (NARRATIVE, TEXT, MAP)}
    books = {p: openpyxl.load_workbook(p) for p in (NARRATIVE, TEXT)}
    wb = books[NARRATIVE]
    before = {(s.title, c.coordinate): signature(c) for s in wb for row in s for c in row}
    owned_cells = set()
    for ws in wb:
        for row in ws.iter_rows(min_row=5):
            key, comment = row[1].value, row[1].comment
            if isinstance(key, str) and (key == root or key.startswith(root + ".")) and comment:
                try:
                    if json.loads(comment.text).get("batch") == batch:
                        owned_cells.update((ws.title, c.coordinate) for c in row)
                except (ValueError, AttributeError):
                    pass
    dimensions = {s.title: (s.max_row, s.max_column) for s in wb}
    edits, locations, cleared = [], [], []
    added = Counter()

    def owned(ws, number, key):
        require(isinstance(key, str) and (key == root or key.startswith(root + ".")),
                f"Outside owned root: {ws.title}/{key}")
        comment = ws.cell(number, 2).comment
        require(comment is not None, f"Missing ownership metadata: {ws.title}/{key}")
        try:
            owner = json.loads(comment.text).get("batch")
        except (ValueError, AttributeError):
            owner = None
        require(owner == batch, f"Wrong row owner: {ws.title}/{key}: {owner}")

    def change(cell, value, key, field):
        # Store cleared fields as None, matching XLSX round-trip semantics.
        if value == "":
            value = None
        if cell.value == value:
            return
        edits.append({"sheet": cell.parent.title, "cell": cell.coordinate,
                      "key": key, "field": field, "before": cell.value, "after": value})
        cell.value = value

    try:
        require(manifest["source_map_sha256"] == hashes[MAP], "Source map changed after build")
        for item in rows:
            sheet, values = item["sheet"], item["values"]
            ws = books[TEXT if sheet == "Text" else NARRATIVE][sheet]
            headers = {c.value: c.column for c in ws[1]}
            require(sheet in HEADERS and set(values) <= set(HEADERS[sheet].split()),
                    f"Undeclared fields: {sheet}")
            require(set(values) <= headers.keys(), f"Missing schema fields: {sheet}")
            key_field = HEADERS[sheet].split()[0]
            key = values[key_field]
            matches = [r for r in range(5, ws.max_row + 1) if ws.cell(r, headers[key_field]).value == key]
            require(len(matches) <= 1, f"Duplicate key: {sheet}/{key}")
            if matches:
                row_number = matches[0]
                for field, value in values.items():
                    cell = ws.cell(row_number, headers[field])
                    if sheet in REPAIR_FIELDS and field in REPAIR_FIELDS[sheet]:
                        owned(ws, row_number, key)
                        change(cell, value, key, field)
                    elif value is not None and field != "comment":
                        require((cell.value or "") == value if value == "" else cell.value == value,
                                f"Existing non-repair field differs: {sheet}/{key}/{field}")
            else:
                require(sheet not in {"Text", "Object"}, f"Shared row must already exist: {sheet}/{key}")
                require(key == root or key.startswith(root + "."), f"Cannot append shared row: {sheet}/{key}")
                row_number = ws.max_row + 1
                for col in range(1, ws.max_column + 1):
                    ws.cell(row_number, col)._style = copy(ws.cell(max(5, row_number - 1), col)._style)
                for field, value in values.items():
                    if value is not None:
                        ws.cell(row_number, headers[field], value)
                ws.cell(row_number, 2).comment = Comment(json.dumps({"batch": batch,
                    "source_map_line": manifest["source_map_sha256"]}, ensure_ascii=False), "A-neutral preview")
                added[sheet] += 1
            locations.append({"sheet": sheet, "key": key, "row": row_number})

        ws = wb["ConversationLine"]
        for row in manifest["rows"]:
            if row["kind"] == "正文":
                continue
            key = row["key"].replace(".text.", ".line.")
            matches = [r for r in range(5, ws.max_row + 1) if ws.cell(r, 2).value == key]
            require(len(matches) <= 1, f"Duplicate obsolete option-line: {key}")
            for number in matches:
                owned(ws, number, key)
                cleared.append({"key": key, "row": number})
                for cell in ws[number][1:]:
                    if ws.cell(1, cell.column).value:
                        change(cell, None, key, ws.cell(1, cell.column).value)

        for sheet, field in (("Result", "eventId"), ("EventDate", "eventId"), ("Action", "actionId")):
            if sheet in wb.sheetnames:
                ws = wb[sheet]
                headers = {c.value: c.column for c in ws[1]}
                if field in headers:
                    require(not any(ws.cell(r, headers[field]).value == root for r in range(5, ws.max_row + 1)),
                            f"Result/date/action row already exists for {root}")
        for path, sha in hashes.items():
            require(digest(path) == sha, f"Concurrent change during plan: {path}")
        if not edits and not added:
            return {"idempotent": True, "changed": False, "added": {}, "edits": [],
                    "before_sha256": hashes[NARRATIVE], "after_sha256": hashes[NARRATIVE]}

        out = ROOT / "Library/Locus/tmp/neutral-content" / (datetime.now().strftime("%Y%m%d-%H%M%S-%f") + "-" + batch + "-repair")
        out.mkdir(parents=True, exist_ok=False)
        backup = out / "Narrative.before.xlsx"
        shutil.copy2(NARRATIVE, backup)
        require(digest(backup) == hashes[NARRATIVE], "Backup SHA mismatch")
        candidate = out / "Narrative.candidate.xlsx"
        wb.save(candidate)
        saved = openpyxl.load_workbook(candidate)
        allowed = {(e["sheet"], e["cell"]) for e in edits}
        try:
            require(saved.sheetnames == wb.sheetnames, "Candidate sheet order changed")
            for sheet, old_size in dimensions.items():
                require((saved[sheet].max_row, saved[sheet].max_column) ==
                        (wb[sheet].max_row, old_size[1]), f"Candidate dimensions changed: {sheet}")
            for key, old in before.items():
                actual = signature(saved[key[0]][key[1]])
                if key in allowed:
                    expected = signature(wb[key[0]][key[1]])
                    require(actual == expected, f"Repair readback mismatch: {key}")
                    require(actual[2:] == old[2:], f"Repair changed metadata: {key}")
                else:
                    require(actual == old, f"Unowned or non-repair cell changed: {key}")
            for s in wb:
                for row in s.iter_rows(min_row=dimensions[s.title][0] + 1):
                    for cell in row:
                        require(signature(saved[s.title][cell.coordinate]) == signature(cell),
                                f"Append readback mismatch: {s.title}/{cell.coordinate}")
        finally:
            saved.close()
        require(allowed <= owned_cells, "Repair includes non-owned cells")
        candidate_sha = digest(candidate)
        # Retain the verified candidate as evidence; replace from a separate sibling file.
        replacement = out / "Narrative.replace.xlsx"
        shutil.copy2(candidate, replacement)
        require(digest(replacement) == candidate_sha, "Replacement SHA mismatch")
        for path, sha in hashes.items():
            require(digest(path) == sha, f"Concurrent change before CAS replace: {path}")
        replacement.replace(NARRATIVE)
        require(digest(NARRATIVE) == candidate_sha, "Post-CAS SHA mismatch")
        require(digest(TEXT) == hashes[TEXT], "Text changed during CAS")
        report = {"idempotent": False, "changed": True, "batch": batch,
                  "added": dict(added), "edits": edits, "cleared_option_lines": cleared,
                  "locations": locations, "evidence": str(out),
                  "before_sha256": hashes[NARRATIVE], "after_sha256": candidate_sha,
                  "text_sha256": hashes[TEXT], "source_map_sha256": hashes[MAP],
                  "old_cells_checked": len(before), "changed_cells": len(allowed),
                  "preserved_cells": len(before) - len(allowed),
                  "non_owned_cells_checked_unchanged": len(set(before) - owned_cells),
                  "metadata_preserved": True}
        (out / "repair-evidence.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        return report
    finally:
        for book in books.values():
            book.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--event", choices=sorted(EVENTS), required=True, help="Configure exactly one A02-A05 root")
    parser.add_argument("--apply", action="store_true", help="CAS append/repair owned Narrative rows; Text read-only; never run generation")
    args = parser.parse_args()
    rows, manifest = build(args.event)
    # The default is a read-only in-memory plan; --apply is intentionally explicit.
    if not args.apply:
        counts = Counter(item["sheet"] for item in rows)
        print(json.dumps({**manifest, "would_append": dict(counts), "writes": False}, ensure_ascii=False, indent=2))
        return
    print(json.dumps(apply(args.event, rows, manifest), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
