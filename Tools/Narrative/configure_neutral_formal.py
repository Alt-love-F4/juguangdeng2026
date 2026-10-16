"""Configure authorized A01-A04 SOURCE tables only; never generate or call Unity.

Default is a read-only plan. Save that JSON and pass --apply
--expected-sha256 PLAN.json to require its SHA-256 values for all three books.
Candidates are backed up, reloaded, checked cell-by-cell (including metadata),
and checked for an in-memory idempotent second pass before cooperative CAS.
An already configured source is a byte-preserving no-op. Old blocked builders
are never called. Source-map/DOCX remain immutable evidence, not display edits.
"""
from __future__ import annotations

import argparse
from collections import Counter
from contextlib import contextmanager
from copy import copy
from datetime import datetime
import hashlib
import itertools
import json
import os
from pathlib import Path
import shutil
import uuid
from zipfile import ZIP_DEFLATED, ZipFile
import xml.etree.ElementTree as ET

import openpyxl
from openpyxl.comments import Comment
from openpyxl.writer.excel import ExcelWriter

# Read-only repair primitives; neither repair() nor any blocked build/apply runs.
from repair_neutral_stage import metadata, snapshot, signature_sha256
from import_neutral_text import source_rows, verify_docx

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / 'Assets/ConfigTables/Data'
NARRATIVE = DATA / 'Narrative/Narrative.xlsx'
TEXT = DATA / 'Shared/Text.xlsx'
PARAM = DATA / 'Simulation/GameParam.xlsx'
MAP = ROOT / 'Docs/superpowers/specs/2026-10-16-content-batches/A-neutral/source-map.md'
BOOKS = (NARRATIVE, TEXT, PARAM)
LOCK = ROOT / 'Library/Locus/tmp/narrative-source-write.lock'
A01, A02, A03, A04 = ('neutral.' + x for x in ('1008.s1', '1010.s1', '1010.s2', '1013.s2'))
ROOTS = (A01, A02, A03, A04)
B_COMPLETED = 'story.rose.1009.s1.completed'
AUTHORIZATION = ('User authorization: A01-A04 formal source tables; relationship ranges -2..5, '
                 'init/reverse 0; A02/A03 day5 MorningBeforeMap priority100/90; A04 day8 '
                 'MorningBeforeMap, q2 wisdom>=2, real approval flag, two unconditional waits; '
                 'display-only time-word amendments; economics deferred, GameParam only.')
PARAMS = {
    'neutral.promotion.preparation_cost_multiplier': ('0.8', 'User authorized adjustable placeholder; NOT connected to accounting; no Money result.'),
    'neutral.promotion.revenue_multiplier': ('1.2', 'User authorized adjustable placeholder; NOT connected to accounting; no Money result.'),
    'neutral.reporter.advertisement_fee': ('100', 'User authorized provisional A05 fee; parent-owned consumer must synchronize balance guard/payment before generation.'),
}
# Exact display-only before/after authorization. The original DOCX and map survive.
TEXT_EDITS = {
    A02 + '.text.p066': ('那就定在两天后，怎么样？', '那就定在三天后，怎么样？'),
    A02 + '.text.p067': ('两天后……好！时间应该足够了！', '三天后……好！时间应该足够了！'),
    A02 + '.text.p068': ('那今天营业结束后，我们就一起商量一下具体的安排。', '那就先按这个安排准备，今天营业结束后，我们再一起核对一下细节。'),
    A03 + '.text.p003': ('清晨的咖啡馆还没有正式营业。', '早间的第一拨客人散去后，咖啡馆迎来了一段短暂的整理时间。'),
    A03 + '.text.p004': ('为了准备即将到来的促销活动，我比往常早起了一些。桌椅已经摆放妥当，柜台上堆着几张写了一半的告示，旁边则是昨晚整理好的账簿。', '为了准备即将到来的促销活动，我趁客流稍缓整理了一下店里。桌椅已经摆放妥当，柜台上堆着几张写了一半的告示，旁边则是昨晚整理好的账簿。'),
    A03 + '.text.p007': ('早安。请问这里已经开始营业了吗？', '早上好。请问现在方便接待吗？'),
    A03 + '.text.p011': ('看来我来得稍早了一些。不过，既然门没有锁上，想必也不至于打扰太久。', '看来我来得不太凑巧。不过，既然门没有锁上，想必也不至于打扰太久。'),
    A03 + '.text.p012': ('您好，先生。我们确实还没正式营业。', '您好，先生。我们正在整理店里。'),
    A03 + '.text.p025': ('您是第一次来这里吗？虽然还没正式营业，不过咖啡已经准备好了。若是不介意，不妨进来坐坐。', '您是第一次来这里吗？我们正在整理，不过咖啡已经准备好了。若是不介意，不妨进来坐坐。'),
    A03 + '.text.p026': ('如此说来，我倒没有拒绝的理由了。能在清晨遇到一位热情的店主，总是件令人愉快的事。', '如此说来，我倒没有拒绝的理由了。能在早间遇到一位热情的店主，总是件令人愉快的事。'),
    A03 + '.text.p030': ('您要是愿意的话，也可以先坐一会儿。我们很快就能营业了。', '您要是愿意的话，也可以先坐一会儿。我们很快就能整理好。'),
    A03 + '.text.p034': ('抱歉，先生，我们还需要准备一会儿。如果您不介意，可以等正式营业之后再来。', '抱歉，先生，我们还需要整理一会儿。如果您不介意，可以稍后再来。'),
    A03 + '.text.p035': ('当然，遵守营业时间也是经营的一部分。既然如此，我便不打扰了。', '当然，先把店里安排妥当也是经营的一部分。既然如此，我便不打扰了。'),
    A03 + '.text.p052': ('正是如此。所以我们打算在两天后举办一次促销活动。', '正是如此。所以我们打算在三天后举办一次促销活动。'),
    A03 + '.text.p113': ('距离正式营业已经没有多少时间，我收拾好桌上的账簿，将写了一半的促销告示先放到一旁，准备迎接今天的第一批客人。', '短暂的整理告一段落，我收拾好桌上的账簿，将写了一半的促销告示先放到一旁，准备继续招呼今天的客人。'),
}
A04_PAGES = ('morning cooperate organize_pass organize_fail promote_pass promote_fail '
             'wait_business afternoon visit longterm_pass longterm_fail accounts_pass '
             'accounts_fail present visit_tail wait_night closing').split()
FLAGS = {
    B_COMPLETED: '玫瑰线十月九日密见已完成（仅声明，B负责生产）',
    'story.' + A01 + '.completed': '十月八日市场初遇已完成',
    'story.' + A02 + '.completed': '十月十日促销筹备已完成',
    'story.' + A02 + '.frugal_preparation': '本次促销精打细算准备资格（未接账）',
    'story.' + A03 + '.completed': '十月十日安托万到访已完成',
    'story.' + A03 + '.antoine_approval': '安托万赏识（控制十月十三日到场）',
    'story.' + A04 + '.completed': '十月十三日促销已完成',
    'story.' + A04 + '.orderly_service': '本次促销井然有序资格（未接账）',
    'story.' + A04 + '.warm_recommendation': '本次促销热情推介资格（未接账）',
}


def require(ok, message):
    if not ok:
        raise RuntimeError(message)


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


@contextmanager
def write_guard():
    """Use the same cooperative lock path as other narrative writers, with our identity."""
    LOCK.parent.mkdir(parents=True, exist_ok=True)
    owner = json.dumps(dict(owner='A-neutral-formal', pid=os.getpid(), token=uuid.uuid4().hex,
                           time=datetime.now().astimezone().isoformat()), sort_keys=True)
    try:
        with LOCK.open('x', encoding='utf-8') as stream:
            stream.write(owner)
    except FileExistsError:
        raise RuntimeError(f'Source-write lock occupied; do not remove another owner: {LOCK}') from None
    try:
        yield
    finally:
        require(LOCK.exists() and LOCK.read_text(encoding='utf-8') == owner,
                'Source-write lock ownership changed; preserving it')
        LOCK.unlink()


def load_book(path):
    wb = openpyxl.load_workbook(path)
    with ZipFile(path) as archive:
        if 'docProps/core.xml' not in archive.namelist():
            # openpyxl invents current timestamps when the source has no core part.
            # They are not authored metadata. Normalize to a stable writer-safe
            # sentinel; preserve_core_part removes the invented core part again.
            wb.properties.created = wb.properties.modified = datetime(2000, 1, 1)
    return wb


def preserve_core_part(source, candidate):
    """Keep the authored core XML byte-for-byte, including an absent core part."""
    with ZipFile(source) as archive:
        core = archive.read('docProps/core.xml') if 'docProps/core.xml' in archive.namelist() else None
        # Some source producers author empty t="str" cells. openpyxl drops them,
        # losing data_type metadata. Retain those exact empty cells in the XML.
        empty_strings = {}
        ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        for name in archive.namelist():
            if name.startswith('xl/worksheets/sheet') and name.endswith('.xml'):
                cells = [c for c in ET.fromstring(archive.read(name)).findall('.//s:sheetData/s:row/s:c', ns)
                         if c.get('t') == 'str' and not list(c)]
                if cells:
                    empty_strings[name] = cells
    repack = candidate.with_suffix('.core-preserved.xlsx')
    with ZipFile(candidate) as old, ZipFile(repack, 'w', ZIP_DEFLATED, allowZip64=True) as new:
        for info in old.infolist():
            data = old.read(info.filename)
            if info.filename == 'docProps/core.xml':
                if core is None:
                    continue
                data = core
            elif info.filename in empty_strings:
                tree = ET.fromstring(data)
                rows = {row.get('r'): row for row in tree.findall('.//s:sheetData/s:row', ns)}
                for original in empty_strings[info.filename]:
                    coordinate = original.get('r')
                    row = rows[''.join(c for c in coordinate if c.isdigit())]
                    current = next((c for c in row if c.get('r') == coordinate), None)
                    if current is not None:
                        require(not list(current), f'Empty-string preservation encountered value {coordinate}')
                        row.remove(current)
                    row.insert(0, original)
                data = ET.tostring(tree, encoding='utf-8')
            elif core is None and info.filename in ('[Content_Types].xml', '_rels/.rels'):
                tree = ET.fromstring(data)
                for child in list(tree):
                    if child.get('PartName') == '/docProps/core.xml' or child.get('Type', '').endswith('/metadata/core-properties'):
                        tree.remove(child)
                data = ET.tostring(tree, encoding='utf-8')
            new.writestr(info, data)
    repack.replace(candidate)


def heads(ws):
    return {cell.value: cell.column for cell in ws[1] if cell.value}


def records(ws):
    h = heads(ws)
    return [dict(_row=r, **{key: ws.cell(r, col).value for key, col in h.items() if key != '##var'})
            for r in range(5, ws.max_row + 1) if ws.cell(r, 2).value is not None]


def locate(ws, key, owner=None):
    rows = [r for r in range(5, ws.max_row + 1) if ws.cell(r, 2).value == key
            and (owner is None or ws.cell(r, heads(ws)['ownerId']).value == owner)]
    require(len(rows) <= 1, f'Duplicate source key: {ws.title}/{owner}/{key}')
    return rows[0] if rows else None


class Configuration:
    def __init__(self, books):
        self.books = books
        self.n, self.t, self.p = (books[p] for p in BOOKS)
        self.edits, self.added, self.locations = [], [], []

    def append(self, book, sheet, values):
        ws = book[sheet]
        key = next(iter(values.values()))
        owner = values.get('ownerId') if sheet == 'Bar' else None
        row = locate(ws, key, owner)
        require(set(values) <= heads(ws).keys(), f'Schema mismatch {sheet}: {set(values) - heads(ws).keys()}')
        if row is None:
            row = ws.max_row + 1
            for col in range(1, ws.max_column + 1):
                ws.cell(row, col)._style = copy(ws.cell(row - 1, col)._style)
            for field, value in values.items():
                ws.cell(row, heads(ws)[field], value)
            ws.cell(row, 2).comment = Comment(json.dumps(dict(batch='A01-A04-formal',
                authorization='User A01-A04 formal source-table decision', owner=owner), ensure_ascii=False), 'A-neutral-formal')
            self.added.append(dict(book=self.path(book).name, sheet=sheet, row=row, key=key, values=values))
        else:
            # Existing shared keys are read-only. Never silently replace another batch.
            for field, value in values.items():
                require(ws.cell(row, heads(ws)[field]).value == value,
                        f'Existing different row: {sheet}/{key}/{field}')
        self.locations.append(dict(book=self.path(book).name, sheet=sheet, key=key, row=row, owner=owner))
        return row

    def path(self, book):
        return next(p for p, b in self.books.items() if b is book)

    def change(self, book, sheet, key, field, after, allowed_before):
        ws = book[sheet]
        row = locate(ws, key)
        require(row is not None, f'Missing required existing row {sheet}/{key}')
        cell = ws.cell(row, heads(ws)[field])
        require(cell.value in (*allowed_before, after), f'Unexpected old value: {sheet}/{key}/{field}: {cell.value!r}')
        if cell.value != after:
            self.edits.append(dict(book=self.path(book).name, sheet=sheet, row=row, cell=cell.coordinate,
                key=key, field=field, before=cell.value, after=after))
            cell.value = after
        return row

    def condition(self, key, owner_type, owner_id, kind, target, value):
        self.append(self.n, 'Condition', dict(conditionId=key, ownerType=owner_type, ownerId=owner_id,
                    groupIndex=1, condType=kind, target=target, value=value))

    def outcome(self, root, seq, suffix, kind, target, value, choice=None, check=None):
        key = root + '.result.' + suffix
        self.append(self.n, 'Result', dict(resultId=key, eventId=root, seq=seq, resultType=kind,
            target=target, value=value, note='User-authorized A01-A04 root-tail settlement; no accounting'))
        for label, guard in (('choice', 'choice:' + root + '.choice.' + choice if choice else None),
                             ('check', 'check:' + root + '.' + check if check else None)):
            if guard:
                self.condition(key + '.' + label, 'result', key, 'PRESENTATION', guard, 'true')

    def configure(self):
        for actor, label in (('colette', '柯莱特对玩家的好感'), ('adele', '阿黛尔对玩家的好感'), ('antoine', '安托万对玩家的好感')):
            require(locate(self.n['Object'], actor) is not None, f'Missing formal Object {actor}')
            name = f'bar.{actor}.trust_player.name'
            self.append(self.t, 'Text', dict(key=name, text=label, comment='User-authorized formal relation label'))
            self.append(self.n, 'Bar', dict(barId='trust_player', ownerId=actor, targetId='player',
                barType='relation', min=-2, max=5, direction='highIsGood', advanceMode='event_driven',
                initValue=0, reverseInitValue=0, nameText=name, autoStep=0, sourceMode='ncs'))
        for key, label in FLAGS.items():
            name = 'condition.' + key
            self.append(self.t, 'Text', dict(key=name, text=label, comment='User-authorized A01-A04 flag label; B completion is declared only'))
            self.append(self.n, 'ConditionKey', dict(key=key, nameText=name, initialValue=False, onlySetTrue=True))
        for root, old in ((A01, 'entry=MapAction;blocksSave=true;blockedBy=A-B010-relationship-definition'),
                          (A02, 'entry=MapAction;blocksSave=true;blockedBy=A-B002,A-B003,A-B004,A-B005,A-B010'),
                          (A03, 'entry=MorningBeforeMap;blocksSave=true;blockedBy=A-B003,A-B004,A-B010'),
                          (A04, 'entry=MapAction;blocksSave=true;blockedBy=A-B002,A-B004,A-B005,A-B006,A-B010,A-CAP01,A-CAP04')):
            self.change(self.n, 'Event', root, 'eventProperty',
                        f'entry={"MapAction" if root == A01 else "MorningBeforeMap"};blocksSave=true', (old,))
        # Only A01 gets a map Action; A02-A04 are morning dates, never blank shells.
        action = A01 + '.action'
        self.append(self.t, 'Text', dict(key=action + '.name', text='大堂区采购', comment='A01 P008 source-backed map action label; no purchase settlement'))
        self.append(self.n, 'Action', dict(actionKey=action, type='ONCE', mustExecute=True,
            nameText=action + '.name', requirements=None, resultEvents=A01, durationDays=0,
            progressTarget=0, deliveryMode='None', replacesLocationActions=False, focusLocationOnComplete=False))
        ws = self.n['MapPosition']
        row = locate(ws, 'market')
        require(row is not None, 'Missing market map position')
        old_mount = ws.cell(row, heads(ws)['extendedActions']).value
        mounts = [x.strip() for x in (old_mount or '').split(',') if x.strip()]
        require(mounts.count(action) <= 1, 'Duplicate A01 map mount')
        if action not in mounts:
            mounts.append(action)
        self.change(self.n, 'MapPosition', 'market', 'extendedActions', ','.join(mounts), (old_mount,))
        for owner_type, owner_id, suffix in (('event', A01, ''), ('action', action, '.action')):
            self.condition(A01 + suffix + '.day', owner_type, owner_id, 'BAR_BETWEEN', 'world.calendar', '3,3')
            self.condition(A01 + suffix + '.map', owner_type, owner_id, 'BOOL_EQ', 'world.map.active', 'true')
            self.condition(A01 + suffix + '.uncompleted', owner_type, owner_id, 'BOOL_EQ', 'story.' + A01 + '.completed', 'false')
        for root, day, priority, previous in ((A02, 5, 100, B_COMPLETED),
                                             (A03, 5, 90, 'story.' + A02 + '.completed'),
                                             (A04, 8, 100, 'story.' + A02 + '.completed')):
            self.append(self.n, 'EventDate', dict(eventId=root, day=day, timeSlot='morning', repeatRule='once', priority=priority))
            self.condition(root + '.day', 'event', root, 'BAR_BETWEEN', 'world.calendar', f'{day},{day}')
            self.condition(root + '.uncompleted', 'event', root, 'BOOL_EQ', 'story.' + root + '.completed', 'false')
            self.condition(root + '.previous', 'event', root, 'BOOL_EQ', previous, 'true')
        self.append(self.n, 'AbilityCheck', dict(checkId=A04 + '.q2_wisdom', mode='Inline',
                    comment='User resolves A04 q2B discrepancy: wisdom >= 2; source evidence unchanged'))
        self.append(self.n, 'AbilityCheckCondition', dict(conditionId=A04 + '.q2_wisdom.1',
            checkId=A04 + '.q2_wisdom', attributeId='wisdom', threshold=2, groupIndex=1, seq=1))
        # Preserve every other field, including stage/background and existing comments.
        for suffix, kind in (('wait_business', 'WaitBusiness'), ('wait_night', 'WaitNight')):
            self.append(self.n, 'ConversationPage', dict(pageId=A04 + '.' + suffix,
                presentationKey=A04 + '.conversation', seq=A04_PAGES.index(suffix) + 1, stepType=kind))
        original_pages = [x for x in A04_PAGES if not x.startswith('wait_')]
        for seq, suffix in enumerate(A04_PAGES, 1):
            if suffix.startswith('wait_'):
                continue
            self.change(self.n, 'ConversationPage', A04 + '.' + suffix, 'seq', seq,
                        (original_pages.index(suffix) + 1,))
            if suffix in ('visit', 'longterm_pass', 'longterm_fail', 'accounts_pass', 'accounts_fail', 'present', 'visit_tail'):
                ws = self.n['ConversationPage']
                row = locate(ws, A04 + '.' + suffix)
                current = ws.cell(row, heads(ws)['pageCondKeys']).value
                updated = current.replace('dependency:A04.antoine_approval', 'story.' + A03 + '.antoine_approval')
                updated = updated.replace('dependency:A04.q2B.wisdom-pass', 'check:' + A04 + '.q2_wisdom')
                self.change(self.n, 'ConversationPage', A04 + '.' + suffix, 'pageCondKeys', updated, (current,))
        self.outcome(A01, 1, 'q1a.colette', 'REL_ADD', 'rel:colette>player', '+2', 'q1.a')
        self.outcome(A01, 2, 'q1b.colette', 'REL_ADD', 'rel:colette>player', '+1', 'q1.b', 'language')
        self.outcome(A02, 1, 'q1a.adele', 'REL_ADD', 'rel:adele>player', '+1', 'q1.a')
        self.outcome(A02, 2, 'q1b.frugal_preparation', 'SET_FLAG', 'story.' + A02 + '.frugal_preparation', 'true', 'q1.b', 'wisdom')
        self.outcome(A02, 3, 'q1c.adele', 'REL_ADD', 'rel:adele>player', '+2', 'q1.c', 'language')
        self.outcome(A03, 1, 'q1b.antoine', 'REL_ADD', 'rel:antoine>player', '+1', 'q1.b', 'q1_language')
        self.outcome(A03, 2, 'q2a.antoine', 'REL_ADD', 'rel:antoine>player', '+2', 'q2.a', 'q2_language')
        self.outcome(A03, 3, 'q2a.antoine_approval', 'SET_FLAG', 'story.' + A03 + '.antoine_approval', 'true', 'q2.a', 'q2_language')
        self.outcome(A04, 1, 'q1a.adele', 'REL_ADD', 'rel:adele>player', '+1', 'q1.a')
        self.outcome(A04, 2, 'q1b.orderly_service', 'SET_FLAG', 'story.' + A04 + '.orderly_service', 'true', 'q1.b', 'q1_wisdom')
        self.outcome(A04, 3, 'q1c.warm_recommendation', 'SET_FLAG', 'story.' + A04 + '.warm_recommendation', 'true', 'q1.c', 'q1_language')
        self.outcome(A04, 4, 'q2a.antoine', 'REL_ADD', 'rel:antoine>player', '+2', 'q2.a', 'q2_language')
        self.outcome(A04, 5, 'q2b.antoine', 'REL_ADD', 'rel:antoine>player', '+2', 'q2.b', 'q2_wisdom')
        for root, seq in ((A01, 3), (A02, 4), (A03, 4), (A04, 6)):
            self.outcome(root, seq, 'completed', 'SET_FLAG', 'story.' + root + '.completed', 'true')
        for key, (before, after) in TEXT_EDITS.items():
            self.change(self.t, 'Text', key, 'text', after, (before,))
        for key, (value, comment) in PARAMS.items():
            self.append(self.p, self.p.sheetnames[0], dict(key=key, value=value, comment=comment))
        return self


def validate(books):
    """Static source contract + branch traversal. Not an Editor/runtime acceptance test."""
    n, t, p = (books[x] for x in BOOKS)
    texts = {r['key']: r['text'] for r in records(t['Text'])}
    conditions, results = records(n['Condition']), records(n['Result'])
    own_results = [r for r in results if r['eventId'] in ROOTS]
    require(len(own_results) == 17, 'Unexpected A01-A04 Result count')
    require(not any(r['target'].removeprefix('flag:') == B_COMPLETED for r in own_results), 'This package cannot produce B completion')
    require(not any('antoine_dialogue_unlocked' in str(r) for r in own_results), 'No fictional regular-activity unlock')
    keys = {r['key']: r for r in records(n['ConditionKey'])}
    for key in FLAGS:
        require(keys[key]['initialValue'] is False and keys[key]['onlySetTrue'] is True, f'Flag definition: {key}')
        require(keys[key]['nameText'] in texts, f'Missing flag label {key}')
    for key, (_, after) in TEXT_EDITS.items():
        require(texts[key] == after, f'Display amendment differs: {key}')
    for root in ROOTS:
        ev = next(r for r in records(n['Event']) if r['eventId'] == root)
        require(ev['onceOnly'] is True and 'blockedBy' not in ev['eventProperty'], f'Invalid formal Event {root}')
        group = [r for r in own_results if r['eventId'] == root]
        require(len({r['seq'] for r in group}) == len(group), f'Duplicate result sequence {root}')
        tail = max(group, key=lambda r: r['seq'])
        require(tail['target'] == 'story.' + root + '.completed', f'Completed must be root-tail {root}')
        require(not any(c['ownerType'] == 'result' and c['ownerId'] == tail['resultId'] for c in conditions), 'Completion must be unconditional')
    for result in own_results:
        require(result['resultType'] in ('REL_ADD', 'SET_FLAG'), 'No money/item/unlock result allowed')
        guards = [r for r in conditions if r['ownerType'] == 'result' and r['ownerId'] == result['resultId']]
        for guard in guards:
            require(guard['groupIndex'] == 1 and guard['condType'] == 'PRESENTATION' and guard['value'] == 'true', 'Result guard contract differs')
            require(guard['target'].startswith(('choice:', 'check:')), 'Unsupported result guard')
        if guards:
            require(guards[0]['target'].startswith('choice:'), 'Choice guard must precede check')
    dates = [r for r in records(n['EventDate']) if r['eventId'] in ROOTS]
    require([(r['eventId'], r['day'], r['timeSlot'], r['repeatRule'], r['priority']) for r in dates] ==
            [(A02, 5, 'morning', 'once', 100), (A03, 5, 'morning', 'once', 90),
             (A04, 8, 'morning', 'once', 100)], 'Formal dates/priorities differ')
    action = next(r for r in records(n['Action']) if r['actionKey'] == A01 + '.action')
    require((action['type'], action['mustExecute'], action['requirements'], action['resultEvents'],
             action['replacesLocationActions'], action['focusLocationOnComplete']) ==
            ('ONCE', True, None, A01, False, False), 'A01 map action differs')
    require(not any(r['actionKey'].startswith((A02, A03, A04)) for r in records(n['Action'])),
            'Morning events must not create map actions')
    for root, previous in ((A02, B_COMPLETED), (A03, 'story.' + A02 + '.completed'),
                           (A04, 'story.' + A02 + '.completed')):
        prev = [c for c in conditions if c['ownerType'] == 'event' and c['ownerId'] == root and c['target'] == previous]
        require(len(prev) == 1 and prev[0]['condType'] == 'BOOL_EQ' and prev[0]['value'] == 'true', 'Upstream guard differs')
    require(not any(c['ownerType'] == 'event' and c['ownerId'] == A04 and c['target'] == 'story.' + A03 + '.antoine_approval' for c in conditions),
            'Approval cannot gate entire A04 root')
    param_values = {r['key']: r['value'] for r in records(p[p.sheetnames[0]])}
    require(all(param_values[k] == v[0] for k, v in PARAMS.items()), 'GameParam value differs')
    relation_rows = [r for r in records(n['Bar']) if r['ownerId'] in ('colette', 'adele', 'antoine') and r['targetId'] == 'player']
    require(len(relation_rows) == 3, 'Exactly three formal relationships required')
    for r in relation_rows:
        if r['ownerId'] in ('colette', 'adele', 'antoine') and r['targetId'] == 'player':
            require((r['barType'], r['min'], r['max'], r['initValue'], r['reverseInitValue'], r['advanceMode'], r['sourceMode'], r['autoStep']) ==
                    ('relation', -2, 5, 0, 0, 'event_driven', 'ncs', 0), 'Relation contract mismatch')
    pages = records(n['ConversationPage'])
    a04 = sorted((r for r in pages if r['presentationKey'] == A04 + '.conversation'), key=lambda r: r['seq'])
    require([r['pageId'] for r in a04] == [A04 + '.' + s for s in A04_PAGES], 'A04 order differs')
    for row in a04:
        require('dependency:' not in (row['pageCondKeys'] or ''), 'Unresolved A04 dependency')
        guards = (row['pageCondKeys'] or '').split('|')
        for index, guard in enumerate(guards):
            if 'check:' in guard:
                require(any(g.startswith('choice:') for g in guards[:index]), 'Page check precedes selected choice')
        if row['stepType'] in ('WaitBusiness', 'WaitNight'):
            require(not any(row.get(f) for f in ('pageCondKeys', 'pageBarConds', 'optionGroupKey', 'readingKey', 'checkId')), 'Wait must be unconditional and empty')
            require(not any(r['pageId'] == row['pageId'] for r in records(n['ConversationLine'])), 'Wait contains dialogue')
    # All pre-existing source content/choices retain exactly the mapped Text IDs.
    for event in ('1008', '1010.s1', '1010.s2', '1013.s2'):
        for _, kind, key, _ in source_rows(event):
            require(key in texts, f'Missing source Text {key}')
            if kind == '正文':
                line = locate(n['ConversationLine'], key.replace('.text.', '.line.'))
                require(line is not None, f'Missing source line {key}')
                require(n['ConversationLine'].cell(line, heads(n['ConversationLine'])['textKey']).value == key, f'Line/Text mismatch {key}')
    # Exhaustive stable-attribute route model reading ACTUAL page and result guards.
    checks = {r['conditionId'].removesuffix('.1'): (r['attributeId'], r['threshold'])
              for r in records(n['AbilityCheckCondition']) if r['checkId'].startswith(ROOTS)}
    trace, unique_routes = [], {root: set() for root in ROOTS}
    for root in ROOTS:
        presentation = root if root == A01 else root + '.conversation'
        ordered = sorted((r for r in pages if r['presentationKey'] == presentation), key=lambda r: r['seq'])
        for language, wisdom, approval, first, second in itertools.product(range(1, 4), range(1, 4), (False, True), 'abc', 'abc'):
            attrs = dict(language=language, wisdom=wisdom)
            choices, cached, shown, waits = set(), {}, [], []
            choice_events = []
            def read(key):
                if key.startswith('not:'):
                    return not read(key[4:])
                if key.startswith('choice:'):
                    return key[7:] in choices
                if key.startswith('anyChoice:'):
                    return any(x in choices for x in key[10:].split(';'))
                if key.startswith('check:'):
                    name = key[6:]
                    attribute, threshold = checks[name]
                    if name not in cached:
                        cached[name] = attrs[attribute] >= threshold
                    return cached[name]
                require(key == 'story.' + A03 + '.antoine_approval', f'Unknown model key {key}')
                return approval
            for page in ordered:
                if not all(read(k) for k in (page['pageCondKeys'] or '').split('|') if k):
                    continue
                shown.append(page['pageId'].removeprefix(root + '.'))
                if page['stepType'] in ('WaitBusiness', 'WaitNight'):
                    waits.append(page['stepType'])
                if page['optionGroupKey']:
                    q = page['optionGroupKey'].split('.')[-1]
                    choice = root + '.choice.' + q + '.' + (first if q == 'q1' else second)
                    require(locate(n['Choice'], choice) is not None, f'Unknown selected choice {choice}')
                    choices.add(choice)
                    choice_events.append(q)
            applied = []
            for result in (r for r in own_results if r['eventId'] == root):
                guards = [c for c in conditions if c['ownerType'] == 'result' and c['ownerId'] == result['resultId']]
                # Result reads do not create missing snapshots, exactly as controller.
                if all((g['target'][7:] in choices if g['target'].startswith('choice:') else cached.get(g['target'][6:]) is True) for g in guards):
                    applied.append(result)
            deltas = {actor: sum(int(r['value']) for r in applied if r['resultType'] == 'REL_ADD' and r['target'] == 'rel:' + actor + '>player') for actor in ('colette', 'adele', 'antoine')}
            flags = {r['target'] for r in applied if r['resultType'] == 'SET_FLAG'}
            require('story.' + root + '.completed' in flags, 'Missing common completion')
            if root == A01:
                expected = 2 if first == 'a' else 1 if first == 'b' and language >= 2 else 0
                require(deltas['colette'] == expected, 'A01 relation mismatch')
                require(len(shown) == 5 and 'join' in shown and 'outro' in shown, 'A01 route mismatch')
            elif root == A02:
                require(deltas['adele'] == (1 if first == 'a' else 2 if first == 'c' and language >= 2 else 0), 'A02 relation mismatch')
                require(('story.' + A02 + '.frugal_preparation' in flags) == (first == 'b' and wisdom >= 3), 'Frugal flag mismatch')
                require(len(shown) == 3 and shown[-1] == 'join', 'A02 route mismatch')
            elif root == A03:
                expected = (1 if first == 'b' and language >= 2 else 0) + (2 if first != 'c' and second == 'a' and language >= 3 else 0)
                require(deltas['antoine'] == expected, 'A03 relation mismatch')
                require(('story.' + A03 + '.antoine_approval' in flags) == (first != 'c' and second == 'a' and language >= 3), 'Approval mismatch')
                require(('serve' in shown) == (first != 'c') and ('refuse_tail' in shown) == (first == 'c') and shown[-1] == 'common_end', 'A03 branch isolation mismatch')
                require(not (first == 'c' and (cached or 'q2' in choice_events)), 'Refusal executed a check/q2')
            else:
                require(waits == ['WaitBusiness', 'WaitNight'] and shown[-1] == 'closing', 'Both A04 waits required for every path')
                require(('visit' in shown) == approval and ('visit_tail' in shown) == approval, 'A04 approval branch mismatch')
                require(deltas['adele'] == (1 if first == 'a' else 0), 'A04 Adele mismatch')
                expected = 2 if approval and ((second == 'a' and language >= 3) or (second == 'b' and wisdom >= 2)) else 0
                require(deltas['antoine'] == expected, 'A04 Antoine mismatch')
                require(('story.' + A04 + '.orderly_service' in flags) == (first == 'b' and wisdom >= 2), 'Orderly flag mismatch')
                require(('story.' + A04 + '.warm_recommendation' in flags) == (first == 'c' and language >= 3), 'Warm flag mismatch')
                require(not (not approval and any('.q2_' in c for c in cached)), 'No-approval path executed q2 check')
            for check in cached:
                suffix = check.removeprefix(root + '.')
                expected_choice = {'language': 'q1.b' if root == A01 else 'q1.c', 'wisdom': 'q1.b',
                    'q1_language': 'q1.b' if root == A03 else 'q1.c', 'q1_wisdom': 'q1.b',
                    'q2_language': 'q2.a', 'q2_wisdom': 'q2.b'}[suffix]
                require(root + '.choice.' + expected_choice in choices, f'Unselected check executed {check}')
            unique_routes[root].add(tuple(shown))
            trace.append(dict(root=root, language=language, wisdom=wisdom, approval=approval, first=first, second=second,
                              pages=shown, checks=cached, relations=deltas, flags=sorted(flags)))
    # A04 has 30 abstract combinations but only 26 with stable attributes:
    # same-attribute q1/q2 tests cannot disagree at an identical threshold.
    expected_counts = {A01: 4, A02: 5, A03: 12, A04: 26}
    actual_counts = {root: len(routes) for root, routes in unique_routes.items()}
    require(actual_counts == expected_counts, f'Stable-attribute route coverage differs: {actual_counts}')
    return dict(kind='Static source contract/model only; no Unity/runtime claim', scenarios=len(trace),
                routes=expected_counts, skipped_branch_checks_not_evaluated=True,
                a04_both_waits_all_paths=True, trace_sha256=hashlib.sha256(json.dumps(trace, sort_keys=True).encode()).hexdigest())


def run(apply=False, expected=None):
    initial = {p.name: digest(p) for p in BOOKS}
    if expected is not None:
        require(expected == initial, 'Source SHA differs from expected plan; re-plan all three books')
    require(not apply or expected is not None, '--apply requires --expected-sha256 PLAN.json')
    # Verify the original evidence without staging or importing any source-table rows.
    sources = {MAP: digest(MAP)}
    for event in ('1008', '1010.s1', '1010.s2', '1013.s2'):
        path = verify_docx(event, source_rows(event))
        sources[path] = digest(path)
    books = {p: load_book(p) for p in BOOKS}
    try:
        before = {p: snapshot(w) for p, w in books.items()}
        meta = {p: metadata(w) for p, w in books.items()}
        sizes = {p: {s.title: (s.max_row, s.max_column) for s in w} for p, w in books.items()}
        config = Configuration(books).configure()
        validation = validate(books)
        second = Configuration(books).configure()
        require(not second.edits and not second.added, 'In-memory idempotence failed')
        allowed = {p: {(e['sheet'], e['cell']) for e in config.edits if e['book'] == p.name} for p in BOOKS}
        unchanged = {p: {k: v for k, v in before[p].items() if k not in allowed[p]} for p in BOOKS}
        report = dict(tool='configure_neutral_formal', authorization=AUTHORIZATION, writes=False,
            changed=False, idempotent=not config.edits and not config.added,
            before_sha256=initial, after_sha256=initial.copy(),
            source_evidence_sha256={str(p): h for p, h in sources.items()},
            added_counts=dict(Counter(e['book'] + '/' + e['sheet'] for e in config.added)),
            added_rows=config.added, edits=config.edits, locations=config.locations,
            text_authorization=[dict(key=k, before=b, after=a, authority='Explicit user authorization; display Text only, original evidence unchanged') for k, (b, a) in TEXT_EDITS.items()],
            preserved_cells={p.name: dict(count=len(unchanged[p]), signature_sha256=signature_sha256(unchanged[p])) for p in BOOKS},
            validation=validation, generation_run=False, unity_called=False,
            deferred=['No promotion accounting connected; parameters/qualification flags only, no Money results.',
                      'No regular Antoine activity/dialogue unlock invented; approval is consumed by actual A04 visit/no-visit.',
                      'B completion is declared false only, never produced here. A02 cannot be claimed naturally reachable while B is blocked.',
                      'A05 fee parameter is parent-consumed; consumer must synchronize price/condition/payment and regenerate when changed.'])
        def cas(expected_hashes):
            for p in BOOKS:
                require(digest(p) == expected_hashes[p.name], f'Concurrent workbook change: {p}; CAS aborted')
            for p, sha in sources.items():
                require(digest(p) == sha, f'Original evidence changed: {p}; CAS aborted')
        cas(initial)
        if not apply or report['idempotent']:
            return report
        out = ROOT / 'Library/Locus/tmp/neutral-content' / (datetime.now().strftime('%Y%m%d-%H%M%S-%f') + '-A01-A04-formal')
        out.mkdir(parents=True, exist_ok=False)
        candidates, candidate_hashes, backups = {}, {}, {}
        for p, wb in books.items():
            backup = out / (p.stem + '.before.xlsx')
            shutil.copy2(p, backup)
            require(digest(backup) == initial[p.name], f'Backup SHA mismatch {p}')
            backups[p] = backup
            candidate = out / (p.stem + '.candidate.xlsx')
            # Workbook.save mutates core modified metadata; use helper's writer approach.
            with ZipFile(candidate, 'w', ZIP_DEFLATED, allowZip64=True) as archive:
                ExcelWriter(wb, archive).save()
            preserve_core_part(backup, candidate)
            candidates[p] = candidate
            candidate_hashes[p.name] = digest(candidate)
        reloaded = {p: load_book(c) for p, c in candidates.items()}
        try:
            for p, saved in reloaded.items():
                require(saved.sheetnames == books[p].sheetnames, f'Sheet order changed {p}')
                require(metadata(saved) == meta[p], f'Workbook/sheet metadata changed {p}')
                actual = snapshot(saved)
                intended = snapshot(books[p])
                require(actual == intended, f'Candidate readback differs from intended {p}')
                for key, old in before[p].items():
                    require(key in actual, f'Existing coordinate lost {p}/{key}')
                    if key in allowed[p]:
                        require(actual[key][2:] == old[2:], f'Edited cell metadata changed {p}/{key}')
                    else:
                        require(actual[key] == old, f'Non-target cell value/metadata changed {p}/{key}')
                for sheet, (old_rows, old_cols) in sizes[p].items():
                    require(saved[sheet].max_column == old_cols, f'Column schema changed {p}/{sheet}')
                    appended = [r['row'] for r in config.added if r['book'] == p.name and r['sheet'] == sheet]
                    require(saved[sheet].max_row == old_rows + len(appended), f'Unexpected append dimensions {p}/{sheet}')
                    require(appended == list(range(old_rows + 1, old_rows + len(appended) + 1)), f'Non-contiguous append {p}/{sheet}')
                actual_unchanged = {k: actual[k] for k in unchanged[p]}
                require(signature_sha256(actual_unchanged) == signature_sha256(unchanged[p]), f'Preserved-cell SHA differs {p}')
            require(validate(reloaded) == validation, 'Candidate contract differs from in-memory contract')
            again = Configuration(reloaded).configure()
            require(not again.edits and not again.added, 'Candidate readback not idempotent')
        finally:
            for wb in reloaded.values():
                wb.close()
        report.update(evidence=str(out), backup_sha256={p.name: digest(b) for p, b in backups.items()},
            candidate_sha256=candidate_hashes, metadata_preserved=True, candidate_readback=True,
            candidate_idempotent=True, status='verified-candidates')
        journal = out / 'formal-evidence.json'
        journal.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        replacement = {}
        for p, candidate in candidates.items():
            replacement[p] = out / (p.stem + '.replace.xlsx')
            shutil.copy2(candidate, replacement[p])
            require(digest(replacement[p]) == candidate_hashes[p.name], f'Replacement SHA mismatch {p}')
        current, committed = initial.copy(), []
        try:
            for p in BOOKS:
                cas(current)
                replacement[p].replace(p)
                committed.append(p)
                current[p.name] = candidate_hashes[p.name]
                require(digest(p) == current[p.name], f'Post-CAS hash mismatch {p}')
            cas(current)
        except BaseException:
            # A filesystem has no multi-file replace transaction. Roll back only our
            # exact writes; do not overwrite an intervening writer on failure.
            for p in reversed(committed):
                require(digest(p) == candidate_hashes[p.name], f'Cannot safely roll back concurrent change {p}; inspect {journal}')
                rollback = out / (p.stem + '.rollback.xlsx')
                shutil.copy2(backups[p], rollback)
                require(digest(rollback) == initial[p.name], 'Rollback backup mismatch')
                rollback.replace(p)
            report.update(status='rolled-back', writes=False)
            journal.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
            raise
        report.update(writes=True, changed=True, after_sha256=current, status='committed')
        journal.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        return report
    finally:
        for wb in books.values():
            wb.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--apply', action='store_true')
    parser.add_argument('--expected-sha256', type=Path, help='JSON plan containing before_sha256 for all three source books')
    args = parser.parse_args()
    expected = json.loads(args.expected_sha256.read_text(encoding='utf-8'))['before_sha256'] if args.expected_sha256 else None
    if args.apply:
        with write_guard():
            report = run(True, expected)
    else:
        report = run(False, expected)
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
