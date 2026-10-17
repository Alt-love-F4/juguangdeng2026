"""Populate all Rose 1011 display content and keep unavailable gameplay bridges blocked."""
from copy import copy
from datetime import datetime
import hashlib
import shutil
from pathlib import Path

import openpyxl

import import_rose_text as source

ROOT = source.ROOT
NARRATIVE = ROOT / 'Assets/ConfigTables/Data/Narrative/Narrative.xlsx'
TEXT = ROOT / 'Assets/ConfigTables/Data/Shared/Text.xlsx'
source_hashes = {path: hashlib.sha256(path.read_bytes()).hexdigest() for path in (NARRATIVE, TEXT)}
book = openpyxl.load_workbook(NARRATIVE)
text_book = openpyxl.load_workbook(TEXT)
entries = source.expected_rows('1011')
originals = {row[0]: row[4] for row in source.source_rows('1011')}


def append_row(workbook, sheet_name, values, keyfield=None):
    ws = workbook[sheet_name]
    headers = [cell.value for cell in ws[1]]
    key_name = keyfield or headers[1]
    key_index = headers.index(key_name)
    matches = [row for row in ws.iter_rows(min_row=5) if row[key_index].value == values[key_name]]
    if len(matches) > 1:
        raise AssertionError(f'Duplicate row {sheet_name}/{values[key_name]}')
    if matches:
        row_number = matches[0][0].row
        for key, value in values.items():
            existing = ws.cell(row_number, headers.index(key) + 1).value
            assert existing == value, f'Existing different row {sheet_name}/{values[key_name]}/{key}: {existing!r} != {value!r}'
        return
    row_number = ws.max_row + 1
    for column in range(1, ws.max_column + 1):
        ws.cell(row_number, column)._style = copy(ws.cell(5, column)._style)
    for key, value in values.items():
        ws.cell(row_number, headers.index(key) + 1, value)


def upsert_row(sheet_name, values):
    ws = book[sheet_name]
    headers = [cell.value for cell in ws[1]]
    key_index = 1
    matches = [row for row in ws.iter_rows(min_row=5) if row[key_index].value == values[headers[key_index]]]
    if len(matches) > 1:
        raise AssertionError(f'Duplicate row {sheet_name}/{values[headers[key_index]]}')
    if matches:
        row_number = matches[0][0].row
    else:
        row_number = ws.max_row + 1
        for column in range(1, ws.max_column + 1):
            ws.cell(row_number, column)._style = copy(ws.cell(5, column)._style)
    for key, value in values.items():
        ws.cell(row_number, headers.index(key) + 1, value)


def ensure_text(key, text, comment):
    ws = text_book['Text']
    headers = [cell.value for cell in ws[1]]
    matches = [row for row in ws.iter_rows(min_row=5) if row[1].value == key]
    if len(matches) > 1:
        raise AssertionError(f'Duplicate Text/{key}')
    if matches:
        assert matches[0][2].value == text, f'Existing text differs: {key}'
        return
    row_number = ws.max_row + 1
    for column in range(1, ws.max_column + 1):
        ws.cell(row_number, column)._style = copy(ws.cell(5, column)._style)
    ws.cell(row_number, 2, key)
    ws.cell(row_number, 3, text)
    ws.cell(row_number, 4, comment)


def purge_narrative_rows():
    # This script owns the Rose 1011 narrative rows. Rebuild them so reruns do
    # not leave obsolete page ids or duplicate sequence numbers from staging.
    owned_sheets = [
        'Event', 'EventDate', 'SetupParams', 'PresentationIndex',
        'ConversationPage', 'ConversationLine', 'Choice', 'Condition',
        'AbilityCheck', 'AbilityCheckCondition', 'ReadingPage', 'ReadingLine',
    ]
    prefix = 'rose.1011.s1'
    for sheet_name in owned_sheets:
        ws = book[sheet_name]
        rows = [row[0].row for row in ws.iter_rows(min_row=5)
                if isinstance(row[1].value, str)
                and (row[1].value == prefix or row[1].value.startswith(prefix + '.'))]
        for row_number in reversed(rows):
            ws.delete_rows(row_number, 1)


purge_narrative_rows()


for entry in entries:
    ensure_text(entry['key'], entry['text'], f"B1011 source {entry['source']}; source verified by import_rose_text.py")

for key, display in [
    ('condition.story.neutral.1010.s1.completed', '10月10日中立筹备已完成'),
    ('condition.story.rose.1011.s1.completed', '10月11日玫瑰剧情已完成'),
    ('condition.story.rose.1011.s1.sylvain_appreciation', '西尔万已认可店长'),
    ('rose.1011.s1.preview.choice_label', '選項'),
]:
    ensure_text(key, display, 'Rose 1011 source/condition display text')

for key, only_set_true in [
    ('story.neutral.1010.s1.completed', True),
    ('story.rose.1009.s1.completed', True),
    ('story.rose.1011.s1.completed', True),
    ('story.rose.1011.s1.sylvain_appreciation', True),
]:
    append_row(book, 'ConditionKey', dict(key=key, nameText='condition.' + key, initialValue=False, onlySetTrue=only_set_true))

root_key = 'rose.1011.s1.conversation'
upsert_row('Event', dict(
    eventId='rose.1011.s1', eventKind='actor', actorId='player', targetId=None, titleText=None,
    contentKey=None, setupType='Tpl_ArrivalConversation', setupKey='rose.1011.s1', onceOnly=True,
    logTag='B1011', presentationClass='Conversation', presentationKey=root_key, presentationLogTag='B1011',
    eventProperty='entry=MorningBeforeMap;blocksSave=true;blockedBy=B-B05,B-C02,B-C04,B-C07;contentBoundary=P047;previewTail=P052'))
append_row(book, 'EventDate', dict(eventId='rose.1011.s1', day=6, timeSlot='morning', repeatRule='once', priority=0))
append_row(book, 'SetupParams', dict(setupKey='rose.1011.s1', tplName='Tpl_ArrivalConversation', slotValues=None,
                                    skipAble=False, autoAdvance=False, checkId=None))
upsert_row('PresentationIndex', dict(presentationKey=root_key, presentationClass='Conversation',
                                     logTag='B1011 full source configured; production bridge blocked at P047'))

for condition_id, cond_type, target, value in [
    ('rose.1011.s1.day', 'BAR_BETWEEN', 'world.calendar', '6,6'),
    ('rose.1011.s1.neutral_ready', 'BOOL_EQ', 'story.neutral.1010.s1.completed', 'true'),
    ('rose.1011.s1.rose1009_ready', 'BOOL_EQ', 'story.rose.1009.s1.completed', 'true'),
    ('rose.1011.s1.incomplete', 'BOOL_EQ', 'story.rose.1011.s1.completed', 'false'),
]:
    append_row(book, 'Condition', dict(conditionId=condition_id, ownerType='event', ownerId='rose.1011.s1',
                                      groupIndex=1, condType=cond_type, target=target, value=value))

check_ids = {
    'sparrow': 'rose.1011.s1.check.sparrow',
    'portrait': 'rose.1011.s1.check.portrait',
    'wisdom': 'rose.1011.s1.check.wisdom',
}
for check_id, note, attribute, threshold in [
    (check_ids['sparrow'], 'P014-P017; perception >5 means int >=6; failure skips the insight page.', 'perception', 6),
    (check_ids['portrait'], 'P111-P114; perception >4 means int >=5; failure skips the memory page.', 'perception', 5),
    (check_ids['wisdom'], 'P080 correct answer; user-confirmed wisdom >=3 enters the long route, failure goes to source early_end.', 'wisdom', 3),
]:
    append_row(book, 'AbilityCheck', dict(checkId=check_id, successEventId=None, failureEventId=None,
        successResultOwnerId=None, failureResultOwnerId=None, comment=note, mode='Inline'))
    append_row(book, 'AbilityCheckCondition', dict(conditionId=check_id + '.1', checkId=check_id,
        attributeId=attribute, threshold=threshold, groupIndex=1, seq=1))

mixed_speakers = {
    52: {'a': 'sylvain', 'b': None, 'c': 'sylvain', 'd': 'adele', 'e': 'sylvain', 'f': None, 'g': 'sylvain'},
    85: {'a': 'player', 'b': 'player', 'c': 'player'},
    87: {'a': 'sylvain', 'b': 'sylvain'},
    110: {'a': 'sylvain', 'b': 'sylvain', 'c': 'sylvain'},
    116: {'a': 'sparrow', 'b': 'sparrow'},
    117: {'a': 'sylvain', 'b': 'sylvain', 'c': 'sylvain', 'd': 'sylvain'},
    122: {'a': 'sylvain', 'b': 'sylvain'},
    125: {'a': 'sylvain', 'b': 'sylvain'},
}


def base_speaker(raw):
    raw = raw.strip()
    if raw.startswith('{playername}'):
        return 'player'
    for prefix, speaker in [('阿黛尔', 'adele'), ('麻雀', 'sparrow'), ('西尔万', 'sylvain')]:
        if raw.startswith(prefix):
            return speaker
    return None


# pageId, sequence, condition, choice group, step, background, source paragraphs
page_specs = [
    ('arrival', 1, None, None, None, 'oct07.cafe.day', [3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13]),
    ('sparrow_perception', 2, None, None, 'PassiveCheck', None, [15, 16, 17]),
    ('drink_choice', 3, None, 'rose.1011.s1.choice.drink', None, None, [18]),
    ('drink_black', 4, 'choice:rose.1011.s1.choice.p020', None, None, None, [21]),
    ('drink_rich', 5, 'choice:rose.1011.s1.choice.p022', None, None, None, [23]),
    ('drink_milk', 6, 'choice:rose.1011.s1.choice.p024', None, None, None, [25]),
    ('news_choice', 7, None, 'rose.1011.s1.choice.news', None, None, [27, 28, 29, 30]),
    ('news_known', 8, 'choice:rose.1011.s1.choice.p032', None, None, None, [35]),
    ('news_unknown', 9, 'choice:rose.1011.s1.choice.p033', None, None, None, [37]),
    ('sylvain_arrives', 10, None, None, None, None, [38, 39, 40, 42, 43, 44, 45, 46, 47]),
    # P048-P051 remain represented in source text and the full preview. The
    # formal production handoff is intentionally blocked until the real cafe
    # production callback and grade mapping exist, so no fake page condition is
    # serialized here.
    ('brew_gate', 11, None, None, None, None, []),
    ('brew_low', 12, None, None, None, None, [50]),
    ('nature_choice', 13, None, 'rose.1011.s1.choice.nature', None, None, [52, 53, 54, 55, 56, 57, 58]),
    ('early_end_nature', 14, 'anyChoice:rose.1011.s1.choice.p060;rose.1011.s1.choice.p061', None, None, None, [68, 69, 70, 71, 72, 73]),
    ('nature_not_person', 15, 'choice:rose.1011.s1.choice.p062', 'rose.1011.s1.choice.reason', None, None, [75, 76, 77]),
    ('reason_early_end_wrong', 16, 'anyChoice:rose.1011.s1.choice.p079;rose.1011.s1.choice.p081', None, None, None, [68, 69, 70, 71, 72, 73]),
    ('reason_early_end_wisdom', 17, 'choice:rose.1011.s1.choice.p080|not:check:rose.1011.s1.check.wisdom', None, None, None, [68, 69, 70, 71, 72, 73]),
    ('reason_title', 18, 'choice:rose.1011.s1.choice.p080|check:rose.1011.s1.check.wisdom', None, None, None, [85, 86, 87, 88, 89, 90, 91, 92, 93]),
    ('history_yes', 19, 'specialization.player.letter.politics', None, None, None, [95, 96, 97, 98]),
    ('history_no', 20, 'not:specialization.player.letter.politics', None, None, None, [100]),
    ('history_lecture', 21, None, None, None, None, [103, 104, 105, 106, 107, 108, 109, 110]),
    ('portrait_memory', 22, None, None, 'PassiveCheck', None, [112, 113, 114]),
    ('history_return', 23, None, None, None, None, [116, 117, 118, 119, 120, 121, 122, 123, 124, 125]),
    ('card_notice', 24, None, None, None, None, [126]),
    ('card_read', 25, None, None, 'LetterRead', None, []),
    ('history_recognition', 26, 'specialization.player.letter.politics', None, None, None, [132, 133, 134]),
    ('full_end', 27, None, None, None, None, [136, 137, 138, 139, 140]),
]

reading_key = 'rose.1011.s1.card'
upsert_row('PresentationIndex', dict(presentationKey=reading_key, presentationClass='Reading',
                                     logTag='B1011 card reading; completion/cancel bridge remains blocked'))
append_row(book, 'ReadingPage', dict(pageId=reading_key + '.page', presentationKey=reading_key, seq=1,
    titleKey=None, imageKey=None, pageCondKeys=None, pageBarConds=None))
for seq, number in enumerate([128, 129, 130], 1):
    append_row(book, 'ReadingLine', dict(lineId=f'{reading_key}.p{number}', pageId=reading_key + '.page',
        seq=seq, mode='narration', speakerKey=None, textKey=f'rose.1011.s1.text.p{number}', voiceKey=None))

for page, seq, conditions, option_group, step_type, background, source_numbers in page_specs:
    page_id = f'rose.1011.s1.{page}'
    upsert_row('ConversationPage', dict(pageId=page_id, presentationKey=root_key, seq=seq,
        pageCondKeys=conditions, pageBarConds=None, optionGroupKey=option_group, bgOverride=background,
        bgmOverride=None, stepType=step_type, readingKey=reading_key if page == 'card_read' else None,
        checkId=check_ids['sparrow'] if page == 'sparrow_perception' else check_ids['portrait'] if page == 'portrait_memory' else None))
    line_seq = 0
    for entry in entries:
        number = int(entry['source'].split('P')[-1])
        if number not in source_numbers or entry['kind'] in ('选项', '名片正文'):
            continue
        line_seq += 1
        suffix = entry['key'].rsplit('.', 1)[-1]
        speaker = mixed_speakers.get(number, {}).get(suffix, base_speaker(originals[number]))
        upsert_row('ConversationLine', dict(
            lineId=f'{page_id}.p{number:03}.{suffix}', pageId=page_id, seq=line_seq,
            speakerKey=speaker, side='none', textKey=entry['key'], voiceKey=None, emote=None,
            sfxKey=None, accentFxKey=None, accentTargetKey=None, speakerNameKey=None))

choice_groups = [
    ('rose.1011.s1.choice.drink', [20, 22, 24]),
    ('rose.1011.s1.choice.news', [32, 33]),
    ('rose.1011.s1.choice.nature', [60, 61, 62]),
    ('rose.1011.s1.choice.reason', [79, 80, 81]),
]
for group, numbers in choice_groups:
    for seq, number in enumerate(numbers, 1):
        append_row(book, 'Choice', dict(choiceId=f'rose.1011.s1.choice.p{number:03}', groupKey=group,
            textKey=f'rose.1011.s1.text.p{number:03}', resultOwnerId=None, seq=seq,
            checkId=check_ids['wisdom'] if number == 80 else None))

# Menu-discoverable source preview, with no Event or Result and no gameplay state changes.
preview_key = 'rose.1011.s1.fulltext.preview'
preview_page = preview_key + '.page'
upsert_row('PresentationIndex', dict(presentationKey=preview_key, presentationClass='Conversation',
    logTag='B1011 all source-mapped dialogue and choices in source order; isolated preview'))
append_row(book, 'ConversationPage', dict(pageId=preview_page, presentationKey=preview_key, seq=1,
    pageCondKeys=None, pageBarConds=None, optionGroupKey=None, bgOverride='cafe.day', bgmOverride=None,
    stepType=None, readingKey=None, checkId=None))
for seq, entry in enumerate(entries, 1):
    number = int(entry['source'].split('P')[-1])
    suffix = entry['key'].rsplit('.', 1)[-1]
    is_choice = entry['kind'] == '选项'
    speaker = 'player' if is_choice else mixed_speakers.get(number, {}).get(suffix, base_speaker(originals[number]))
    append_row(book, 'ConversationLine', dict(
        lineId=f'rose.1011.s1.preview.p{number:03}.{suffix}', pageId=preview_page, seq=seq,
        speakerKey=speaker, side='none', textKey=entry['key'], voiceKey=None, emote=None,
        sfxKey=None, accentFxKey=None, accentTargetKey=None,
        speakerNameKey='rose.1011.s1.preview.choice_label' if is_choice else None))

stamp = datetime.now().strftime('%Y%m%d-%H%M%S')
out = ROOT / 'Library/Locus/tmp/rose-content' / f'{stamp}-1011-full'
out.mkdir(parents=True)
for path, name in ((NARRATIVE, 'Narrative'), (TEXT, 'Text')):
    shutil.copy2(path, out / f'{name}.before.xlsx')
book.save(out / 'Narrative.candidate.xlsx')
text_book.save(out / 'Text.candidate.xlsx')
if any(hashlib.sha256(path.read_bytes()).hexdigest() != digest for path, digest in source_hashes.items()):
    raise RuntimeError('Concurrent workbook change; candidate files preserved and sources not replaced')
(out / 'Narrative.candidate.xlsx').replace(NARRATIVE)
(out / 'Text.candidate.xlsx').replace(TEXT)
print(f'Configured Rose 1011: {len(entries)} source-mapped display lines/choices, {len(page_specs)} staged pages, four choice groups, three confirmed checks.')
print('Full source preview is available as conversation:rose.1011.s1.fulltext.preview.')
print('Formal event remains blocked at P047 by B-B05/B-C02/B-C04/B-C07; no completed/appreciation result is configured.')
print('Backup:', out)
