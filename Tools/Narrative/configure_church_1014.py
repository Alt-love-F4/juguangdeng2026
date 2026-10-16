"""Append blocked church.1014.s2 configuration with CAS backup; no generation or Editor work."""
from __future__ import annotations
import argparse, hashlib, json, shutil, re
from collections import Counter
from copy import copy
from datetime import datetime
from pathlib import Path
import openpyxl
from import_church_text import expected_rows, MAP, SOURCES

ROOT = Path(__file__).resolve().parents[2]
NARRATIVE = ROOT / 'Assets/ConfigTables/Data/Narrative/Narrative.xlsx'
TEXT = ROOT / 'Assets/ConfigTables/Data/Shared/Text.xlsx'
EVENT = 'church.1014.s2'
CONV = EVENT + '.conversation'
CHECK = EVENT + '.perception'
PROPERTY = 'entry=AfterBusiness;blocksSave=true;blockedBy=D-B02-faction-source,D-B03-activity-letter'
HEADERS = {
    'Object': 'objId objType nameText iconPath portraitPath portraitIndoorPath avtarPath descText tags',
    'Event': 'eventId eventKind actorId targetId titleText contentKey setupType setupKey onceOnly logTag presentationClass presentationKey presentationLogTag eventProperty',
    'PresentationIndex': 'presentationKey presentationClass logTag',
    'SetupParams': 'setupKey tplName slotValues skipAble autoAdvance checkId',
    'ConversationPage': 'pageId presentationKey seq pageCondKeys pageBarConds optionGroupKey bgOverride bgmOverride stepType readingKey checkId',
    'ConversationLine': 'lineId pageId seq speakerKey side textKey voiceKey emote sfxKey accentFxKey accentTargetKey speakerNameKey stageLeftKey stageRightKey stageCenterKey',
    'Choice': 'choiceId groupKey textKey resultOwnerId seq checkId',
    'AbilityCheck': 'checkId successEventId failureEventId successResultOwnerId failureResultOwnerId comment mode',
    'AbilityCheckCondition': 'conditionId checkId attributeId threshold groupIndex seq',
}

def digest(path): return hashlib.sha256(path.read_bytes()).hexdigest()
def signature(cell):
    return (cell.value, cell.data_type, tuple(cell._style) if cell.has_style else (0,) * 9,
            (cell.comment.text, cell.comment.author) if cell.comment else None,
            copy(cell.hyperlink))

def text_entries():
    wb = openpyxl.load_workbook(TEXT, read_only=True, data_only=False)
    try:
        found = {}
        for row in wb['Text'].iter_rows(min_row=5, values_only=True):
            key = row[1]
            if isinstance(key, str) and key.startswith(EVENT + '.'):
                assert key not in found, key
                found[key] = row[2]
        assert len(found) == 57, f'expected 51 dialogue + 6 choice Text rows, got {len(found)}'
        return found
    finally: wb.close()

def plan(legacy=False):
    found = text_entries()
    entries = expected_rows('1014.s2')
    assert found == {e['key']: e['text'] for e in entries}, 'Existing Text/source-map/DOCX mismatch'
    sources = {e['key']: e for e in entries}
    pages = [
        ('arrival', [4,6,8,9,10,11,12,14,15], None, 'Dialogue', None, None),
        ('notice_pass', [17,18,19,20,21], None, 'PassiveCheck', None, CHECK),
        ('notice_fail', [23,24,25,26], f'not:check:{CHECK}', 'Dialogue', None, None),
        ('letter_question', [27,29,30,'031.a','031.b',32,33,34,35], None, 'Dialogue', EVENT+'.service.choices', None),
        ('relief_question', [41,42,43,44,45,46,47,48], None, 'Dialogue', EVENT+'.relief.choices', None),
        ('praise', [54], 'choice:'+EVENT+'.p050', 'Dialogue', None, None),
        ('doubt', [57,58,59,60,61], 'choice:'+EVENT+'.p051', 'Dialogue', None, None),
        ('wish', [64,65], 'choice:'+EVENT+'.p052', 'Dialogue', None, None),
        ('depart', [67,68,69,70,71,73,74,75], None, 'Dialogue', None, None),
    ]
    if not legacy:
        pages[3] = ('letter_question', [29,30,'031.a','031.b',32,33,34,35], None, 'Dialogue', EVENT+'.service.choices', None)
        pages.insert(3, ('letter_intro', [27], None, 'Dialogue', None, None))
        pages.insert(4, ('envelope', [], None, 'LetterReveal', None, None))
        pages.insert(6, ('envelope_hide', [], None, 'LetterHide', None, None))
    rows=[]
    def add(sheet, **values): rows.append({'sheet':sheet,'values':values})
    add('Object', objId='church.thibaut', objType='npc', nameText='church.thibaut.name')
    add('Event', eventId=EVENT, eventKind='world', actorId='world', setupKey=EVENT, onceOnly=True,
        logTag='D1014S2', presentationClass='Conversation', presentationKey=CONV,
        presentationLogTag='D1014S2', eventProperty=PROPERTY)
    add('PresentationIndex', presentationKey=CONV, presentationClass='Conversation', logTag='D1014S2')
    add('SetupParams', setupKey=EVENT, tplName='Tpl_ArrivalConversation', slotValues='opts=', skipAble=False, autoAdvance=False)
    add('AbilityCheck', checkId=CHECK, comment='D1014 inline perception snapshot; blocked preview only', mode='Inline')
    add('AbilityCheckCondition', conditionId=CHECK+'.1', checkId=CHECK, attributeId='perception', threshold=2, groupIndex=1, seq=1)
    for seq,(short,nums,guard,step,choices,checkid) in enumerate(pages,1):
        page=EVENT+'.'+short
        add('ConversationPage', pageId=page, presentationKey=CONV, seq=seq, pageCondKeys=guard,
            optionGroupKey=choices, bgOverride='oct07.cafe.day', stepType=step, checkId=checkid)
        for line_seq,num in enumerate(nums,1):
            suffix = f'p{num:03}' if isinstance(num,int) else 'p'+num
            key=EVENT+'.'+suffix
            speaker = 'player' if int(suffix.split('.')[0][1:]) in {8,10,12,15,27,32,45,47,60,69} else None
            if not speaker and int(suffix.split('.')[0][1:]) in {9,11,14,31,35,41,46,48,54,58,61,65,68}:
                speaker='church.thibaut'
            source = sources[key+'.text']
            assert source['page'] == (EVENT+'.letter_question' if short=='letter_intro' else page)
            emote = re.search(r'（#([^）]+)）', source['original'])
            add('ConversationLine', lineId=key, pageId=page, seq=line_seq + (1 if short=='letter_question' and not legacy else 0), speakerKey=speaker, side='none', textKey=key+'.text', emote=emote.group(1) if emote else None, speakerNameKey='church.thibaut.name' if speaker=='church.thibaut' else None)
    for group, nums in [('service',[37,38,39]),('relief',[50,51,52])]:
        for seq,num in enumerate(nums,1):
            key=f'{EVENT}.p{num:03}'
            add('Choice', choiceId=key, groupKey=f'{EVENT}.{group}.choices', textKey=key+'.text', seq=seq)
    assert len(rows)==1+1+1+1+1+1+(9 if legacy else 12)+51+6
    return rows, found

def prepare(wb, rows, text_wb, migrations=None):
    locations=[]; added=Counter()
    legacy={(r['sheet'],next(iter(r['values'].values()))):r['values'] for r in plan(legacy=True)[0]}
    for item in rows:
        ws=wb[item['sheet']]; values=item['values']; heads={c.value:c.column for c in ws[1]}
        assert [c.value for c in ws[1]] == ['##var'] + HEADERS[item['sheet']].split(), item['sheet']
        key=next(iter(values.values())); matches=[r for r in range(5,ws.max_row+1) if ws.cell(r,2).value==key]
        assert len(matches)<=1,(item['sheet'],key)
        if matches:
            r=matches[0]
            for field in HEADERS[item['sheet']].split():
                cell=ws.cell(r,heads[field]); desired=values.get(field)
                if cell.value!=desired:
                    old=legacy.get((item['sheet'],key),{}).get(field)
                    assert cell.value==old and (item['sheet']=='ConversationPage' and field=='seq'
                        or item['sheet']=='ConversationLine' and key==EVENT+'.p027' and field=='pageId'),(item['sheet'],key,field)
                    if migrations is not None: migrations.append({'sheet':item['sheet'],'cell':cell.coordinate,'before':old,'after':desired})
                    cell.value=desired
        else:
            r=ws.max_row+1
            for col in range(1,ws.max_column+1): ws.cell(r,col)._style=copy(ws.cell(r-1,col)._style)
            for field,value in values.items(): ws.cell(r,heads[field],value)
            added[item['sheet']]+=1
        locations.append({'sheet':item['sheet'],'key':key,'row':r})
    for sheet in ('EventDate','Result','Ending','ConditionKey','Condition','Action'):
        assert not any(EVENT in str(c.value) for row in wb[sheet].iter_rows(min_row=5) for c in row), sheet
    # Identity Text is the only Text change and must already match source-map import.
    t=text_wb['Text']; matches=[r for r in range(5,t.max_row+1) if str(t.cell(r,2).value).startswith(EVENT+'.')]
    assert len(matches)==57
    name_matches=[r for r in range(5,t.max_row+1) if t.cell(r,2).value=='church.thibaut.name']
    assert len(name_matches)<=1
    if name_matches:
        assert t.cell(name_matches[0],3).value=='客人'
    else:
        r=t.max_row+1
        for col in range(1,t.max_column+1): t.cell(r,col)._style=copy(t.cell(r-1,col)._style)
        t.cell(r,2,'church.thibaut.name'); t.cell(r,3,'客人')
        t.cell(r,4,'D1014 source guest alias; no hidden identity or portrait substitute')
        added['Text']=1
    return added, locations

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--apply',action='store_true'); args=ap.parse_args()
    initial={p:digest(p) for p in (NARRATIVE,TEXT,MAP)}
    rows, found=plan()
    books={NARRATIVE:openpyxl.load_workbook(NARRATIVE),TEXT:openpyxl.load_workbook(TEXT)}
    try:
        before={p:{(s.title,c.coordinate):signature(c) for s in w for row in s for c in row} for p,w in books.items()}
        migrations=[]
        added,locations=prepare(books[NARRATIVE],rows,books[TEXT],migrations)
        if not args.apply or not (any(added.values()) or migrations):
            print(json.dumps({'writes':False,'added':added,'migrations':migrations,'locations':locations},ensure_ascii=False)); return
        out=ROOT/'Library/Locus/tmp/church-content'/(datetime.now().strftime('%Y%m%d-%H%M%S-%f')+'-1014-config'); out.mkdir(parents=True)
        write_paths=[NARRATIVE] + ([TEXT] if added.get('Text',0) else [])
        for p in write_paths:
            assert digest(p)==initial[p], f'Concurrent change before backup: {p}'
            shutil.copy2(p,out/(p.stem+'.before.xlsx'))
            cand=out/(p.stem+'.candidate.xlsx'); books[p].save(cand)
            saved=openpyxl.load_workbook(cand)
            try:
                changed={(m['sheet'],m['cell']):m for m in migrations} if p==NARRATIVE else {}
                for (s,c),v in before[p].items():
                    actual=signature(saved[s][c])
                    if (s,c) in changed:
                        m=changed[(s,c)]
                        assert v[0]==m['before'] and actual[0]==m['after'] and actual[1:]==v[1:],(p,s,c)
                    else: assert actual==v,(p,s,c)
                if p==NARRATIVE:
                    check_added,_=prepare(saved,rows,books[TEXT])
                    assert not check_added,'Candidate not idempotent'
            finally: saved.close()
        manifest={'time':datetime.now().astimezone().isoformat(),'event':EVENT,'eventProperty':PROPERTY,'added':dict(added),'locations':locations,'migrations':migrations,'pages':12,'lines':51,'choices':6,'source_paragraphs':75,'source_sha256':SOURCES['1014.s2'][2],'blocked_preview_only':True,'preview_bg':'oct07.cafe.day','speaker_object':'church.thibaut','speaker_name':'客人','portrait_bound':False,'omitted_sheets':['EventDate','Result','ConditionKey','Ending','Condition','Action'],'pending':{'P028/envelope':'Closed prop-only LetterReveal after P027, no Reading data; LetterHide after service choice before P041. Existing prefab closed envelope art, non-modal/non-interactive; no LetterRead/name/signature','P005,P013,P031/animations':'door/service/drinking/cup placement animations missing; P031.a then cup action then P031.b; only dialogue configured','P070-P072/departure':'departure animation and closed-door guest silhouette composite missing; cafe background retained as preview placeholder','P007/portrait':'church.thibaut has no bound portrait; guest uses side=none, no substitute; bind formal identity resource before staging'},'generation_run':False,'editor_operations':False,'before_sha256':{str(p):h for p,h in initial.items()},'preserved_existing_cells':{str(p):len(v) for p,v in before.items()}}
        (out/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
        for p in initial: assert digest(p)==initial[p], f'Concurrent change: {p}'
        for p in write_paths: (out/(p.stem+'.candidate.xlsx')).replace(p)
        manifest['replaced']=True; manifest['after_sha256']={str(p):digest(p) for p in books}; manifest['text_unchanged']=digest(TEXT)==initial[TEXT]
        (out/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
        print(json.dumps({'added':dict(added),'manifest':str(out/'manifest.json'),'after_sha256':manifest['after_sha256']},ensure_ascii=False))
    finally:
        for w in books.values(): w.close()
if __name__=='__main__': main()
