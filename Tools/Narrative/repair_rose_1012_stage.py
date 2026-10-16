"""Bounded B1012 stage repair. Default is a read-only JSON plan; no generation.

Only 43 explicit ConversationLine cells may change: 27 speaker sides and 16
stage slots. All 71 source identities, emotes, aliases and text keys are audited.
No Text, background, global names, guard, root or Result is authored.
--apply requires --expected-sha256 from a fresh plan and the shared rose lock.
--self-test uses memory only (including XLSX serialization); run with python -B.

Reviewed anchors: P004.a/P016/P014 rebuild independent preparation/retry/post-
brew entries; P032 is the first alienor appearance and shifts focus from adele;
P050 clears alienor after she goes inside; P052 returns adele with the tray;
P060 brings her back on the long route; P064 clears outside adele on entering;
P085 returns her at this existing paragraph's entrance cue. P032 removal is a
layout decision, not an invented exit sentence. P085 has a late entrance inside
one unsplit source paragraph, so its slot applies at paragraph start. Parent
review accepted these source-granularity choices; no text is expanded/split.
P019 uses the same character portrait only as handoff B-R010 permits: CG and
paper-flower/civilian-outfit resources remain unverified placeholders.
"""
from __future__ import annotations

import sys
sys.dont_write_bytecode = True

import argparse
from collections import Counter
from datetime import datetime
from io import BytesIO
import json
from pathlib import Path
import re
import shutil

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

import import_rose_text as source
import repair_rose_stage as shared
from rose_write_guard import source_write_guard

ROOT = source.ROOT
NARRATIVE = shared.NARRATIVE
SCHEMA = shared.SCHEMA
EVENT = "rose.1012.s1"
ALIAS = EVENT + ".alias.alienor.name"
HEADERS = shared.HEADERS
SLOTS = shared.SLOTS
require = shared.require
EVIDENCE = {
    source.MAP: "94c70d2e272b20fbd4f83d08a5aaa50292ea203bb22bd3b2c7d2807006360c72",
    source.MAP.with_name("spec.md"): "12d2f98cc4869d4f6e10649d36ca72dc3463a8a52529af2a4035c711ec72925d",
    source.MAP.with_name("handoff.md"): "96db3d28eaf64c9f30dd8904bd8929388e159daab0368720b05ff26305788b9c",
}
# Exact observed source fixture, independently checked against every original
# paragraph and generated line. Tokens without /speaker are narration.
GROUPS = (
    ("preparation", "p004.a/adele/兴奋 p004.b/adele/兴奋 p005/player p006 p007 p008 p009.a/adele p009.b/adele p010/player p011/adele"),
    ("brew_pass", "p014/adele/兴奋"),
    ("brew_retry", "p016/adele/不屑"),
    ("sampling_adele", "p019/adele/兴奋 p020 p021 p022 p023/adele/兴奋 p024"),
    ("sampling_player", "p026/player p027 p028 p029/adele p030 p031 p032/alienor p033 p034"),
    ("hat_memory", "p036"),
    ("visit_brief", "p039/player p040/alienor p041/player p042 p043/alienor p044/player p045/alienor p046 p047 p048/player p049/alienor p050 p051 p052 p053"),
    ("visit_referred", "p056/alienor p057/player p058/alienor p059/player p060 p061/adele p062"),
    ("inside", "p064/player p065/alienor p066/player p067.a/alienor p067.b/alienor p067.c/alienor p068/player p069 p070"),
    ("hat_helped", "p072/alienor p073/player p074/alienor/尴尬 p075"),
    ("hat_not_helped", "p077/player p078/alienor/尴尬 p079/player p080/alienor p081/player p082 p083/alienor"),
    ("full_end", "p085"),
)
ANCHORS = (
    ("p004.a", ("adele", "-", "-"), "preparation_entry", "P003-P004; spec 4.4; adele left follows B1011"),
    ("p014", ("adele", "-", "-"), "post_brew_entry", "P013-P014; independent post_brew fixture"),
    ("p016", ("adele", "-", "-"), "retry_entry", "P015-P016; independent brew_retry fixture"),
    ("p032", ("-", "alienor", None), "visitor_focus", "P029-P034/P060; first alienor, clear adele after her P029 speech"),
    ("p050", (None, "-", None), "visitor_enters_shop", "P050: she and her attendants enter; no attendant portrait"),
    ("p052", ("adele", None, None), "short_tray_handoff", "P052: hand tray back to adele"),
    ("p060", ("adele", None, None), "adele_returns", "P060: adele rushes over and takes tray"),
    ("p064", ("-", None, None), "inside_clear_adele", "P061-P064: player takes visitor inside; adele stays serving"),
    ("p085", ("adele", None, None), "adele_inside_entrance", "P085: adele comes in late within this unsplit paragraph"),
)


def baseline():
    rows = {}
    for page, tokens in GROUPS:
        for seq, token in enumerate(tokens.split(), 1):
            parts = token.split("/")
            key = EVENT + "." + parts[0]
            speaker = parts[1] if len(parts) > 1 else None
            values = dict.fromkeys(HEADERS)
            values.update(lineId=key, pageId=EVENT + "." + page, seq=seq,
                          speakerKey=speaker, side="none", textKey=EVENT + ".text." + parts[0],
                          emote=parts[2] if len(parts) > 2 else None,
                          speakerNameKey=ALIAS if speaker == "alienor" else None)
            require(key not in rows, "Duplicate fixture")
            rows[key] = values
    require(len(rows) == 71, "B1012 fixture coverage changed")
    require(Counter(r["speakerKey"] for r in rows.values()) ==
            Counter({"adele": 11, "alienor": 16, "player": 16, None: 28}), "Identity coverage changed")
    return rows


def whitelist(old_rows):
    """Return exact (lineId, field): (old value, target, category, evidence)."""
    edits = {}

    def add(key, field, target, category, evidence):
        require(key in old_rows and field in ("side", *SLOTS), "Unknown whitelist cell")
        require((key, field) not in edits, "Overlapping whitelist")
        edits[key, field] = (old_rows[key][field], target, category, evidence)

    for key, row in old_rows.items():
        if row["speakerKey"] in ("adele", "alienor"):
            add(key, "side", "left" if row["speakerKey"] == "adele" else "right",
                "speaker_side", "source-map exact dialogue identity; spec 3.4/4.4; ordinary two-person layout")
    for short, slots, category, evidence in ANCHORS:
        for field, target in zip(SLOTS, slots):
            if target is not None:
                add(EVENT + "." + short, field, target, category, evidence)
    require(len(edits) == 43, "Whitelist coverage changed")
    return edits


def evidence_texts(old_rows):
    """Pinned mapped source only: no DOCX import and no filesystem writes."""
    for path, expected in EVIDENCE.items():
        require(source.digest(path) == expected, f"Reviewed evidence changed: {path}")
    texts = {}
    mapped = set()
    for number, kind, page, key, original in source.source_rows("1012"):
        if kind != "正文":
            continue
        raw = source.clean(original)
        parts = source.split_text("1012", number, raw) if number in source.MIXED["1012"] else [raw]
        # P004/P009 parts are all adele speech; P067 parts all alienor.
        speaker = ("adele" if original.startswith("阿黛尔：") else
                   "alienor" if original.startswith("埃莉诺：") else
                   "player" if original.startswith("{playername}") else None)
        emote = re.search(r'[（(]#([^）)]+)[）)]', original)
        for index, text in enumerate(parts):
            suffix = "." + chr(97 + index) if number in source.MIXED["1012"] else ""
            text_key = key + suffix
            line_key = text_key.replace(".text.", ".")
            require(line_key in old_rows and line_key not in mapped, "Unreviewed/duplicate mapped line")
            row = old_rows[line_key]
            require(row["pageId"] == page and row["speakerKey"] == speaker and
                    row["emote"] == (emote[1] if emote else None), f"Source identity/page drift: {line_key}")
            mapped.add(line_key)
            texts[text_key] = text
    require(mapped == old_rows.keys(), "B1012 source coverage changed")
    texts[ALIAS] = "埃莉诺"
    return texts


def prepare(wb):
    """Validate all 71 rows before assigning any of the 43 allowed values."""
    old_rows = baseline()
    allowed = whitelist(old_rows)
    ws = wb["ConversationLine"]
    require([c.value for c in ws[1]] == HEADERS, "ConversationLine schema drift")
    locations = {}
    for row in ws.iter_rows(min_row=5):
        key = row[1].value
        if not isinstance(key, str) or not key.startswith(EVENT + "."):
            continue
        require(key in old_rows, f"Unknown B1012 line: {key}")
        require(key not in locations, f"Duplicate B1012 line: {key}")
        locations[key] = row[0].row
    require(locations.keys() == old_rows.keys(), "Missing B1012 line(s)")
    for key, row in old_rows.items():
        for field, old in row.items():
            value = ws.cell(locations[key], HEADERS.index(field) + 1).value
            permitted = (old, allowed[key, field][1]) if (key, field) in allowed else (old,)
            require(any(type(value) is type(v) and value == v for v in permitted),
                    f"Unexpected value: {key}/{field}: {value!r}; expected {permitted!r}")
    targets, changes = [], []
    for (key, field), (_, target, category, evidence) in allowed.items():
        cell = ws.cell(locations[key], HEADERS.index(field) + 1)
        item = dict(sheet=ws.title, cell=cell.coordinate, lineId=key, field=field,
                    before=cell.value, after=target, category=category, source=evidence)
        targets.append(item)
        if cell.value != target:
            changes.append(item)
    for item in changes:
        ws[item["cell"]].value = item["after"]
    return targets, changes


def repair(apply=False, expected_sha256=None):
    require(not apply or expected_sha256 is not None, "--apply requires --expected-sha256 from a fresh plan")
    paths = tuple(dict.fromkeys((NARRATIVE, source.TEXT, SCHEMA, *EVIDENCE,
                                Path(__file__), Path(shared.__file__), Path(source.__file__),
                                Path(source_write_guard.__wrapped__.__code__.co_filename))))
    hashes = {p: source.digest(p) for p in paths}
    shared.validate_schema(SCHEMA.read_text(encoding="utf-8"))
    if expected_sha256 is not None:
        require(hashes[NARRATIVE] == expected_sha256.lower(), "Narrative SHA changed; re-plan")
    texts = evidence_texts(baseline())
    tw = openpyxl.load_workbook(source.TEXT, read_only=True)
    try:
        staged = {}
        for row in tw["Text"].iter_rows(min_row=5, values_only=True):
            if row[1] in texts:
                require(row[1] not in staged, "Duplicate B1012 Text key")
                staged[row[1]] = row[2]
        require(staged == texts, "B1012 source Text/alias mismatch")
    finally:
        tw.close()
    wb = openpyxl.load_workbook(NARRATIVE)
    try:
        before, old_metadata = shared.snapshot(wb), shared.metadata(wb)
        targets, changes = prepare(wb)
        expected = shared.snapshot(wb)
        report = dict(tool="repair_rose_1012_stage", writes=bool(apply and changes), changed=False,
                      formal_enabled=False, before_sha256=hashes[NARRATIVE], idempotent=not changes,
                      audited_rows=71, identities={"adele": 11, "alienor": 16, "player": 16, "narrator": 28},
                      changed_cells=len(changes), changed_rows=len({x["lineId"] for x in changes}),
                      changes_by_category=dict(Counter(x["category"] for x in changes)),
                      whitelist_by_category=dict(Counter(x["category"] for x in targets)),
                      cells_checked=len(before), targets=targets, changes=changes,
                      resource_status="Existing backgrounds unchanged; CG, decorations and civilian outfit unverified",
                      layout_decisions=["P032 clears adele after P029; visitor focus, no invented exit text",
                                        "P085 entrance applies at paragraph start; text describes late entrance"],
                      evidence_sha256={str(p): h for p, h in hashes.items() if p != NARRATIVE})

        def cas():
            shared.compare_hashes(hashes, {p: source.digest(p) for p in paths})

        cas()
        if not apply or not changes:
            return report
        out = ROOT / "Library/Locus/tmp/rose-content" / (
            datetime.now().strftime("%Y%m%d-%H%M%S-%f") + "-1012-stage-repair")
        out.mkdir(parents=True, exist_ok=False)
        backup, candidate = out / "Narrative.before.xlsx", out / "Narrative.candidate.xlsx"
        shutil.copy2(NARRATIVE, backup)
        require(source.digest(backup) == hashes[NARRATIVE], "Backup SHA mismatch; not applied")
        shared.save_preserving_properties(wb, candidate)
        saved = openpyxl.load_workbook(candidate)
        try:
            shared.verify_candidate(saved, before, old_metadata, expected, changes)
            require(not prepare(saved)[1], "Candidate is not idempotent")
        finally:
            saved.close()
        candidate_sha = source.digest(candidate)
        report.update(status="verified-candidate", evidence=str(out), metadata_preserved=True,
                      backup_sha256=source.digest(backup), candidate_sha256=candidate_sha)
        journal = out / "repair-evidence.json"
        journal.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        replacement = out / "Narrative.replace.xlsx"
        shutil.copy2(candidate, replacement)
        require(source.digest(replacement) == candidate_sha and source.digest(candidate) == candidate_sha
                and source.digest(backup) == hashes[NARRATIVE], "Evidence/replacement SHA mismatch")
        cas()
        replacement.replace(NARRATIVE)
        require(source.digest(NARRATIVE) == candidate_sha, "Post-replacement SHA mismatch")
        shared.compare_hashes({p: h for p, h in hashes.items() if p != NARRATIVE},
                              {p: source.digest(p) for p in paths if p != NARRATIVE})
        report.update(changed=True, status="committed", after_sha256=candidate_sha)
        journal.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        return report
    finally:
        wb.close()


def fixture():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ConversationLine"
    ws.append(HEADERS)
    for _ in range(3):
        ws.append([None] * len(HEADERS))
    for values in baseline().values():
        ws.append([values[h] for h in HEADERS])
    ws.freeze_panes = "C5"
    ws.column_dimensions["B"].width = 32
    for row in ws.iter_rows(min_row=5):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.fill = PatternFill("solid", fgColor="F0F0F0")
            cell.number_format = "@"
            cell.comment = Comment("retain metadata", "B1012 fixture")
    ws.cell(5, HEADERS.index("stageLeftKey") + 1)._style = None
    ws.cell(5, HEADERS.index("stageLeftKey") + 1).comment = None
    other = wb.create_sheet("Untouched")
    other["A1"] = "Text/guard/root/Result/background/global-name sentinel"
    other["A1"].hyperlink = "https://example.invalid/keep"
    other["B2"] = "=1+2"
    buffer = BytesIO()
    shared.save_preserving_properties(wb, buffer)
    wb.close()
    buffer.seek(0)
    return openpyxl.load_workbook(buffer)


def self_test():
    """Only memory; no path reads, writes, tempfile, lock or source access."""
    passed = []
    wb = fixture()
    before, meta = shared.snapshot(wb), shared.metadata(wb)
    targets, changes = prepare(wb)
    require(len(targets) == len(changes) == 43, "Exact baseline change count changed")
    expected = shared.snapshot(wb)
    require(not prepare(wb)[1] and shared.snapshot(wb) == expected, "Not idempotent")
    passed.extend(("exact-43-whitelist-cells", "idempotent-memory-no-op"))
    buffer = BytesIO()
    shared.save_preserving_properties(wb, buffer)
    buffer.seek(0)
    saved = openpyxl.load_workbook(buffer)
    shared.verify_candidate(saved, before, meta, expected, changes)
    require(not prepare(saved)[1], "Roundtrip not idempotent")
    passed.append("memory-xlsx-roundtrip-all-cells-and-metadata")
    ws = wb["ConversationLine"]
    loc = {r[1].value: r[0].row for r in ws.iter_rows(min_row=5)}

    def cell(key, field):
        return ws.cell(loc[key], HEADERS.index(field) + 1)

    def rejects(operation):
        try:
            operation()
        except RuntimeError:
            return
        raise RuntimeError("Self-test failed to reject invalid input")

    for key, field in whitelist(baseline()):
        c = cell(key, field)
        old = c.value
        c.value = "UNREVIEWED"
        unchanged = shared.snapshot(wb)
        rejects(lambda: prepare(wb))
        require(shared.snapshot(wb) == unchanged, "Rejected prepare mutated cells")
        c.value = old
    passed.append("reject-third-value-all-43-cells-without-partial-mutation")
    for short, field, bad in (("p004.a", "speakerKey", None), ("p004.b", "emote", None),
                              ("p009.a", "speakerKey", None), ("p009.b", "speakerKey", None),
                              ("p020", "speakerKey", "adele"), ("p060", "speakerKey", "adele"),
                              ("p085", "speakerKey", "adele"), ("p032", "speakerNameKey", "character.alienor.name"),
                              ("p005", "side", "left"), ("p006", "side", "left"),
                              ("p029", "stageLeftKey", "-"), ("p019", "stageRightKey", "alienor"),
                              ("p004.a", "seq", True), ("p004.a", "textKey", "changed"),
                              ("p004.a", "pageId", "changed")):
        c = cell(EVENT + "." + short, field)
        old = c.value
        c.value = bad
        rejects(lambda: prepare(wb))
        c.value = old
    passed.append("reject-narrator-speaker-alias-text-order-player-side-and-early-visitor-drift")
    c = cell(EVENT + ".p004.a", "lineId")
    for bad in (None, EVENT + ".p999", EVENT + ".p004.b"):
        c.value = bad
        rejects(lambda: prepare(wb))
    c.value = EVENT + ".p004.a"
    ws.cell(1, len(HEADERS) + 1).value = "unknownField"
    rejects(lambda: prepare(wb))
    del ws._cells[1, len(HEADERS) + 1]
    passed.append("reject-missing-unknown-duplicate-row-and-header-drift")
    # Simulate only the reviewed slot semantics, across all six entry/routes.
    pages = {page: [EVENT + "." + t.split("/")[0] for t in tokens.split()] for page, tokens in GROUPS}
    for stage, appreciation, hat in (("preparation", False, False), ("brew_retry", False, False),
                                      *(("post_brew", a, h) for a in (False, True) for h in (False, True))):
        route = [stage] if stage != "post_brew" else ["brew_pass", "sampling_adele", "sampling_player"]
        if stage == "post_brew":
            if hat:
                route.append("hat_memory")
            route += (["visit_referred", "inside", "hat_helped" if hat else "hat_not_helped", "full_end"]
                      if appreciation else ["visit_brief"])
        slots = ["STALE", "STALE", "STALE"]
        observed = {}
        for page in route:
            for key in pages[page]:
                for i, field in enumerate(SLOTS):
                    value = cell(key, field).value
                    if value is not None:
                        slots[i] = None if value == "-" else value
                observed[key.rsplit(EVENT + ".", 1)[-1]] = tuple(slots)
                speaker = cell(key, "speakerKey").value
                if speaker in ("adele", "alienor"):
                    require(slots[0 if speaker == "adele" else 1] == speaker, "Speaker absent from slot")
                if key in [EVENT + "." + s for s in ("p004.a", "p016", "p014")]:
                    require(slots == ["adele", None, None], "Independent entry retained stale slots")
        if stage == "post_brew":
            require(observed["p029"] == ("adele", None, None), "P029 adele missing/early visitor")
            require(observed["p032"] == (None, "alienor", None), "P032 focus wrong")
            if appreciation:
                require(observed["p060"] == ("adele", "alienor", None), "P060 return wrong")
                require(observed["p064"] == (None, "alienor", None), "Outside adele carried inside")
                require(observed["p085"] == ("adele", "alienor", None), "P085 entrance wrong")
            else:
                require(observed["p050"] == (None, None, None), "Visitor retained outside")
                require(observed["p052"] == ("adele", None, None), "Short handoff missing")
    passed.append("six-memory-stage-routes-and-all-speaking-slots")
    saved["Untouched"]["A1"].value = "tampered"
    rejects(lambda: shared.verify_candidate(saved, before, meta, expected, changes))
    saved["Untouched"]["A1"].value = "Text/guard/root/Result/background/global-name sentinel"
    target = changes[0]
    saved[target["sheet"]][target["cell"]].comment = Comment("retain metadata", "wrong author")
    rejects(lambda: shared.verify_candidate(saved, before, meta, expected, changes))
    passed.append("reject-unlisted-candidate-value-and-target-metadata-tampering")
    shared.compare_hashes({"input": "a"}, {"input": "a"})
    rejects(lambda: shared.compare_hashes({"input": "a"}, {"input": "b"}))
    passed.append("CAS-reject-concurrent-input")
    c = cell(EVENT + ".p004.a", "side")
    c.value = "none"
    require(len(prepare(wb)[1]) == 1 and not prepare(wb)[1], "Partial approved state failed")
    passed.append("partial-approved-state-idempotent")
    wb.close()
    saved.close()
    return dict(self_test="PASS", files_written=0, tests=passed, tests_passed=len(passed),
                audited_rows=71, old_state_changed_cells=len(changes),
                whitelist_by_category=dict(Counter(x["category"] for x in targets)))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--apply", action="store_true")
    mode.add_argument("--self-test", action="store_true")
    parser.add_argument("--expected-sha256", help="Narrative SHA from the read-only plan")
    args = parser.parse_args()
    if args.self_test:
        require(args.expected_sha256 is None, "--self-test does not read sources or SHA")
        report = self_test()
    elif args.apply:
        require(args.expected_sha256 is not None and re.fullmatch(r"[0-9a-fA-F]{64}", args.expected_sha256),
                "--apply requires a 64-digit --expected-sha256")
        with source_write_guard("repair rose 1012 stage whitelist"):
            report = repair(True, args.expected_sha256)
    else:
        report = repair(False, args.expected_sha256)
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
