"""Configure A05 reporter SOURCE tables only; never generate or call Unity.

Default is a read-only plan (including static route checks). After parent approval:
  python -B Tools/Narrative/configure_neutral_reporter.py > PLAN.json
  python -B Tools/Narrative/configure_neutral_reporter.py --apply --expected-sha256 PLAN.json
The plan's three workbook SHAs and immutable source evidence are required by apply.
No A01-A04 configure module is imported or called. Existing body/choice/stage/emote
cells and GameParam are read-only. Changing the single fee GameParam requires a
fresh plan/apply to synchronize the choice guard and payment before generation.
"""
from __future__ import annotations

import argparse
from collections import Counter
from contextlib import contextmanager
from copy import copy, deepcopy
from datetime import datetime
import hashlib
import json
import os
from pathlib import Path
import shutil
import uuid
from zipfile import ZIP_DEFLATED, ZipFile
import xml.etree.ElementTree as ET

import openpyxl
from openpyxl.comments import Comment
from openpyxl.writer.excel import ExcelWriter
from openpyxl.xml.functions import tostring

# Immutable source verification only; never stage/import Text or configure other roots.
from import_neutral_text import source_rows, verify_docx, clean

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / 'Assets/ConfigTables/Data'
NARRATIVE = DATA / 'Narrative/Narrative.xlsx'
TEXT = DATA / 'Shared/Text.xlsx'
PARAM = DATA / 'Simulation/GameParam.xlsx'
BOOKS = (NARRATIVE, TEXT, PARAM)
MAP = ROOT / 'Docs/superpowers/specs/2026-10-16-content-batches/A-neutral/source-map.md'
SCHEMA = ROOT / 'Assets/ConfigTables/Defines/schema.xml'
LOCK = ROOT / 'Library/Locus/tmp/narrative-source-write.lock'
A05 = 'neutral.1018.s1'
VISIT = A05 + '.reporter'
SCHEDULE = A05 + '.schedule'
NOTICE_EVENTS = (A05 + '.notice.recommendation', A05 + '.notice.advertisement')
OWN_EVENTS = (A05, SCHEDULE, *NOTICE_EVENTS)
FEE = 'neutral.reporter.advertisement_fee'
PREVIOUS = 'story.neutral.1013.s2.completed'
COMPLETED = 'story.' + A05 + '.completed'
RECOMMEND = 'story.' + A05 + '.recommendation_requested'
ADVERT = 'story.' + A05 + '.advertisement_requested'
PRINTED = 'story.' + A05 + '.printed'
YES = A05 + '.choice.q_ad.yes'
NO = A05 + '.choice.q_ad.no'
PAGES = ('arrival wait_service serve taste_pass taste_fail identity recommend '
         'advert_offer advert_yes advert_no closing').split()
OLD_PROPERTY = 'entry=Immediate;blocksSave=true;blockedBy=A-B002,A-B007,A-B008,A-B009,A-CAP02,A-CAP05'
OLD_GUARDS = {
    'taste_pass': 'dependency:A05.quality.pass',
    'recommend': 'dependency:A05.quality.pass',
    'taste_fail': 'dependency:A05.quality.fail',
    'advert_offer': 'dependency:A05.quality.fail',
    'advert_yes': 'choice:' + YES + '|dependency:A05.quality.fail|dependency:A05.advertisement.fee-and-balance',
    'advert_no': 'choice:' + NO + '|dependency:A05.quality.fail',
}
GUARDS = {
    'taste_pass': 'service:passed', 'recommend': 'service:passed',
    'taste_fail': 'not:service:passed', 'advert_offer': 'not:service:passed',
    'advert_yes': 'choice:' + YES + '|not:service:passed',
    'advert_no': 'choice:' + NO + '|not:service:passed',
}
FLAG_LABELS = {
    COMPLETED: '记者到访已完成', RECOMMEND: '咖啡馆推荐待刊',
    ADVERT: '咖啡馆广告待刊', PRINTED: '咖啡馆报刊通知已刊出',
}
NOTICE_TEXT = {
    A05 + '.notice.title': '《圣诺特之声》',
    A05 + '.notice.recommendation.body': '本期侧栏刊登了沸点咖啡馆的推荐文章。',
    A05 + '.notice.advertisement.body': '本期商业版刊登了沸点咖啡馆的广告。',
}


def require(ok, message):
    if not ok:
        raise RuntimeError(message)


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def heads(ws):
    return {cell.value: cell.column for cell in ws[1] if cell.value}


def records(ws):
    h = heads(ws)
    return [dict(_row=r, **{k: ws.cell(r, c).value for k, c in h.items() if k != '##var'})
            for r in range(5, ws.max_row + 1) if ws.cell(r, 2).value is not None]


def locate(ws, key, identity=None):
    h = heads(ws)
    matches = [r for r in range(5, ws.max_row + 1) if ws.cell(r, 2).value == key
               and all(ws.cell(r, h[k]).value == v for k, v in (identity or {}).items())]
    require(len(matches) <= 1, f'Duplicate source row: {ws.title}/{key}/{identity}')
    return matches[0] if matches else None


def one(ws, key):
    row = locate(ws, key)
    require(row is not None, f'Missing source row: {ws.title}/{key}')
    return next(r for r in records(ws) if r['_row'] == row)


def fee_value(book):
    value = one(book[book.sheetnames[0]], FEE)['value']
    require(isinstance(value, str) and value.isascii() and value.isdecimal()
            and 0 < int(value) <= 2147483647, 'Fee GameParam must be a positive int32 string')
    return str(int(value))


def signature(cell):
    typ = 'empty' if cell.value is None and cell.data_type in ('inlineStr', 'n') else cell.data_type
    return (cell.value, typ, tuple(cell._style) if cell.has_style else None,
            (cell.comment.text, cell.comment.author, cell.comment.width, cell.comment.height) if cell.comment else None,
            tuple(sorted(dict(cell.hyperlink).items())) if cell.hyperlink else None)


def snapshot(wb):
    return {(ws.title, c.coordinate): signature(c) for ws in wb for row in ws for c in row}


def signature_sha256(cells):
    return hashlib.sha256(json.dumps([(s, c, v) for (s, c), v in sorted(cells.items())],
        ensure_ascii=False, default=str, separators=(',', ':')).encode('utf-8')).hexdigest()


def metadata(wb):
    def dimensions(items):
        return [(k, tuple(sorted(dict(v).items())), tuple(v._style) if v.has_style else None)
                for k, v in items.items()]
    return deepcopy((wb.properties, tostring(wb.custom_doc_props.to_tree()), wb.defined_names,
        wb.calculation, wb.security, wb.views,
        [(ws.title, ws.sheet_state, ws.sheet_properties, ws.sheet_format, ws.views,
          ws.freeze_panes, str(ws.merged_cells), dimensions(ws.row_dimensions),
          dimensions(ws.column_dimensions), ws.auto_filter, ws.data_validations,
          ws.print_options, ws.page_margins, ws.page_setup, ws.protection) for ws in wb]))


def load_book(path):
    wb = openpyxl.load_workbook(path)
    with ZipFile(path) as archive:
        if 'docProps/core.xml' not in archive.namelist():
            wb.properties.created = wb.properties.modified = datetime(2000, 1, 1)
    return wb


def preserve_core_part(source, candidate):
    """Preserve authored core XML (including absence) and empty t=str cell metadata."""
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with ZipFile(source) as archive:
        core = archive.read('docProps/core.xml') if 'docProps/core.xml' in archive.namelist() else None
        empty = {}
        for name in archive.namelist():
            if name.startswith('xl/worksheets/sheet') and name.endswith('.xml'):
                cells = [c for c in ET.fromstring(archive.read(name)).findall('.//s:sheetData/s:row/s:c', ns)
                         if c.get('t') == 'str' and not list(c)]
                if cells:
                    empty[name] = cells
    repack = candidate.with_suffix('.core-preserved.xlsx')
    with ZipFile(candidate) as old, ZipFile(repack, 'w', ZIP_DEFLATED, allowZip64=True) as new:
        for info in old.infolist():
            data = old.read(info.filename)
            if info.filename == 'docProps/core.xml':
                if core is None:
                    continue
                data = core
            elif info.filename in empty:
                tree = ET.fromstring(data)
                rows = {r.get('r'): r for r in tree.findall('.//s:sheetData/s:row', ns)}
                for original in empty[info.filename]:
                    coord = original.get('r')
                    row = rows[''.join(c for c in coord if c.isdigit())]
                    current = next((c for c in row if c.get('r') == coord), None)
                    if current is not None:
                        require(not list(current), f'Empty-string preservation met a value: {coord}')
                        row.remove(current)
                    row.insert(0, original)
                data = ET.tostring(tree, encoding='utf-8')
            elif core is None and info.filename in ('[Content_Types].xml', '_rels/.rels'):
                tree = ET.fromstring(data)
                for child in list(tree):
                    if child.get('PartName') == '/docProps/core.xml' or child.get('Type', '').endswith('/metadata/core-properties'):
                        tree.remove(child)
                data = ET.tostring(tree, encoding='utf-8')
            new.writestr(info, data)
    repack.replace(candidate)


@contextmanager
def write_guard():
    LOCK.parent.mkdir(parents=True, exist_ok=True)
    owner = json.dumps(dict(owner='A05-neutral-reporter', pid=os.getpid(), token=uuid.uuid4().hex))
    try:
        with LOCK.open('x', encoding='utf-8') as stream:
            stream.write(owner)
    except FileExistsError:
        raise RuntimeError(f'Source-write lock occupied; do not remove another owner: {LOCK}') from None
    try:
        yield
    finally:
        require(LOCK.exists() and LOCK.read_text(encoding='utf-8') == owner, 'Lock ownership changed; preserving lock')
        LOCK.unlink()


class Configuration:
    def __init__(self, books):
        self.books = books
        self.n, self.t, self.p = (books[p] for p in BOOKS)
        self.edits, self.added = [], []
        self.fee = fee_value(self.p)

    def path(self, book):
        return next(p for p, w in self.books.items() if w is book)

    def append(self, book, sheet, values, identity=None):
        ws = book[sheet]
        require(set(values) <= heads(ws).keys(), f'Schema mismatch {sheet}: {set(values) - heads(ws).keys()}')
        key = next(iter(values.values()))
        row = locate(ws, key, identity)
        if row is None:
            row = ws.max_row + 1
            for col in range(1, ws.max_column + 1):
                ws.cell(row, col)._style = copy(ws.cell(row - 1, col)._style)
            for field, value in values.items():
                ws.cell(row, heads(ws)[field], value)
            ws.cell(row, 2).comment = Comment(json.dumps(dict(batch='A05-reporter',
                authorization='Parent-authorized A05 formal source tables'), ensure_ascii=False), 'A05-neutral-reporter')
            self.added.append(dict(book=self.path(book).name, sheet=sheet, row=row, key=key, values=values))
        else:
            for field, value in values.items():
                require(ws.cell(row, heads(ws)[field]).value == value, f'Existing different row: {sheet}/{key}/{field}')
        return row

    def change(self, book, sheet, key, field, value, allowed):
        ws = book[sheet]
        row = locate(ws, key)
        require(row is not None and field in heads(ws), f'Missing required cell {sheet}/{key}/{field}')
        cell = ws.cell(row, heads(ws)[field])
        require(cell.value in (*allowed, value), f'Unexpected source value {sheet}/{key}/{field}: {cell.value!r}')
        if cell.value != value:
            self.edits.append(dict(book=self.path(book).name, sheet=sheet, row=row, cell=cell.coordinate,
                                   key=key, field=field, before=cell.value, after=value))
            cell.value = value

    def condition(self, key, owner_type, owner, kind, target, value):
        self.append(self.n, 'Condition', dict(conditionId=key, ownerType=owner_type, ownerId=owner,
            groupIndex=1, condType=kind, target=target, value=value))

    def result(self, event, seq, suffix, kind, target, value, guards=()):
        key = event + '.result.' + suffix
        self.append(self.n, 'Result', dict(resultId=key, eventId=event, seq=seq, resultType=kind,
            target=target, value=value, note='A05 reporter: atomic root settlement; no economy multiplier/reward'))
        for index, guard in enumerate(guards, 1):
            self.condition(key + '.guard.' + str(index), 'result', key, 'PRESENTATION', guard, 'true')

    def configure(self):
        # Reuse the existing schema; never create headers, columns or alter schema.xml.
        require({'orderProductId', 'requireRealProduction'} <= heads(self.n['CafeVisit']).keys(),
                'CafeVisit product/real-production columns must already exist')
        self.change(self.n, 'Event', A05, 'eventProperty', 'entry=Immediate;blocksSave=true', (OLD_PROPERTY,))
        self.change(self.n, 'Object', 'neutral.reporter', 'portraitIndoorPath',
                    'Characters/neutral_reporter/placeholder', (None,))
        for key, label in FLAG_LABELS.items():
            self.append(self.t, 'Text', dict(key='condition.' + key, text=label, comment='A05 reporter state label'))
            self.append(self.n, 'ConditionKey', dict(key=key, nameText='condition.' + key,
                                                    initialValue=False, onlySetTrue=True))
        # Upstream belongs to A01-A04; consume, never declare/reconfigure it here.
        one(self.n['ConditionKey'], PREVIOUS)
        self.condition(A05 + '.day', 'event', A05, 'BAR_BETWEEN', 'world.calendar', '13,13')
        self.condition(A05 + '.uncompleted', 'event', A05, 'BOOL_EQ', COMPLETED, 'false')
        self.append(self.n, 'Event', dict(eventId=SCHEDULE, eventKind='world', actorId='world',
            onceOnly=True, logTag='A05', eventProperty='entry=MorningBeforeMap'))
        self.append(self.n, 'EventDate', dict(eventId=SCHEDULE, day=13, timeSlot='morning', repeatRule='once', priority=100))
        self.condition(SCHEDULE + '.day', 'event', SCHEDULE, 'BAR_BETWEEN', 'world.calendar', '13,13')
        self.condition(SCHEDULE + '.previous', 'event', SCHEDULE, 'BOOL_EQ', PREVIOUS, 'true')
        self.condition(SCHEDULE + '.uncompleted', 'event', SCHEDULE, 'BOOL_EQ', COMPLETED, 'false')
        self.result(SCHEDULE, 1, 'reserve', 'RESERVE_CAFE_VISIT', 'obj:neutral.reporter', 'visit:' + VISIT)
        self.result(SCHEDULE, 2, 'spawn', 'SPAWN_NPC', 'obj:neutral.reporter', 'scene:cafe')
        self.append(self.n, 'CafeVisit', dict(visitKey=VISIT, characterId='neutral.reporter',
            arrivalDelaySeconds='150', patienceSeconds='180', leaveAfterServed='false',
            orderProductId='espresso', requireRealProduction=True))
        self.append(self.n, 'CafeVisitHook', dict(visitKey=VISIT, stage='arrived', trigger='auto', seq=1, eventId=A05),
            identity=dict(stage='arrived', trigger='auto', seq=1))
        self.append(self.n, 'ConversationPage', dict(pageId=A05 + '.wait_service',
            presentationKey=A05 + '.conversation', seq=2, stepType='WaitService'))
        old_pages = [p for p in PAGES if p != 'wait_service']
        for index, suffix in enumerate(PAGES, 1):
            if suffix == 'wait_service':
                continue
            self.change(self.n, 'ConversationPage', A05 + '.' + suffix, 'seq', index, (old_pages.index(suffix) + 1,))
            if suffix in GUARDS:
                self.change(self.n, 'ConversationPage', A05 + '.' + suffix, 'pageCondKeys', GUARDS[suffix], (OLD_GUARDS[suffix],))
        # Fee is not a default owned by this tool. Synchronize existing managed consumers
        # only when their old values agree; never overwrite the editable GameParam.
        guard_key, pay_key = YES + '.balance', A05 + '.result.payment'
        old_guard = locate(self.n['Condition'], guard_key)
        old_pay = locate(self.n['Result'], pay_key)
        require((old_guard is None) == (old_pay is None), 'Partial fee consumers; inspect before repair')
        if old_guard is not None:
            previous_fee = one(self.n['Condition'], guard_key)['value']
            require(isinstance(previous_fee, str) and previous_fee.isascii() and previous_fee.isdecimal()
                    and 0 < int(previous_fee) <= 2147483647, 'Invalid previous fee')
            require(one(self.n['Result'], pay_key)['value'] == previous_fee, 'Old fee consumers disagree')
            self.change(self.n, 'Condition', guard_key, 'value', self.fee, (previous_fee,))
            self.change(self.n, 'Result', pay_key, 'value', self.fee, (previous_fee,))
        self.condition(guard_key, 'choice', YES, 'BAR_GTE', 'cafe.money', self.fee)
        self.result(A05, 1, 'recommendation', 'SET_FLAG', RECOMMEND, 'true', ('service:passed',))
        paid_guards = ('choice:' + YES, 'not:service:passed')
        self.result(A05, 2, 'advertisement', 'SET_FLAG', ADVERT, 'true', paid_guards)
        self.result(A05, 3, 'payment', 'MONEY_REMOVE', 'cafe', self.fee, paid_guards)
        self.result(A05, 4, 'despawn', 'DESPAWN_NPC', 'obj:neutral.reporter', None)
        self.result(A05, 5, 'completed', 'SET_FLAG', COMPLETED, 'true')
        for key, text in NOTICE_TEXT.items():
            self.append(self.t, 'Text', dict(key=key, text=text, comment='A05 authorized minimal next-issue Notice; no reward'))
        for event, request, opposite in zip(NOTICE_EVENTS, (RECOMMEND, ADVERT), (ADVERT, RECOMMEND)):
            presentation = event + '.presentation'
            self.append(self.n, 'Event', dict(eventId=event, eventKind='world', actorId='world', onceOnly=True,
                logTag='A05', presentationClass='Notice', presentationKey=presentation,
                presentationLogTag='A05', eventProperty='entry=MorningBeforeMap;blocksSave=true'))
            self.append(self.n, 'PresentationIndex', dict(presentationKey=presentation, presentationClass='Notice', logTag='A05'))
            self.append(self.n, 'Notice', dict(presentationKey=presentation, titleKey=A05 + '.notice.title',
                textKey=event + '.body'))
            self.append(self.n, 'EventDate', dict(eventId=event, day=14, timeSlot='morning', repeatRule='once', priority=100))
            self.condition(event + '.day', 'event', event, 'BAR_BETWEEN', 'world.calendar', '14,14')
            self.condition(event + '.requested', 'event', event, 'BOOL_EQ', request, 'true')
            self.condition(event + '.opposite', 'event', event, 'BOOL_EQ', opposite, 'false')
            self.condition(event + '.unprinted', 'event', event, 'BOOL_EQ', PRINTED, 'false')
            self.result(event, 1, 'printed', 'SET_FLAG', PRINTED, 'true')
        return self


def validate(books):
    """Static source contract and route/transaction model; not a Unity/runtime claim."""
    n, t, p = (books[x] for x in BOOKS)
    fee = fee_value(p)
    texts = {r['key']: r['text'] for r in records(t['Text'])}
    rows = source_rows('1018.s1')
    for _, kind, key, original in rows:
        require(texts.get(key) == clean(original, kind), f'Original A05 Text differs: {key}')
        if kind == '正文':
            require(one(n['ConversationLine'], key.replace('.text.', '.line.'))['textKey'] == key,
                    f'Original body line binding differs: {key}')
    require([r['textKey'] for r in records(n['Choice']) if r['groupKey'] == A05 + '.q_ad'] ==
            [A05 + '.text.p069', A05 + '.text.p078'], 'A05 original choices differ')
    require(one(n['Object'], 'neutral.reporter')['portraitIndoorPath'] == 'Characters/neutral_reporter/placeholder', 'Indoor portrait differs')
    for key in FLAG_LABELS:
        row = one(n['ConditionKey'], key)
        require(row['initialValue'] is False and row['onlySetTrue'] is True and row['nameText'] in texts, f'Flag definition {key}')
    require(one(n['Event'], A05)['eventProperty'] == 'entry=Immediate;blocksSave=true', 'Root entry differs')
    for event in OWN_EVENTS:
        require(one(n['Event'], event)['onceOnly'] is True, f'Event must be once: {event}')
    schedule = one(n['Event'], SCHEDULE)
    require(schedule['eventProperty'] == 'entry=MorningBeforeMap' and
            not any(schedule[k] for k in ('titleText', 'contentKey', 'setupType', 'setupKey', 'presentationClass', 'presentationKey')), 'Schedule must be result-only')
    dates = [r for r in records(n['EventDate']) if r['eventId'] in OWN_EVENTS]
    require([(r['eventId'], r['day'], r['timeSlot'], r['repeatRule'], r['priority']) for r in dates] ==
            [(SCHEDULE, 13, 'morning', 'once', 100), *[(e, 14, 'morning', 'once', 100) for e in NOTICE_EVENTS]], 'Schedule/Notice dates differ')
    visit = one(n['CafeVisit'], VISIT)
    require(tuple(visit[k] for k in ('characterId', 'arrivalDelaySeconds', 'patienceSeconds', 'leaveAfterServed', 'orderProductId', 'requireRealProduction')) ==
            ('neutral.reporter', '150', '180', 'false', 'espresso', True), 'Real visit contract differs')
    hooks = [r for r in records(n['CafeVisitHook']) if r['visitKey'] == VISIT or r['eventId'] == A05]
    require(len(hooks) == 1 and tuple(hooks[0][k] for k in ('visitKey', 'stage', 'trigger', 'seq', 'eventId')) ==
            (VISIT, 'arrived', 'auto', 1, A05), 'Exact arrived/auto hook differs')
    pages = sorted((r for r in records(n['ConversationPage']) if r['presentationKey'] == A05 + '.conversation'), key=lambda r: r['seq'])
    require([r['pageId'] for r in pages] == [A05 + '.' + s for s in PAGES] and [r['seq'] for r in pages] == list(range(1, 12)), 'Wait must be arrival -> wait_service -> serve')
    for page in pages:
        suffix = page['pageId'].removeprefix(A05 + '.')
        require(page['pageCondKeys'] == GUARDS.get(suffix), f'Unexpected page guard {suffix}')
    wait = pages[1]
    require(wait['stepType'] == 'WaitService' and not any(wait[k] for k in ('pageCondKeys', 'pageBarConds', 'optionGroupKey', 'readingKey', 'checkId', 'bgOverride', 'bgmOverride')), 'Wait must be empty/unconditional')
    require(not any(r['pageId'] == wait['pageId'] for r in records(n['ConversationLine'])), 'Wait contains dialogue')
    require(one(n['ConversationLine'], A05 + '.line.p011')['pageId'] == A05 + '.arrival' and
            all(one(n['ConversationLine'], A05 + '.line.p' + str(i).zfill(3))['pageId'] == A05 + '.serve' for i in (14, 15, 16)), 'P013/P014 source seam differs')
    conditions = records(n['Condition'])
    results = sorted((r for r in records(n['Result']) if r['eventId'] in OWN_EVENTS), key=lambda r: (r['eventId'], r['seq']))
    def guards(owner_type, owner):
        return [r for r in conditions if r['ownerType'] == owner_type and r['ownerId'] == owner]
    def guard_contract(owner_type, owner, expected):
        actual = guards(owner_type, owner)
        require([(r['groupIndex'], r['condType'], r['target'], r['value']) for r in actual] ==
                [(1, kind, target, value) for kind, target, value in expected], f'Exact guard contract differs: {owner}')
    guard_contract('event', A05, [('BAR_BETWEEN', 'world.calendar', '13,13'), ('BOOL_EQ', COMPLETED, 'false')])
    guard_contract('event', SCHEDULE, [('BAR_BETWEEN', 'world.calendar', '13,13'), ('BOOL_EQ', PREVIOUS, 'true'), ('BOOL_EQ', COMPLETED, 'false')])
    guard_contract('choice', YES, [('BAR_GTE', 'cafe.money', fee)])
    guard_contract('choice', NO, [])
    expected_results = {
        SCHEDULE: [('RESERVE_CAFE_VISIT', 'obj:neutral.reporter', 'visit:' + VISIT), ('SPAWN_NPC', 'obj:neutral.reporter', 'scene:cafe')],
        A05: [('SET_FLAG', RECOMMEND, 'true'), ('SET_FLAG', ADVERT, 'true'), ('MONEY_REMOVE', 'cafe', fee),
              ('DESPAWN_NPC', 'obj:neutral.reporter', None), ('SET_FLAG', COMPLETED, 'true')],
        **{e: [('SET_FLAG', PRINTED, 'true')] for e in NOTICE_EVENTS},
    }
    root_guards = [('service:passed',), ('choice:' + YES, 'not:service:passed'), ('choice:' + YES, 'not:service:passed'), (), ()]
    for event, expected in expected_results.items():
        group = [r for r in results if r['eventId'] == event]
        require([(r['resultType'], r['target'], r['value']) for r in group] == expected and
                [r['seq'] for r in group] == list(range(1, len(group) + 1)), f'Result contract differs: {event}')
        for index, result in enumerate(group):
            expected_guards = root_guards[index] if event == A05 else ()
            guard_contract('result', result['resultId'], [('PRESENTATION', g, 'true') for g in expected_guards])
    for event, request, opposite in zip(NOTICE_EVENTS, (RECOMMEND, ADVERT), (ADVERT, RECOMMEND)):
        guard_contract('event', event, [('BAR_BETWEEN', 'world.calendar', '14,14'), ('BOOL_EQ', request, 'true'),
                                     ('BOOL_EQ', opposite, 'false'), ('BOOL_EQ', PRINTED, 'false')])
        ev = one(n['Event'], event)
        require(ev['presentationClass'] == 'Notice' and ev['eventProperty'] == 'entry=MorningBeforeMap;blocksSave=true', 'Notice entry differs')
        require(one(n['PresentationIndex'], ev['presentationKey'])['presentationClass'] == 'Notice', 'Notice index differs')
        notice = one(n['Notice'], ev['presentationKey'])
        require(texts[notice['titleKey']] == NOTICE_TEXT[A05 + '.notice.title'] and
                texts[notice['textKey']] == NOTICE_TEXT[event + '.body'], 'Minimal Notice text differs')
        require(not any(notice[k] for k in ('timeKey', 'imageKey', 'barConds', 'noticeBarConds')), 'Notice adds unauthorized payload')
    # Traverse ACTUAL configured page and result guards. Payment is never filtered by
    # balance: a balance drop after choosing yes must fail the entire preflight.
    traces, routes = [], set()
    def read(key, passed, choice):
        if key.startswith('not:'):
            return not read(key[4:], passed, choice)
        if key.startswith('choice:'):
            return key[7:] == choice
        require(key == 'service:passed', f'Unsupported model guard {key}')
        return passed
    def event_eligible(event, day, flags):
        for condition in guards('event', event):
            if condition['condType'] == 'BAR_BETWEEN':
                low, high = map(int, condition['value'].split(','))
                if not low <= day <= high:
                    return False
            else:
                require(condition['condType'] == 'BOOL_EQ', 'Unsupported model event guard')
                if (condition['target'] in flags) != (condition['value'] == 'true'):
                    return False
        return True
    require(event_eligible(SCHEDULE, 13, {PREVIOUS}) and
            not event_eligible(SCHEDULE, 13, set()) and
            not event_eligible(SCHEDULE, 13, {PREVIOUS, COMPLETED}) and
            not event_eligible(SCHEDULE, 12, {PREVIOUS}), 'Schedule prerequisites/day differ')
    for passed in (True, False):
        for choice in (None,) if passed else (YES, NO):
            shown = [r['pageId'].removeprefix(A05 + '.') for r in pages if all(read(g, passed, choice) for g in (r['pageCondKeys'] or '').split('|') if g)]
            require(shown[:3] == ['arrival', 'wait_service', 'serve'] and shown[-1] == 'closing', 'Source seam/closing route differs')
            require(('taste_pass' in shown) == passed and ('taste_fail' in shown) != passed, 'Taste branches not exclusive')
            require(('recommend' in shown) == passed and ('advert_offer' in shown) != passed, 'Offer branches not exclusive')
            require(('advert_yes' in shown) == (choice == YES) and ('advert_no' in shown) == (choice == NO), 'Choice branches differ')
            routes.add(tuple(shown))
            applied = [r for r in results if r['eventId'] == A05 and all(read(g['target'], passed, choice) for g in guards('result', r['resultId']))]
            fee_due = sum(int(r['value']) for r in applied if r['resultType'] == 'MONEY_REMOVE')
            require(fee_due == (int(fee) if choice == YES else 0), 'Payment route differs')
            for balance in sorted({0, int(fee) - 1, int(fee), int(fee) + 1}):
                # No partial effects survive a failed aggregate preflight.
                committed = applied if balance >= fee_due else []
                flags = {r['target'] for r in committed if r['resultType'] == 'SET_FLAG'}
                require((COMPLETED in flags) == (balance >= fee_due), 'Failed preflight marked completed')
                require((RECOMMEND in flags) == (passed and balance >= fee_due), 'Recommendation mismatch')
                require((ADVERT in flags) == (choice == YES and balance >= fee_due), 'Ad request mismatch')
                eligible = [e for e in NOTICE_EVENTS if event_eligible(e, 14, flags)]
                require(len(eligible) == (1 if (RECOMMEND in flags or ADVERT in flags) else 0), 'Notice mutual exclusion differs')
                require(not any(event_eligible(e, 13, flags) or event_eligible(e, 15, flags) for e in NOTICE_EVENTS), 'Notice eligible on wrong day')
                flags.add(PRINTED)
                require(not any(event_eligible(e, 14, flags) for e in NOTICE_EVENTS), 'Notice can print twice')
                traces.append(dict(passed=passed, choice=choice, balance=balance, fee_due=fee_due,
                                   preflight_success=balance >= fee_due, pages=shown, notices=eligible))
    require(len(routes) == 3, 'Exactly recommendation/paid-ad/refusal routes required')
    return dict(kind='Static source contract/model only; no Unity/runtime claim', source_display_rows=len(rows),
        routes=3, transaction_scenarios=len(traces), fee=fee, single_source_fee=FEE,
        seam='arrival(P013) -> WaitService -> serve(P014-P016)',
        payment_has_no_balance_filter=True, insufficient_preflight_commits_nothing=True,
        notice_mutually_exclusive_and_once=True,
        trace_sha256=hashlib.sha256(json.dumps(traces, sort_keys=True).encode()).hexdigest())


def run(apply=False, expected_plan=None):
    initial = {p.name: digest(p) for p in BOOKS}
    require(not apply or expected_plan is not None, '--apply requires --expected-sha256 PLAN.json')
    if expected_plan is not None:
        require(expected_plan.get('tool') == 'configure_neutral_reporter', 'Expected an A05 reporter plan')
        require(expected_plan['before_sha256'] == initial, 'Source SHA differs from plan; re-plan all three books')
    source = verify_docx('1018.s1', source_rows('1018.s1'))
    evidence = {p: digest(p) for p in (MAP, source, SCHEMA)}
    evidence_hashes = {str(p): sha for p, sha in evidence.items()}
    if expected_plan is not None:
        require(expected_plan['source_evidence_sha256'] == evidence_hashes, 'Source/schema evidence differs from plan')
    books = {p: load_book(p) for p in BOOKS}
    try:
        before = {p: snapshot(w) for p, w in books.items()}
        meta = {p: metadata(w) for p, w in books.items()}
        sizes = {p: {s.title: (s.max_row, s.max_column) for s in w} for p, w in books.items()}
        config = Configuration(books).configure()
        validation = validate(books)
        second = Configuration(books).configure()
        require(not second.edits and not second.added, 'In-memory idempotence failed')
        allowed = {p: {(e['sheet'], e['cell']) for e in config.edits if e['book'] == p.name} for p in BOOKS}
        preserved = {p: {k: v for k, v in before[p].items() if k not in allowed[p]} for p in BOOKS}
        require(not allowed[PARAM] and not any(e['book'] == PARAM.name for e in config.added), 'GameParam is read-only')
        report = dict(tool='configure_neutral_reporter', writes=False, changed=False,
            idempotent=not config.edits and not config.added, before_sha256=initial, after_sha256=initial.copy(),
            source_evidence_sha256=evidence_hashes, fee=config.fee, single_source_fee=FEE,
            added_counts=dict(Counter(e['book'] + '/' + e['sheet'] for e in config.added)),
            added_rows=config.added, edits=config.edits,
            preserved_cells={p.name: dict(count=len(preserved[p]), signature_sha256=signature_sha256(preserved[p])) for p in BOOKS},
            validation=validation, generation_run=False, unity_called=False,
            protected=['All existing body/choice/stage/emote cells and comments unchanged.',
                       'GameParam (including editable fee and unused economy multipliers) is read-only.',
                       'Object portraitPath unchanged; only portraitIndoorPath is configured.',
                       'A01-A04 and unrelated batches unchanged; exact CafeVisit context enforced by service code.'])
        def cas(hashes):
            for p in BOOKS:
                require(digest(p) == hashes[p.name], f'Concurrent workbook change; CAS aborted: {p}')
            for p, sha in evidence.items():
                require(digest(p) == sha, f'Immutable evidence changed; CAS aborted: {p}')
        cas(initial)
        if not apply or report['idempotent']:
            return report
        out = ROOT / 'Library/Locus/tmp/neutral-content' / (datetime.now().strftime('%Y%m%d-%H%M%S-%f') + '-A05-reporter')
        out.mkdir(parents=True, exist_ok=False)
        changed_books = [p for p in BOOKS if any(e['book'] == p.name for e in config.edits + config.added)]
        candidates, candidate_hashes, backups = {}, {}, {}
        for p in BOOKS:
            backup = out / (p.stem + '.before.xlsx')
            shutil.copy2(p, backup)
            require(digest(backup) == initial[p.name], f'Backup SHA mismatch {p}')
            backups[p] = backup
            if p not in changed_books:
                continue
            candidate = out / (p.stem + '.candidate.xlsx')
            with ZipFile(candidate, 'w', ZIP_DEFLATED, allowZip64=True) as archive:
                ExcelWriter(books[p], archive).save()
            preserve_core_part(backup, candidate)
            candidates[p] = candidate
            candidate_hashes[p.name] = digest(candidate)
        reloaded = {p: load_book(candidates.get(p, backups[p])) for p in BOOKS}
        try:
            for p, saved in reloaded.items():
                require(saved.sheetnames == books[p].sheetnames and metadata(saved) == meta[p], f'Workbook/sheet metadata changed: {p}')
                actual = snapshot(saved)
                require(actual == snapshot(books[p]), f'Candidate readback differs from intended: {p}')
                for key, old in before[p].items():
                    require(key in actual, f'Existing coordinate lost: {p}/{key}')
                    require(actual[key][2:] == old[2:] if key in allowed[p] else actual[key] == old,
                            f'Non-target value/metadata changed: {p}/{key}')
                for sheet, (rows, cols) in sizes[p].items():
                    added = [e['row'] for e in config.added if e['book'] == p.name and e['sheet'] == sheet]
                    require(saved[sheet].max_column == cols and saved[sheet].max_row == rows + len(added), f'Schema/dimensions changed: {p}/{sheet}')
                    require(added == list(range(rows + 1, rows + len(added) + 1)), 'Non-contiguous append')
                require(signature_sha256({k: actual[k] for k in preserved[p]}) == signature_sha256(preserved[p]), f'Preserved-cell SHA differs: {p}')
            require(validate(reloaded) == validation, 'Candidate static contract differs')
            again = Configuration(reloaded).configure()
            require(not again.edits and not again.added, 'Candidate readback not idempotent')
        finally:
            for wb in reloaded.values():
                wb.close()
        report.update(evidence=str(out), backup_sha256={p.name: digest(b) for p, b in backups.items()},
            candidate_sha256=candidate_hashes, metadata_preserved=True, candidate_readback=True,
            candidate_idempotent=True, status='verified-candidates')
        journal = out / 'reporter-evidence.json'
        journal.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        replacements = {}
        for p, candidate in candidates.items():
            replacements[p] = out / (p.stem + '.replace.xlsx')
            shutil.copy2(candidate, replacements[p])
            require(digest(replacements[p]) == candidate_hashes[p.name], 'Replacement SHA mismatch')
        current, committed = initial.copy(), []
        try:
            for p in changed_books:
                cas(current)
                require(digest(candidates[p]) == candidate_hashes[p.name] and digest(backups[p]) == initial[p.name], 'Evidence SHA changed')
                replacements[p].replace(p)
                committed.append(p)
                current[p.name] = candidate_hashes[p.name]
                require(digest(p) == current[p.name], f'Post-CAS SHA mismatch {p}')
            cas(current)
        except BaseException:
            for p in reversed(committed):
                require(digest(p) == candidate_hashes[p.name], f'Cannot safely roll back intervening writer: {p}; inspect {journal}')
                rollback = out / (p.stem + '.rollback.xlsx')
                shutil.copy2(backups[p], rollback)
                require(digest(rollback) == initial[p.name], 'Rollback backup SHA mismatch')
                rollback.replace(p)
            report.update(status='rolled-back', writes=False)
            journal.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
            raise
        report.update(writes=True, changed=True, after_sha256=current, status='committed')
        journal.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        return report
    finally:
        for wb in books.values():
            wb.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--apply', action='store_true')
    parser.add_argument('--expected-sha256', type=Path, help='JSON reporter plan; compares all three books and source evidence')
    args = parser.parse_args()
    expected = json.loads(args.expected_sha256.read_text(encoding='utf-8-sig')) if args.expected_sha256 else None
    if args.apply:
        with write_guard():
            report = run(True, expected)
    else:
        report = run(False, expected)
    print(json.dumps(report, ensure_ascii=True, indent=2))


if __name__ == '__main__':
    main()
