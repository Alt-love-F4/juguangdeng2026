"""Prepare result-free, blocked B1012 source rows. Default: read-only JSON plan.

Parent applies serially with --apply; generation is a separate parent task.
--fixture --stage preparation|brew_retry|post_brew is an ISOLATED plan only,
not a Unity launcher. Post-brew requires explicit A/H and D=true snapshots.
No stage crosses P012, failure retry, or claims a real production completion.
Every page is dependency gated; no formal ConditionKey/Result/Action/Date.
"""
from __future__ import annotations

import argparse
from collections import Counter
from copy import copy
from datetime import datetime
import json
import os
from pathlib import Path
import re
import shutil
import uuid
import xml.etree.ElementTree as ET

import openpyxl
from openpyxl.comments import Comment
import import_rose_text as source

ROOT = source.ROOT
NARRATIVE = ROOT / 'Assets/ConfigTables/Data/Narrative/Narrative.xlsx'
TEXT = source.TEXT
SCHEMA = ROOT / 'Assets/ConfigTables/Defines/schema.xml'
EVENT = 'rose.1012.s1'
PRESENTATION = EVENT + '.conversation'
ALIAS = EVENT + '.alias.alienor.name'
STAGES = ('preparation', 'brew_retry', 'post_brew')
HEADERS = {
    'Event': 'eventId eventKind actorId targetId titleText contentKey setupType setupKey onceOnly logTag presentationClass presentationKey presentationLogTag eventProperty',
    'SetupParams': 'setupKey tplName slotValues skipAble autoAdvance checkId',
    'PresentationIndex': 'presentationKey presentationClass logTag',
    'ConversationPage': 'pageId presentationKey seq pageCondKeys pageBarConds optionGroupKey bgOverride bgmOverride stepType readingKey checkId',
    'ConversationLine': 'lineId pageId seq speakerKey side textKey voiceKey emote sfxKey accentFxKey accentTargetKey speakerNameKey',
    'Text': 'key text comment',
}
UNCONNECTED = {
    'B-C02/B-B05': 'P012: real request/session owner, independent black>=1 and milk>=1 cups, actual product grades and approved ordinary mapping absent. No fake pass, discard or inventory changes.',
    'B-C03': 'Day7 first real business-scene ready boundary and valid idle owner absent. No BeginDay EventDate, AP change or automatic business end.',
    'B-history': 'A=story.rose.1011.s1.sylvain_appreciation; H=story.church.1009.s2.hat_helped. Both explicit immutable snapshots; D completed must be proven by producer, not inferred from H=false or date.',
    'B-fixture': 'Dedicated B1012 isolated runtime adapter is not installed by this tool. Dependency injection must require preview context, exactly one stage, explicit A/H/D, reject unknown keys and stop at stage end. PreviewStartPageId alone is insufficient.',
    'B-resources': 'Paper flowers, sampling CG, civilian outfit and exact stage reconstruction unverified; no guessed visual keys. P054/P086 terminal black fade remains unconnected.',
    'B-results': 'No completion/student-known result or formal entry registered. Formal activation refused until producer evidence, business bridge, ratings, cancellation/retry and terminal fade are integrated and verified.',
}


def require(ok, message):
    if not ok:
        raise RuntimeError(message)


def stage_of(page):
    short = page.rsplit('.', 1)[-1]
    return short if short in ('preparation', 'brew_retry') else 'post_brew'


def guards(page):
    short = page.rsplit('.', 1)[-1]
    result = ['dependency:B-C02.stage.' + stage_of(page)]
    if stage_of(page) == 'post_brew':
        result += ['dependency:B-history.D-completed', 'dependency:B-history.snapshot-ready']
        if short == 'hat_memory':
            result.append('dependency:B-history.H')
        if short == 'visit_brief':
            result.append('not:dependency:B-history.A')
        if short in ('visit_referred', 'inside', 'hat_helped', 'hat_not_helped', 'full_end'):
            result.append('dependency:B-history.A')
        if short in ('hat_helped', 'hat_not_helped'):
            result.append(('not:' if short == 'hat_not_helped' else '') + 'dependency:B-history.H')
    return '|'.join(result)


def build_plan():
    entries = source.expected_rows('1012')  # Full original hash/XML/source-map verification.
    originals = {r[0]: r for r in source.source_rows('1012')}
    rows = []

    def add(sheet, **values):
        rows.append({'sheet': sheet, 'values': values})

    add('Event', eventId=EVENT, eventKind='actor', actorId='player', setupKey=EVENT,
        onceOnly=True, presentationClass='Conversation', presentationKey=PRESENTATION,
        eventProperty='entry=Immediate;blocksSave=true;blockedBy=B-C02+B-C03+B-B05+B-history+B-fixture+B-resources;resultFree=true')
    add('SetupParams', setupKey=EVENT, tplName='Tpl_ArrivalConversation', slotValues='opts=', skipAble=False, autoAdvance=False)
    add('PresentationIndex', presentationKey=PRESENTATION, presentationClass='Conversation', logTag='B1012 blocked construction')
    add('Text', key=ALIAS, text='埃莉诺', comment='B1012 explicit display alias; alienor identity reused')
    pages = list(dict.fromkeys(e['page'] for e in entries))
    for seq, page in enumerate(pages, 1):
        add('ConversationPage', pageId=page, presentationKey=PRESENTATION, seq=seq,
            pageCondKeys=guards(page), stepType='Dialogue')
    seqs = Counter()
    for entry in entries:
        number = int(entry['source'].split('P')[-1])
        raw = originals[number][4]
        speaker = ('adele' if raw.startswith('阿黛尔：') else 'alienor' if raw.startswith('埃莉诺：')
                   else 'player' if raw.startswith('{playername}') else None)
        emote = re.search(r'[（(]#([^）)]+)[）)]', raw)
        seqs[entry['page']] += 1
        add('Text', key=entry['key'], text=entry['text'], comment=entry['source'] + '; blocked B1012 source mapping')
        add('ConversationLine', lineId=entry['key'].replace('.text.', '.'), pageId=entry['page'],
            seq=seqs[entry['page']], speakerKey=speaker, side='none', textKey=entry['key'],
            emote=emote[1] if emote else None, speakerNameKey=ALIAS if speaker == 'alienor' else None)
    require(len(entries) == 71 and len(pages) == 12, 'B1012 source coverage changed')
    return rows, {'event': EVENT, 'blocked': True, 'formal_enabled': False, 'preview_enabled': False,
                  'counts': dict(Counter(r['sheet'] for r in rows)), 'unsupported': UNCONNECTED,
                  'source_map_sha256': source.digest(source.MAP), 'docx_sha256': source.SOURCES['1012'][2],
                  'source_paragraphs': [dict(number=n, kind=r[1], page=r[2], original=r[4]) for n, r in originals.items()],
                  'stage_boundaries': {'preparation': 'P004-P011 -> STOP before P012 real production',
                                       'brew_retry': 'P016 only -> STOP before P012 retry; no automatic retry',
                                       'post_brew': 'Explicit isolated snapshot only: P014 then selected history route; no real production claim; terminal fade unconnected'}}


def signature(cell):
    return (cell.value, cell.data_type, tuple(cell._style) if cell._style else (0,) * 9,
            (cell.comment.text, cell.comment.author) if cell.comment else None, copy(cell.hyperlink))


def validate_schema():
    schema = ET.parse(SCHEMA).getroot()
    for sheet, headers in HEADERS.items():
        if sheet == 'Text':
            continue
        bean = schema.find(".//bean[@name='Ncs" + sheet + "Entry']")
        require(bean is not None, 'Missing schema bean: ' + sheet)
        require(set(headers.split()).issubset({v.get('name') for v in bean.findall('var')}), 'Schema fields missing: ' + sheet)
    alias = schema.find(".//bean[@name='NcsConversationLineEntry']/var[@name='speakerNameKey']")
    require(alias is not None and alias.get('type') == 'string#ref=TbText?', 'speakerNameKey contract unavailable')


def prepare(books, rows):
    validate_schema()
    indexes, locations, added = {}, [], Counter()
    for path, wb in books.items():
        # Check unique-ID sheets consumed by this plan; Bar uses composite IDs.
        for ws in wb:
            if ws.title not in set(HEADERS) | {'Object'}:
                continue
            index = {}
            for row in ws.iter_rows(min_row=5):
                key = row[1].value if len(row) > 1 else None
                if key is not None:
                    require(key not in index, f'Duplicate source key: {ws.title}/{key}')
                    index[key] = row[1].row
            indexes[(path, ws.title)] = index
    for identity in ('player', 'adele', 'alienor'):
        index = indexes[(NARRATIVE, 'Object')]
        require(identity in index, 'Required shared Object missing: ' + identity)
        ws = books[NARRATIVE]['Object']
        heads = {c.value: c.column for c in ws[1]}
        require(ws.cell(index[identity], heads['objType']).value == ('player' if identity == 'player' else 'npc'), 'Object identity conflict: ' + identity)
    for sheet in ('Result', 'EventDate', 'Action', 'ConditionKey', 'CafeVisitHook'):
        if sheet in books[NARRATIVE].sheetnames:
            for row in books[NARRATIVE][sheet].iter_rows(min_row=5):
                require(not any(isinstance(c.value, str) and EVENT in c.value for c in row), 'Blocked root has external activation/result/fact: ' + sheet)
    expected = {}
    for item in rows:
        sheet, values = item['sheet'], item['values']
        path = TEXT if sheet == 'Text' else NARRATIVE
        ws = books[path][sheet]
        headers = [c.value for c in ws[1]]
        require(headers[:2] == ['##var', HEADERS[sheet].split()[0]] and len(headers) == len(set(headers))
                and set(HEADERS[sheet].split()).issubset(headers), 'Workbook header drift: ' + sheet)
        index = indexes[(path, sheet)]
        key = values[headers[1]]
        require((sheet, key) not in expected, 'Duplicate planned key')
        expected[(sheet, key)] = values
        if key in index:
            number = index[key]
            fields = ['key', 'text'] if sheet == 'Text' else headers[1:]
            for field in fields:
                require(ws.cell(number, headers.index(field) + 1).value == values.get(field), f'Existing row differs: {sheet}/{key}/{field}')
            operation = 'reuse'
        else:
            number = ws.max_row + 1
            for col in range(1, ws.max_column + 1):
                ws.cell(number, col)._style = copy(ws.cell(number - 1, col)._style)
            for field, value in values.items():
                ws.cell(number, headers.index(field) + 1, value)
            ws.cell(number, 2).comment = Comment('B1012 source-map; result-free blocked construction', 'B-rose')
            index[key] = number
            added[sheet] += 1
            operation = 'append'
        locations.append(dict(sheet=sheet, key=key, row=number, operation=operation, values=values))
    for sheet in HEADERS:
        path = TEXT if sheet == 'Text' else NARRATIVE
        for key in indexes[(path, sheet)]:
            if str(key).startswith(EVENT) and (sheet, key) not in expected:
                raise RuntimeError(f'Unexpected existing root row: {sheet}/{key}')
    for row in rows:
        if row['sheet'] == 'ConversationLine' and row['values'].get('speakerKey') == 'alienor':
            require(row['values']['speakerNameKey'] == ALIAS, 'Alienor alias missing')
    return dict(added), locations


def fixture_plan(rows, stage, appreciation, hat_helped, church_completed):
    if stage == 'post_brew':
        require(all(v is not None for v in (appreciation, hat_helped, church_completed)), 'Post-brew requires explicit A/H/D snapshots')
        require(church_completed, 'Missing D completed source: refuse post-brew even if H=false')
    deps = {'dependency:B-C02.stage.' + s: s == stage for s in STAGES}
    deps.update({'dependency:B-history.A': appreciation, 'dependency:B-history.H': hat_helped,
                 'dependency:B-history.D-completed': church_completed,
                 'dependency:B-history.snapshot-ready': stage == 'post_brew'})
    selected = []
    for row in rows:
        if row['sheet'] != 'ConversationPage':
            continue
        visible = True
        for guard in row['values']['pageCondKeys'].split('|'):
            negate = guard.startswith('not:')
            key = guard[4:] if negate else guard
            require(key in deps, 'Unknown fixture dependency: ' + key)
            value = deps[key]
            # Never interpret an unknown fact as a negative snapshot.
            visible = visible and value is not None and (not value if negate else value)
        if visible:
            selected.append(row['values']['pageId'])
    lines = [r['values']['lineId'] for r in rows if r['sheet'] == 'ConversationLine' and r['values']['pageId'] in selected]
    return dict(isolated_preview_only=True, runtime_adapter_connected=False, stage=stage,
                dependencies=deps, pages=selected, lines=lines, writes=False, real_production_completed=False)


def apply_plan():
    base = ROOT / 'Library/Locus/tmp'
    base.mkdir(parents=True, exist_ok=True)
    lock = base / 'narrative-source-write.lock'
    token = 'B1012:' + str(os.getpid()) + ':' + uuid.uuid4().hex
    # Occupied lock fails before entering cleanup; never remove another writer's lock.
    with lock.open('x', encoding='utf-8') as handle:
        handle.write(token)
        handle.flush()
        os.fsync(handle.fileno())
    books = {}
    try:
        paths = (TEXT, NARRATIVE)
        watched = (*paths, SCHEMA, source.MAP, Path(source.verify_docx('1012', source.source_rows('1012'))), Path(source.__file__), Path(__file__))
        hashes = {p: source.digest(p) for p in watched}
        rows, manifest = build_plan()
        books = {p: openpyxl.load_workbook(p) for p in paths}
        before = {p: {(s.title, c.coordinate): signature(c) for s in wb for row in s for c in row} for p, wb in books.items()}
        added, locations = prepare(books, rows)
        if not added:
            print(json.dumps({'idempotent': True, 'changed': False}))
            return
        out = base / 'rose-content' / (datetime.now().strftime('%Y%m%d-%H%M%S-%f') + '-1012-config')
        out.mkdir(parents=True)
        candidates = {}
        for p, wb in books.items():
            require(source.digest(p) == hashes[p], 'Concurrent source change before backup: ' + str(p))
            backup = out / (p.stem + '.before.xlsx')
            shutil.copy2(p, backup)
            require(source.digest(backup) == hashes[p], 'Backup race: ' + str(p))
            candidate = out / (p.stem + '.candidate.xlsx')
            wb.save(candidate)
            saved = openpyxl.load_workbook(candidate)
            try:
                for (sheet, coordinate), old in before[p].items():
                    require(signature(saved[sheet][coordinate]) == old, f'Existing cell changed: {sheet}!{coordinate}')
                for item in locations:
                    if (item['sheet'] == 'Text') != (p == TEXT):
                        continue
                    for cell in wb[item['sheet']][item['row']]:
                        require(signature(saved[item['sheet']][cell.coordinate]) == signature(cell), 'Candidate row mismatch')
            finally:
                saved.close()
            candidates[p] = candidate
        manifest.update(added=added, locations=locations, before_sha256={str(p): h for p, h in hashes.items()}, replaced=[])
        journal = out / 'manifest.json'

        def record():
            journal.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')

        record()
        require(all(source.digest(p) == h for p, h in hashes.items()), 'Concurrent input change; candidates retained: ' + str(out))
        # Cooperative CAS, not an OS transaction across two files. Journal partial success;
        # never roll back over another writer. Exact rerun can finish a partial append.
        for p in paths:
            require(all(source.digest(q) == h for q, h in hashes.items() if str(q) not in manifest['replaced']), 'Concurrent input change; see ' + str(journal))
            if not any(n for s, n in added.items() if (s == 'Text') == (p == TEXT)):
                continue
            candidates[p].replace(p)
            manifest['replaced'].append(str(p))
            manifest.setdefault('after_sha256', {})[str(p)] = source.digest(p)
            record()
        print(json.dumps({'added': added, 'evidence': str(out), 'formal_enabled': False}, ensure_ascii=False))
    finally:
        for wb in books.values():
            wb.close()
        if lock.exists() and lock.read_text(encoding='utf-8') == token:
            lock.unlink()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--apply', action='store_true')
    parser.add_argument('--fixture', action='store_true', help='Isolated preview row plan only; cannot apply')
    parser.add_argument('--stage', choices=STAGES)
    for name in ('appreciation', 'hat-helped', 'church-completed'):
        parser.add_argument('--' + name, choices=('true', 'false'))
    args = parser.parse_args()
    fixture_args = (args.stage, args.appreciation, args.hat_helped, args.church_completed)
    require(not args.fixture or (args.stage and not args.apply), 'Fixture requires one explicit stage and forbids --apply')
    require(args.fixture or not any(x is not None for x in fixture_args), 'Snapshot/stage arguments require --fixture')
    if args.apply:
        apply_plan()
        return
    rows, manifest = build_plan()
    books = {p: openpyxl.load_workbook(p) for p in (TEXT, NARRATIVE)}
    try:
        added, locations = prepare(books, rows)
        manifest.update(writes=False, would_append=added, row_plan=locations)
        if args.fixture:
            values = [None if v is None else v == 'true' for v in (args.appreciation, args.hat_helped, args.church_completed)]
            manifest['fixture'] = fixture_plan(rows, args.stage, *values)
        print(json.dumps(manifest, ensure_ascii=False, indent=2))
    finally:
        for wb in books.values():
            wb.close()


if __name__ == '__main__':
    main()
