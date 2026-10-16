"""Prepare only the written P004-P015 fragment. Never crosses brewing P016.
Default: validate a memory-only plan. --apply: append blocked, result-free rows.
This does not implement the unfinished chapter, floor unlock or brewing bridge.
"""
from collections import Counter
from copy import copy
from datetime import datetime
import argparse
import json
import shutil
import openpyxl
from openpyxl.comments import Comment
import import_rose_text as source
from rose_write_guard import source_write_guard

ROOT = source.ROOT
NARRATIVE = ROOT / 'Assets/ConfigTables/Data/Narrative/Narrative.xlsx'
EVENT = 'rose.1016.s1'
PRESENTATION = EVENT + '.conversation'


def plan():
    entries = source.expected_rows('1016')
    assert len(entries) == 9
    rows = []
    def add(sheet, **values): rows.append((sheet, values))
    add('Text', key=EVENT+'.title', text='两位访客', comment='B source fragment, not a formal chapter')
    add('Text', key=EVENT+'.alias.alienor', text='埃莉诺', comment='B P009/P014 alias; no royal identity disclosure')
    add('Object', objId='sade', objType='npc', nameText='character.sade.name', portraitPath='Characters/sade/neutral')
    # alienor already belongs to the shared identity directory. Never change its
    # global name; every spoken line below supplies the local alias explicitly.
    add('Event', eventId=EVENT, eventKind='actor', actorId='player', titleText=EVENT+'.title',
        setupType='Tpl_ArrivalConversation', setupKey=EVENT, onceOnly=True,
        presentationClass='Conversation', presentationKey=PRESENTATION,
        eventProperty='entry=Immediate;blocksSave=true;blockedBy=B-B06+B-B07+B-C02+B-C05;contentBoundary=P015')
    add('SetupParams', setupKey=EVENT, tplName='Tpl_ArrivalConversation', slotValues='opts=', skipAble=False, autoAdvance=False)
    add('PresentationIndex', presentationKey=PRESENTATION, presentationClass='Conversation', logTag='B1016 unfinished: P015 construction boundary')
    for seq, page in enumerate(['arrival', 'wisdom', 'order'], 1):
        values = dict(pageId=EVENT+'.'+page, presentationKey=PRESENTATION, seq=seq)
        if page == 'arrival': values['bgOverride'] = 'oct07.cafe.day'
        if page == 'wisdom':
            values.update(stepType='PassiveCheck', checkId=EVENT+'.check.wisdom')
        add('ConversationPage', **values)
    add('AbilityCheck', checkId=EVENT+'.check.wisdom', mode='Inline', comment='P010 wisdom >6; construction fragment only')
    add('AbilityCheckCondition', conditionId=EVENT+'.check.wisdom.1', checkId=EVENT+'.check.wisdom', attributeId='wisdom', threshold=7, groupIndex=1, seq=1)
    counters = Counter()
    for entry in entries:
        number = int(entry['source'].split('P')[-1])
        assert 4 <= number <= 15
        counters[entry['page']] += 1
        values = dict(lineId=entry['key'].replace('.text.', '.'), pageId=entry['page'], seq=counters[entry['page']], side='none', textKey=entry['key'])
        raw = entry['original'].strip()
        if raw.startswith('{playername}'): values['speakerKey'] = 'player'
        elif raw.startswith('埃莉诺：'):
            values.update(speakerKey='alienor', side='left', speakerNameKey=EVENT+'.alias.alienor')
        elif raw.startswith('萨德：'): values.update(speakerKey='sade', side='right')
        add('ConversationLine', **values)
    return entries, rows


def prepare(books, entries, rows):
    text = books[source.TEXT]['Text']
    texts = {r[1].value:r[2].value for r in text.iter_rows(min_row=5) if r[1].value}
    for entry in entries:
        assert texts.get(entry['key']) == entry['text'], 'Source Text mismatch: '+entry['key']
    assert 'character.sade.name' in texts
    n = books[NARRATIVE]
    assert any(r[1].value == 'alienor' for r in n['Object'].iter_rows(min_row=5)), 'Shared alienor Object not ready'
    for sheet, field in [('Result','eventId'), ('EventDate','eventId'), ('Action','eventId')]:
        if sheet not in n.sheetnames: continue
        ws = n[sheet]; h = {c.value:c.column for c in ws[1]}
        if field in h:
            assert not any(ws.cell(i,h[field]).value == EVENT for i in range(5,ws.max_row+1)), 'Unfinished root has formal '+sheet
    added = Counter()
    for sheet, values in rows:
        book = books[source.TEXT] if sheet == 'Text' else n
        ws = book[sheet]; h = {c.value:c.column for c in ws[1]}
        assert set(values) <= set(h), f'Unsupported header: {sheet}/{set(values)-set(h)}'
        keyfield = ws.cell(1,2).value; key = values[keyfield]
        matches = [r for r in ws.iter_rows(min_row=5) if r[1].value == key]
        assert len(matches) <= 1, f'Duplicate {sheet}/{key}'
        if matches:
            for field,value in values.items():
                if field != 'comment': assert matches[0][h[field]-1].value == value, f'Existing different {sheet}/{key}/{field}'
            continue
        number = ws.max_row+1
        for col in range(1,ws.max_column+1): ws.cell(number,col)._style = copy(ws.cell(5,col)._style)
        for field,value in values.items(): ws.cell(number,h[field],value)
        ws.cell(number,2).comment = Comment('B1016 source mapped; stops at P015; no completion or brewing', 'B-rose')
        added[sheet] += 1
    return dict(added)


def run(apply=False):
    entries, rows = plan()
    paths = (source.TEXT, NARRATIVE)
    hashes = {p:source.digest(p) for p in paths}
    books = {p:openpyxl.load_workbook(p) for p in paths}
    try:
        before = {p:{(s.title,c.coordinate):source.signature(c) for s in w for row in s for c in row} for p,w in books.items()}
        added = prepare(books, entries, rows)
        result = dict(event=EVENT, would_append=added, writes=apply, boundary='P015 before real brewing; no chapter completion')
        if not apply or not added:
            print(json.dumps(result,ensure_ascii=False,indent=2)); return
        out = ROOT/'Library/Locus/tmp/rose-content'/(datetime.now().strftime('%Y%m%d-%H%M%S-%f')+'-1016-wire')
        out.mkdir(parents=True)
        candidates = {}
        for p in paths:
            assert source.digest(p) == hashes[p], 'Concurrent change before backup'
            shutil.copy2(p,out/(p.stem+'.before.xlsx'))
            assert source.digest(out/(p.stem+'.before.xlsx')) == hashes[p], 'Backup race'
            candidate = out/(p.stem+'.candidate.xlsx'); books[p].save(candidate)
            saved = openpyxl.load_workbook(candidate)
            try:
                for (sheet,coordinate),old in before[p].items():
                    assert source.signature(saved[sheet][coordinate]) == old, f'Old cell changed: {sheet}/{coordinate}'
            finally: saved.close()
            candidates[p] = candidate
        result.update(before_sha256={str(p):h for p,h in hashes.items()}, replaced=[])
        journal = out/'manifest.json'
        def record(): journal.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
        record()
        assert all(source.digest(p)==hashes[p] for p in paths), 'Concurrent source change; candidate retained'
        for p in paths:
            assert source.digest(p)==hashes[p], 'Concurrent source change before replace; see manifest'
            candidates[p].replace(p); result['replaced'].append(str(p)); record()
        print(json.dumps({**result,'evidence':str(out)},ensure_ascii=False,indent=2))
    finally:
        for w in books.values(): w.close()


if __name__ == '__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--apply',action='store_true')
    args=parser.parse_args()
    if args.apply:
        with source_write_guard('configure 1016 P004-P015'): run(True)
    else: run()
