"""Plan isolated B1009 portrait and B1012 shared terminal fade in memory.

Default and --self-test never save files or acquire a write lock. Only the parent
may invoke --apply, serially, after review. Apply changes Narrative.xlsx only,
using rose_write_guard, input CAS, verified backup and cell-preserving candidate.
No generation, Unity, Text/schema/code edits, formal activation or results.
"""
from __future__ import annotations

import argparse
from collections import Counter
from copy import copy, deepcopy
from datetime import datetime
from io import BytesIO
import json
from pathlib import Path
import shutil
import sys
import xml.etree.ElementTree as ET

# Read-only invocation must not create/update import caches either.
sys.dont_write_bytecode = True

import openpyxl
from openpyxl.comments import Comment
import import_rose_text as source
import configure_rose_1012 as b1012
from rose_write_guard import source_write_guard

ROOT = source.ROOT
NARRATIVE = ROOT / 'Assets/ConfigTables/Data/Narrative/Narrative.xlsx'
SCHEMA = ROOT / 'Assets/ConfigTables/Defines/schema.xml'
PACKAGE = source.MAP.parent
E09 = 'rose.1009.s1'
E12 = 'rose.1012.s1'
FRONT = 'dependency:B1009.stage.front'
PORTRAIT = 'dependency:B1009.stage.portrait'
POST = 'dependency:B-C02.stage.post_brew'
PORTRAIT_ID = E09 + '.portrait'
FADE_ID = E12 + '.terminal_fade'
OLD_ROOT = 'entry=MapAction;blocksSave=true;blockedBy=B-B01+B-B02;contentBoundary=P071'
NEW_ROOT = OLD_ROOT.replace('contentBoundary=P071',
                           'contentBoundary=isolated-stages;frontBoundary=P071;portraitBoundary=P073-P076')
FRONT_PAGES = ('door', 'blindfold', 'reception', 'origin_letter', 'origin_uncle', 'request')
SHEETS = ('Event', 'SetupParams', 'PresentationIndex', 'ConversationPage',
          'ConversationLine', 'Choice', 'Condition', 'Object')
FORBIDDEN = ('Result', 'Action', 'EventDate', 'ConditionKey', 'CafeVisitHook',
             'AbilityCheck', 'AbilityCheckCondition')


def require(ok, message):
    if not ok:
        raise RuntimeError(message)


def signature(cell):
    link = cell.hyperlink
    return (cell.value, cell.data_type, tuple(cell._style) if cell.has_style else (0,) * 9,
            (cell.comment.text, cell.comment.author) if cell.comment else None,
            (link.target, link.location, link.tooltip, link.display, link.id) if link else None)


def snapshot(book):
    return {(ws.title, c.coordinate): signature(c) for ws in book for row in ws for c in row}


def validate_schema(schema=None):
    schema = ET.parse(SCHEMA).getroot() if schema is None else schema
    for sheet in SHEETS:
        require(schema.find(".//bean[@name='Ncs" + sheet + "Entry']") is not None,
                'Missing schema bean: ' + sheet)
    field = schema.find(".//bean[@name='NcsConversationPageEntry']/var[@name='stepType']")
    require(field is not None and field.get('type', '').startswith('string#set=')
            and 'FadeOut' in field.get('type').split('=', 1)[1].split(','),
            'Current schema does not support FadeOut; refuse without editing schema')
    return schema


def table(book, sheet, schema=None):
    require(sheet in book.sheetnames, 'Required sheet missing: ' + sheet)
    ws = book[sheet]
    heads = [c.value for c in ws[1]]
    require(len(heads) >= 3 and heads[0] == '##var' and all(isinstance(h, str) and h for h in heads)
            and len(heads) == len(set(heads)), 'Invalid/duplicate headers: ' + sheet)
    if schema is not None:
        bean = schema.find(".//bean[@name='Ncs" + sheet + "Entry']")
        fields = [v.get('name') for v in bean.findall('var')]
        require(heads[1] == fields[0] and set(heads[1:]) == set(fields), 'Workbook/schema header drift: ' + sheet)
    elif sheet == 'Text':
        require(heads == ['##var', 'key', 'text', 'comment'], 'Text headers changed')
    index = {}
    for row in ws.iter_rows(min_row=5):
        key = row[1].value
        if key is None:
            require(not any(c.value is not None for c in row), f'Data row without key: {sheet}/{row[0].row}')
            continue
        require(isinstance(key, str) and key and key not in index, f'Invalid/duplicate key: {sheet}/{key}')
        index[key] = {h: c for h, c in zip(heads[1:], row[1:])}
    return ws, heads, index


def values(cells):
    return {h: c.value for h, c in cells.items()}


def build_plan():
    entries = {event: source.expected_rows(event) for event in ('1009', '1012')}
    # Importing B1012 is safe (main is guarded); its canonical row builder verifies
    # DOCX/XML/source-map and preserves its existing stage/history guards exactly.
    rows12, _ = b1012.build_plan()
    baseline = [(r['sheet'], dict(r['values'])) for r in rows12 if r['sheet'] != 'Text']
    baseline += [
        ('Event', dict(eventId=E09, eventKind='actor', actorId='player', titleText=E09 + '.title',
                       setupType='Tpl_ArrivalConversation', setupKey=E09, onceOnly=True,
                       presentationClass='Conversation', presentationKey=E09 + '.conversation', eventProperty=OLD_ROOT)),
        ('SetupParams', dict(setupKey=E09, tplName='Tpl_ArrivalConversation', slotValues='opts=', skipAble=False, autoAdvance=False)),
        ('PresentationIndex', dict(presentationKey=E09 + '.conversation', presentationClass='Conversation', logTag='B1009 blocked P003-P071')),
    ]
    for seq, short in enumerate(FRONT_PAGES, 1):
        page = dict(pageId=E09 + '.' + short, presentationKey=E09 + '.conversation', seq=seq)
        if short in ('door', 'blindfold', 'reception'):
            page['bgOverride'] = 'rose.club.' + ('exterior' if short == 'door' else short)
        if short == 'reception':
            page['optionGroupKey'] = E09 + '.choice.origin'
        if short.startswith('origin_'):
            page['pageCondKeys'] = 'choice:' + E09 + '.choice.p' + ('045' if short == 'origin_letter' else '046')
        baseline.append(('ConversationPage', page))
    seqs = Counter()
    tail = []
    for entry in entries['1009']:
        number = int(entry['source'].rsplit('P', 1)[1])
        if entry['kind'] == '选项':
            continue
        seqs[entry['page']] += 1
        raw = entry['original'].strip()
        speaker = ('player' if raw.startswith('{playername}') else
                   'rose.butake' if raw.startswith(('男人', '布塔克')) and '：' in raw[:8] else
                   'rose.envoy' if raw.startswith('老爷：') else None)
        if number in (38, 43, 51) and entry['key'].endswith('.b'):
            speaker = None
        line = dict(lineId=entry['key'].replace('.text.', '.'), pageId=entry['page'],
                    seq=seqs[entry['page']], side='none', textKey=entry['key'])
        if speaker:
            line['speakerKey'] = speaker
        (tail if number >= 73 else baseline).append(('ConversationLine', line))
    for seq, number in enumerate((45, 46), 1):
        baseline.append(('Choice', dict(choiceId=f'{E09}.choice.p{number:03}', groupKey=E09 + '.choice.origin',
                                        textKey=f'{E09}.text.p{number:03}', seq=seq)))
    for suffix, kind, target, value in (
        ('day', 'BAR_BETWEEN', 'world.calendar', '4,4'),
        ('meeting', 'BOOL_EQ', 'story.oct09.meeting.known', 'true'),
    ):
        baseline.append(('Condition', dict(conditionId=E09 + '.' + suffix, ownerType='event', ownerId=E09,
                                           groupIndex=1, condType=kind, target=target, value=value)))
    require(len(tail) == 4 and [v['textKey'] for _, v in tail] ==
            [f'{E09}.text.p{n:03}' for n in range(73, 77)], 'Portrait source mapping changed')
    require(sum(s == 'ConversationLine' for s, v in baseline if str(next(iter(v.values()))).startswith(E09)) == 62,
            'B1009 front source coverage changed')
    additions = [
        ('ConversationPage', dict(pageId=PORTRAIT_ID, presentationKey=E09 + '.conversation', seq=7,
                                  pageCondKeys=PORTRAIT, bgOverride='rose.club.reception', stepType='Dialogue')),
        *tail,
        ('ConversationPage', dict(pageId=FADE_ID, presentationKey=E12 + '.conversation', seq=13,
                                  pageCondKeys=POST, stepType='FadeOut')),
    ]
    return entries, baseline, additions


def check_row(sheet, key, cells, expected, migrations, stage_objects):
    for field, cell in cells.items():
        want = expected.get(field)
        if field in ('stageLeftKey', 'stageRightKey', 'stageCenterKey') and field not in expected:
            # Existing stage reconstruction is owned by the separate slot tool.
            # Validate references and preserve the exact cell; do not blank it.
            require(cell.value in (None, '-') or cell.value in stage_objects,
                    f'Unknown existing stage Object: {sheet}/{key}/{field}')
            continue
        allowed = migrations.get((sheet, key, field))
        require(cell.value in allowed if allowed else cell.value == want,
                f'Existing row differs: {sheet}/{key}/{field}: {cell.value!r}')
        if field == 'seq':
            require(type(cell.value) is int, f'Non-integer sequence: {sheet}/{key}')


def validate_routes(tables):
    pages = [values(c) for c in tables['ConversationPage'][2].values()]
    lines = [values(c) for c in tables['ConversationLine'][2].values()]
    routes = []
    for event, count in ((E09, 7), (E12, 13)):
        own = sorted((p for p in pages if p['presentationKey'] == event + '.conversation'), key=lambda p: p['seq'])
        require([p['seq'] for p in own] == list(range(1, count + 1)), 'Page sequence drift: ' + event)
        fades = [p for p in own if p['stepType'] == 'FadeOut']
        require(len(fades) == (1 if event == E12 else 0), 'Unexpected fade count: ' + event)
        for p in fades:
            require(p['pageId'] == FADE_ID and p['seq'] == own[-1]['seq']
                    and not any(l['pageId'] == p['pageId'] for l in lines)
                    and all(p[k] is None for k in ('optionGroupKey', 'checkId', 'readingKey')),
                    'FadeOut must be a single final empty step')
    # Exercise the existing B1012 fixture evaluator with the shared final page.
    rows = [{'sheet': 'ConversationPage', 'values': p} for p in pages if p['presentationKey'] == E12 + '.conversation']
    rows += [{'sheet': 'ConversationLine', 'values': l} for l in lines if str(l['lineId']).startswith(E12 + '.')]
    for a in (False, True):
        for h in (False, True):
            fixture = b1012.fixture_plan(rows, 'post_brew', a, h, True)
            selected = fixture['pages']
            require(selected[-1] == FADE_ID and selected.count(FADE_ID) == 1, 'Route missing final shared fade')
            expected_last = E12 + ('.p085' if a else '.p053')
            require(fixture['lines'][-1] == expected_last, 'Route last dialogue changed')
            require((E12 + '.visit_brief' in selected) == (not a) and
                    (E12 + '.full_end' in selected) == a and
                    (E12 + '.hat_helped' in selected) == (a and h) and
                    (E12 + '.hat_not_helped' in selected) == (a and not h), 'History branch crossed')
            routes.append(dict(A=a, H=h, D=True, last_line=expected_last, final_page=FADE_ID))
    for stage in ('preparation', 'brew_retry'):
        require(FADE_ID not in b1012.fixture_plan(rows, stage, None, None, None)['pages'], 'Fade leaked into another stage')
    return routes


def prepare(book, text, entries, baseline, additions):
    schema = validate_schema()
    tables = {sheet: table(book, sheet, schema) for sheet in SHEETS}
    texts = table(text, 'Text')[2]
    for event, expected in entries.items():
        prefix = f'rose.{event}.s1.text.'
        require({k for k in texts if k.startswith(prefix)} == {e['key'] for e in expected}, 'Unexpected/missing source Text keys: ' + event)
        for entry in expected:
            require(texts[entry['key']]['text'].value == entry['text'], 'Source Text differs: ' + entry['key'])
    for key, value in ((E09 + '.title', '赴约'), ('rose.envoy.name', '老爷'), ('rose.butake.name', '男人'),
                       (b1012.ALIAS, '埃莉诺')):
        require(key in texts and texts[key]['text'].value == value, 'Required display Text differs: ' + key)
    objects = tables['Object'][2]
    for key in ('player', 'rose.envoy', 'rose.butake', 'adele', 'alienor'):
        require(key in objects and objects[key]['objType'].value == ('player' if key == 'player' else 'npc'), 'Object identity conflict: ' + key)
        require(objects[key]['nameText'].value in texts, 'Object name Text missing: ' + key)
    for sheet in FORBIDDEN:
        require(sheet in book.sheetnames, 'Required safety sheet missing: ' + sheet)
        for row in book[sheet].iter_rows(min_row=5):
            require(not any(isinstance(c.value, str) and any(e in c.value for e in (E09, E12)) for c in row),
                    'Blocked root has activation/result/fact/check reference: ' + sheet)
    migrations = {('Event', E09, 'eventProperty'): (OLD_ROOT, NEW_ROOT)}
    for sheet, expected in baseline:
        if sheet == 'ConversationPage' and expected['presentationKey'] == E09 + '.conversation':
            old = expected.get('pageCondKeys')
            new = FRONT + ('|' + old if old else '')
            migrations[(sheet, expected['pageId'], 'pageCondKeys')] = (old, new)
    expected_ids = {sheet: set() for sheet in SHEETS if sheet != 'Object'}
    for sheet, expected in baseline + additions:
        ws, heads, index = tables[sheet]
        key = expected[heads[1]]
        require(key not in expected_ids[sheet], 'Duplicate planned key: ' + key)
        expected_ids[sheet].add(key)
        if key in index:
            check_row(sheet, key, index[key], expected, migrations, objects)
        else:
            require((sheet, expected) in additions, f'Required baseline missing: {sheet}/{key}')
    for sheet, ids in expected_ids.items():
        for key, cells in tables[sheet][2].items():
            own = any(key.startswith(e + '.') or key == e for e in (E09, E12))
            linked = any(isinstance(c.value, str) and any(c.value == e or c.value.startswith(e + '.') for e in (E09, E12)) for c in cells.values())
            require(not (own or linked) or key in ids, f'Unexpected root row: {sheet}/{key}')
    changes = []
    for (sheet, key, field), (old, new) in migrations.items():
        cell = tables[sheet][2][key][field]
        if cell.value != new:
            changes.append(dict(sheet=sheet, key=key, cell=cell.coordinate, old=cell.value, new=new))
            cell.value = new
    appended = []
    for sheet, expected in additions:
        ws, heads, index = tables[sheet]
        key = expected[heads[1]]
        if key in index:
            continue
        number = ws.max_row + 1
        for col in range(1, ws.max_column + 1):
            ws.cell(number, col)._style = copy(ws.cell(5, col)._style)
        for field, value in expected.items():
            ws.cell(number, heads.index(field) + 1, value)
        ws.cell(number, 2).comment = Comment(
            'Isolated blocked B preview. Portrait uses existing reception placeholder, NOT CG; '
            'no inferred perception threshold/failure route. Shared final fade represents P054/P086.', 'B-rose')
        index[key] = {h: ws.cell(number, col) for col, h in enumerate(heads, 1) if col > 1}
        appended.append(dict(sheet=sheet, key=key, row=number, values=expected))
    routes = validate_routes(tables)
    return dict(migrations=changes, append=appended, routes=routes,
                would_append=dict(Counter(r['sheet'] for r in appended)),
                formal_enabled=False, runtime_preview_verified=False, schema_fadeout_supported=True,
                portrait_visual=dict(key='rose.club.reception', placeholder=True, is_cg=False),
                blocked=['B1009: B-B01 perception threshold/failure scope and B-B02 map entry remain unresolved',
                         'B1012: real brewing/ratings, business entry, immutable history and runtime fixture remain separate dependencies'])


def verify_preservation(saved, before, expected, migrations):
    allowed = {(m['sheet'], m['cell']): m for m in migrations}
    for address, old in before.items():
        actual = signature(saved[address[0]][address[1]])
        if address in allowed:
            migration = allowed[address]
            require(old[0] == migration['old'] and actual == (migration['new'], 's', *old[2:]),
                    'Explicit migration altered cell metadata: ' + str(address))
        else:
            require(actual == old, 'Existing cell altered: ' + str(address))
    require(snapshot(saved) == expected, 'Candidate does not match full in-memory workbook')


def self_test(book, text, entries, baseline, additions):
    """Corrupt in-memory cells only, then ensure validation refuses every case."""
    passed = []

    def reject(label, mutate):
        restore = mutate()
        try:
            try:
                prepare(book, text, entries, baseline, additions)
            except RuntimeError:
                passed.append(label)
            else:
                raise RuntimeError('Negative self-test accepted: ' + label)
        finally:
            restore()

    def change(wb, sheet, key, field, value):
        cell = table(wb, sheet)[2][key][field]
        old = cell.value
        cell.value = value
        return lambda: setattr(cell, 'value', old)

    for label, wb, sheet, key, field, value in (
        ('front guard drift', book, 'ConversationPage', E09 + '.door', 'pageCondKeys', 'dependency:invented'),
        ('root boundary drift', book, 'Event', E09, 'eventProperty', 'entry=Immediate'),
        ('original sequence drift', book, 'ConversationLine', E09 + '.p071', 'seq', 99),
        ('original dialogue drift', text, 'Text', E12 + '.text.p053', 'text', 'changed'),
        ('portrait source drift', text, 'Text', E09 + '.text.p073', 'text', 'changed'),
        ('portrait check forbidden', book, 'ConversationPage', PORTRAIT_ID, 'checkId', 'invented'),
        ('fade not final', book, 'ConversationPage', FADE_ID, 'seq', 8),
        ('fade options forbidden', book, 'ConversationPage', FADE_ID, 'optionGroupKey', 'invented'),
        ('fade wrong stage', book, 'ConversationPage', FADE_ID, 'pageCondKeys', FRONT),
        ('history guard drift', book, 'ConversationPage', E12 + '.visit_brief', 'pageCondKeys', POST),
        ('unknown stage Object', book, 'ConversationLine', E12 + '.p053', 'stageLeftKey', 'invented'),
    ):
        reject(label, lambda wb=wb, sheet=sheet, key=key, field=field, value=value: change(wb, sheet, key, field, value))

    def extra(sheet, key, field=None, value=None, wb=book):
        ws, heads, _ = table(wb, sheet)
        row = ws.max_row + 1
        ws.cell(row, 2, key)
        if field:
            ws.cell(row, heads.index(field) + 1, value)
        return lambda: ws.delete_rows(row)

    reject('duplicate Text key', lambda: extra('Text', E09 + '.text.p073', wb=text))
    reject('unexpected root page', lambda: extra('ConversationPage', E12 + '.extra'))
    reject('result activation', lambda: extra('Result', 'tail-test', 'eventId', E09))
    reject('action activation', lambda: extra('Action', 'tail-test', 'resultEvents', E12))
    reject('date activation', lambda: extra('EventDate', E12))
    schema = deepcopy(ET.parse(SCHEMA).getroot())
    schema.find(".//bean[@name='NcsConversationPageEntry']/var[@name='stepType']").set('type', 'string#set=,Dialogue')
    try:
        validate_schema(schema)
    except RuntimeError:
        passed.append('schema without FadeOut')
    else:
        raise RuntimeError('Negative self-test accepted missing FadeOut')
    return passed


def run(apply=False, tests=False):
    # Everything affecting validation is watched, including the read-only Text,
    # source documents, package, canonical builder, shared guard and loader rule.
    docx = [Path(source.verify_docx(e, source.source_rows(e))) for e in ('1009', '1012')]
    watched = list(dict.fromkeys([
        NARRATIVE, source.TEXT, SCHEMA, source.MAP, PACKAGE / 'spec.md', PACKAGE / 'handoff.md',
        Path(source.__file__), Path(b1012.__file__), Path(__file__),
        ROOT / 'Tools/Narrative/configure_rose_1009.py', ROOT / 'Tools/Narrative/rose_write_guard.py',
        ROOT / 'Assets/Scripts/Narrative/LubanNarrativeLoader.ConversationSteps.cs', *docx,
    ]))
    initial = {p: source.digest(p) for p in watched}

    def cas():
        require(all(source.digest(p) == h for p, h in initial.items()), 'Concurrent input change; rerun from current sources')

    entries, baseline, additions = build_plan()
    book = openpyxl.load_workbook(NARRATIVE)
    text = openpyxl.load_workbook(source.TEXT)
    try:
        before = snapshot(book)
        text_before = snapshot(text)
        plan = prepare(book, text, entries, baseline, additions)
        expected = snapshot(book)
        manifest = dict(plan, writes=False, changed=bool(plan['append'] or plan['migrations']),
                        parent_serial_apply_only=True, input_sha256={str(p): h for p, h in initial.items()})
        if tests:
            manifest['negative_self_tests'] = self_test(book, text, entries, baseline, additions)
        again = prepare(book, text, entries, baseline, additions)
        require(not again['append'] and not again['migrations'], 'Plan is not idempotent')
        require(snapshot(text) == text_before, 'Read-only Text workbook changed in memory')
        # Round-trip entirely in RAM, including default mode; no temporary files.
        stream = BytesIO()
        book.save(stream)
        stream.seek(0)
        saved = openpyxl.load_workbook(stream)
        try:
            verify_preservation(saved, before, expected, plan['migrations'])
            saved_plan = prepare(saved, text, entries, baseline, additions)
            require(not saved_plan['append'] and not saved_plan['migrations'], 'Round-trip lost idempotency')
        finally:
            saved.close()
        manifest['checks'] = ['source DOCX hash/XML/map and all existing Text', 'schema/workbook headers and unique keys',
                              'strict existing baseline and blocked roots', 'four A/H routes, shared final fade and stage isolation',
                              'idempotency', 'RAM xlsx round-trip and old-cell metadata preservation']
        cas()
        if apply and manifest['changed']:
            out = ROOT / 'Library/Locus/tmp/rose-content' / (datetime.now().strftime('%Y%m%d-%H%M%S-%f') + '-tail-previews')
            out.mkdir(parents=True)
            backup = out / 'Narrative.before.xlsx'
            cas()
            shutil.copy2(NARRATIVE, backup)
            require(source.digest(backup) == initial[NARRATIVE], 'Backup race; no source replacement')
            candidate = out / 'Narrative.candidate.xlsx'
            book.save(candidate)
            saved = openpyxl.load_workbook(candidate)
            try:
                verify_preservation(saved, before, expected, plan['migrations'])
                candidate_plan = prepare(saved, text, entries, baseline, additions)
                require(not candidate_plan['append'] and not candidate_plan['migrations'], 'Candidate validation changed rows')
            finally:
                saved.close()
            manifest.update(evidence=str(out), backup_sha256=source.digest(backup), candidate_sha256=source.digest(candidate), replaced=False)
            journal = out / 'manifest.json'
            journal.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
            cas()
            candidate.replace(NARRATIVE)
            manifest.update(writes=True, replaced=True, after_sha256=source.digest(NARRATIVE))
            journal.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
        print(json.dumps(manifest, ensure_ascii=False, indent=2))
    finally:
        book.close()
        text.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--apply', action='store_true', help='Parent only: serial guarded Narrative.xlsx write; never generates')
    parser.add_argument('--self-test', action='store_true', help='Additional in-memory refusal tests; no files saved')
    args = parser.parse_args()
    require(not (args.apply and args.self_test), '--self-test forbids --apply')
    if args.apply:
        with source_write_guard('B1009/B1012 isolated tail previews; parent serial apply'):
            run(apply=True)
    else:
        run(tests=args.self_test)


if __name__ == '__main__':
    main()
