import json
from pathlib import Path
import re
import tempfile
import unittest
import openpyxl
import cafe_orders
from apply_cafe_orders import apply


class CafeOrderTests(unittest.TestCase):
    def test_counts_ids_references_and_authorship(self):
        data = cafe_orders.build()
        self.assertEqual({k: len(v) for k, v in data.items()},
                         {'OpeningPresentation': 24, 'OpeningLine': 96, 'Text': 120})
        texts = {r['key']: r['text'] for r in data['Text']}
        self.assertEqual(len(texts), 120)
        ids = {r['presentationId'] for r in data['OpeningPresentation']}
        self.assertEqual(ids, {f'cafe.ambient.order.{v}-{i:02}'
                              for v in ('extra', 'solo') for i in range(1, 13)})
        for presentation in data['OpeningPresentation']:
            key = presentation['presentationId']
            self.assertIn(presentation['titleKey'], texts)
            lines = [r for r in data['OpeningLine'] if r['presentationId'] == key]
            self.assertEqual([r['seq'] for r in lines], [1, 2, 3, 4])
            bodies = [texts[r['textKey']] for r in lines]
            self.assertTrue(all(r['lineId'] == f'{key}.{r["seq"]}' for r in lines))
            self.assertEqual(sum(s.count('{饮品名}') for s in bodies), 1)
            self.assertTrue(bodies[0].startswith('顾客：'))
            speakers = {s.split('：', 1)[0] for s in bodies}
            self.assertEqual(speakers, {'顾客', '我', '阿黛尔'} if '.extra-' in key else {'顾客', '我'})
            self.assertFalse(re.search(r'\d+月|\d+日|明天|下次|下雨|晴天|感谢节|阿黛尔不在', ''.join(bodies)))

    def test_append_preserves_existing_values_json_bytes_and_is_idempotent(self):
        data = cafe_orders.build()
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            headers = {
                'OpeningPresentation': ['presentationId', 'titleKey', 'leftNameKey', 'rightNameKey', 'visualsResource', 'centerNameKey'],
                'OpeningLine': ['lineId', 'presentationId', 'seq', 'speaker', 'textKey', 'gesture', 'autoAdvanceSeconds'],
                'Text': ['key', 'text', 'comment'],
            }
            for path, names in [('Narrative/Narrative.xlsx', ('OpeningPresentation', 'OpeningLine')),
                                ('Shared/Text.xlsx', ('Text',))]:
                wb = openpyxl.Workbook(); wb.remove(wb.active)
                for name in names:
                    ws = wb.create_sheet(name)
                    ws.append(['##var'] + headers[name])
                    for _ in range(3): ws.append(['##'] + ['string'] * len(headers[name]))
                    ws.append([None, 'user.original', '用户修改必须保留'])
                extra = wb.create_sheet('Unrelated'); extra['A1'] = '=1+2'
                p = root / 'Assets/ConfigTables/Data' / path
                p.parent.mkdir(parents=True, exist_ok=True); wb.save(p); wb.close()
            prefixes = {}
            for name in data:
                p = root / f'Assets/Settings/Resources/Luban/tb{name.lower()}.json'
                p.parent.mkdir(parents=True, exist_ok=True)
                raw = ('[\n  ' + json.dumps({headers[name][0]: 'user.original', 'text': '用户修改必须保留'}, ensure_ascii=False) + '\n]').encode()
                p.write_bytes(raw); prefixes[p] = raw[:-2]
            self.assertEqual(apply(root), {'OpeningPresentation': 24, 'OpeningLine': 96, 'Text': 120,
                                          'OpeningPresentationJson': 24, 'OpeningLineJson': 96, 'TextJson': 120})
            for p, prefix in prefixes.items(): self.assertTrue(p.read_bytes().startswith(prefix))
            self.assertTrue(all(n == 0 for n in apply(root).values()))


if __name__ == '__main__':
    unittest.main()
