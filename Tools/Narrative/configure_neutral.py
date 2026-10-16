"""Configure source-mapped A01 blocked construction preview using existing NCS guards."""
from copy import copy
from datetime import datetime
import json, re, shutil
import openpyxl
import import_neutral_text as src

ROOT=src.ROOT
P=ROOT/'Assets/ConfigTables/Data/Narrative/Narrative.xlsx'
T=src.TEXT

def main():
    paths=[P,T]; hashes={p:src.digest(p) for p in paths}
    books={p:openpyxl.load_workbook(p) for p in paths}
    before={p:{(s.title,c.coordinate):src.signature(c) for s in w for row in s for c in row} for p,w in books.items()}
    counts={}
    def add(p,s,values):
        ws=books[p][s]; headers={c.value:c.column for c in ws[1]}
        key=next(iter(values.values()))
        found=[r for r in range(5,ws.max_row+1) if ws.cell(r,2).value==key]
        if found:
            for k,v in values.items():
                assert ws.cell(found[0],headers[k]).value==v,(s,key,k)
            return
        row=ws.max_row+1
        for c in range(1,ws.max_column+1):ws.cell(row,c)._style=copy(ws.cell(5,c)._style)
        for k,v in values.items():ws.cell(row,headers[k],v)
        counts[s]=counts.get(s,0)+1
    root='neutral.1008.s1'
    for key,text in [('colette.name','柯莱特·布吕内'),('neutral.shopkeeper.name','店主'),('neutral.unknown.name','？？？')]:
        add(T,'Text',dict(key=key,text=text,comment='A01 source identity; blocked construction'))
    for obj,key in [('colette','colette.name'),('neutral.shopkeeper','neutral.shopkeeper.name')]:
        add(P,'Object',dict(objId=obj,objType='npc',nameText=key))
    add(P,'Event',dict(eventId=root,eventKind='world',actorId='world',setupKey=root,onceOnly=True,presentationClass='Conversation',presentationKey=root,eventProperty='entry=MapAction;blocksSave=true;blockedBy=A-B010-relationship-definition'))
    add(P,'SetupParams',dict(setupKey=root,tplName='Tpl_ArrivalConversation',skipAble=False,autoAdvance=False))
    add(P,'PresentationIndex',dict(presentationKey=root,presentationClass='Conversation',logTag='A01 blocked construction'))
    add(P,'Condition',dict(conditionId=root+'.day',ownerType='event',ownerId=root,groupIndex=1,condType='BAR_BETWEEN',target='world.calendar',value='3,3'))
    add(P,'AbilityCheck',dict(checkId=root+'.language',mode='Inline'))
    add(P,'AbilityCheckCondition',dict(conditionId=root+'.language.1',checkId=root+'.language',attributeId='language',threshold=2,groupIndex=1,seq=1))
    page=''; seq=0; line_seq=0
    for line in src.MAP.read_text(encoding='utf-8').splitlines():
        if line.startswith('## A02'):break
        match=re.match(r'### A01 / (\w+)',line)
        if match:
            page=match.group(1);seq+=1;line_seq=0
            values=dict(pageId=root+'.'+page,presentationKey=root,seq=seq)
            if page in ['market','outro']:values['bgOverride']='neutral.market.day'
            if page=='shop':values.update(bgOverride='neutral.market.shop.day',optionGroupKey=root+'.q1')
            guards={'help':'choice:'+root+'.choice.q1.a','talk_pass':'choice:'+root+'.choice.q1.b|check:'+root+'.language','talk_fail':'choice:'+root+'.choice.q1.b|not:check:'+root+'.language','wait':'choice:'+root+'.choice.q1.c'}
            if page in guards:values['pageCondKeys']=guards[page]
            add(P,'ConversationPage',values)
        if not page or not re.match(r'^\| P\d{3} \| 正文 \|',line):continue
        parts=line.split('|');num=int(re.search(r'P(\d+)',parts[1]).group(1));raw=parts[4].strip();line_seq+=1
        speaker='player' if raw.startswith('{playername}') else 'colette' if raw.startswith(('柯莱特·布吕内','？？？')) else 'neutral.shopkeeper' if raw.startswith('店主：') else None
        emote=re.search(r'（#([^）]+)）',raw)
        values=dict(lineId=root+f'.line.p{num:03}',pageId=root+'.'+page,seq=line_seq,side='right' if speaker=='colette' and num>=23 else 'none',textKey=root+f'.text.p{num:03}')
        if speaker:values['speakerKey']=speaker
        if raw.startswith('？？？'):values['speakerNameKey']='neutral.unknown.name'
        if emote:values['emote']=emote.group(1)
        add(P,'ConversationLine',values)
    for seq,(label,num) in enumerate([('a',36),('b',37),('c',38)],1):
        add(P,'Choice',dict(choiceId=root+'.choice.q1.'+label,groupKey=root+'.q1',textKey=root+f'.text.p{num:03}',seq=seq))
    out=ROOT/'Library/Locus/tmp/neutral-content'/datetime.now().strftime('%Y%m%d-%H%M%S-a01-config');out.mkdir(parents=True)
    candidates={}
    for p,w in books.items():
        shutil.copy2(p,out/(p.stem+'.before.xlsx'));candidate=out/p.name;w.save(candidate);candidates[p]=candidate
        saved=openpyxl.load_workbook(candidate)
        for (s,c),v in before[p].items():assert src.signature(saved[s][c])==v,(s,c)
    for p in paths:assert src.digest(p)==hashes[p],f'Concurrent modification: {p}'
    for p,c in candidates.items():c.replace(p)
    manifest=dict(time=datetime.now().astimezone().isoformat(),counts=counts,scope='Blocked result-free A01; no formal Action, mount, dates or relationship values',before={str(p):h for p,h in hashes.items()})
    (out/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8');print(manifest);print(out)
if __name__=='__main__':main()
