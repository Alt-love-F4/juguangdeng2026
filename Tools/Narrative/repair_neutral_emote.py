"""Repair only A04 P007/P090/P145 emotes, with source verification and SHA CAS.

Default is a read-only plan. --apply requires its --expected-sha256. Patches
three worksheet XML cells only; all other ZIP member bytes and ZIP metadata
are preserved. Never saves through openpyxl, generates data, or calls Unity.
A second apply is a byte-preserving no-op. Evidence is kept under Library.
"""
from __future__ import annotations

import argparse
from copy import deepcopy
from datetime import datetime
import hashlib
import json
from pathlib import Path
import posixpath
import re
import shutil
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import openpyxl
from openpyxl.utils.cell import column_index_from_string

import import_neutral_text as source
from repair_neutral_stage import metadata, signature_sha256
from configure_neutral_remaining import MAP, NARRATIVE, ROOT, digest, require

# No dynamic expansion from the read-only inventory is permitted.
TARGETS = (
    (7, 6, "morning", "adele", "阿黛尔·米尼耶", "开心"),
    (90, 89, "visit", "antoine", "安托万·马尚", "平静"),
    (145, 144, "closing", "adele", "阿黛尔·米尼耶", "疲惫"),
)
ROOT_ID = "neutral.1013.s2"
NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
CELL_RE = re.compile(rb'<c\b[^>]*\br="([A-Z]+)([0-9]+)"[^>]*(?:/>|>.*?</c>)')


def sha_bytes(data):
    return hashlib.sha256(data).hexdigest()


def exact_signature(cell):
    """Unlike the stage helper, do not normalize empty cell data types."""
    return (cell.value, cell.data_type, tuple(cell._style) if cell.has_style else None,
            (cell.comment.text, cell.comment.author, cell.comment.width, cell.comment.height)
            if cell.comment else None,
            tuple(sorted(dict(cell.hyperlink).items())) if cell.hyperlink else None)


def snapshot(wb):
    return {(ws.title, cell.coordinate): exact_signature(cell)
            for ws in wb for row in ws for cell in row}


def mapped_rows():
    rows = []
    root = batch = page = None
    for number, line in enumerate(MAP.read_text(encoding="utf-8").splitlines(), 1):
        match = re.match(r"^## (A\d+)：(neutral\.[^ ]+) /", line)
        if match:
            batch, root = match.groups()
        match = re.match(r"^### A\d+ / (\w+)", line)
        if match:
            page = match.group(1)
        if re.match(r"^\| P\d{3} \|", line):
            cells = [x.strip() for x in line.strip("|").split("|")]
            require(len(cells) == 4, f"Unexpected map row: {number}")
            rows.append(dict(root=root, batch=batch, page=page, map_line=number,
                             paragraph=int(cells[0][1:]), kind=cells[1],
                             location=cells[2], original=cells[3].replace("<br>", "\n")))
    return rows


def source_evidence(rows):
    # Verify every paragraph (including structural cues), not just display rows.
    sources = {}
    for event in source.SOURCES:
        path = source.verify_docx(event, source.source_rows(event))
        sources[str(path)] = digest(path)
    evidence = []
    for number, cue, page, speaker, name, emote in TARGETS:
        found = [r for r in rows if r["root"] == ROOT_ID and r["paragraph"] in (cue, number)]
        require(len(found) == 2, f"Missing source pair: P{number:03}")
        control, body = found
        require(control["paragraph"] == cue and control["kind"] == "结构/说明" and
                control["original"] == f"【立绘出场：{name}（#{emote}）】",
                f"Structural emote differs: P{cue:03}")
        require(body["paragraph"] == number and body["kind"] == "正文" and
                body["original"].startswith(name + "：") and
                all(r["batch"] == "A04" and r["page"] == page for r in found) and
                f"{ROOT_ID}.line.p{number:03}" in body["location"],
                f"Body source differs: P{number:03}")
        evidence.append(dict(lineId=f"{ROOT_ID}.line.p{number:03}",
                             speaker=speaker, emote=emote, paragraphs=found))
    return sources, evidence


def other_omissions(rows, ws, headers):
    """Inventory explicit structural cues only; never adds targets."""
    by_id = {ws.cell(r, headers["lineId"]).value: r for r in range(5, ws.max_row + 1)}
    result = []
    for index, row in enumerate(rows):
        match = re.fullmatch(r"【立绘出场：(.+)（#([^）]+)）】", row["original"])
        if row["kind"] != "结构/说明" or not match:
            continue
        name, emote = match.groups()
        for body in rows[index + 1:]:
            if (body["root"], body["page"]) != (row["root"], row["page"]):
                break
            if body["kind"] != "正文":
                continue
            key = f'{row["root"]}.line.p{body["paragraph"]:03}'
            if row["root"] == ROOT_ID and body["paragraph"] in {t[0] for t in TARGETS}:
                break
            target = by_id.get(key)
            if body["original"].startswith(name + "：") and target is not None:
                cell = ws.cell(target, headers["emote"])
                if cell.value is None:
                    result.append(dict(lineId=key, cell=cell.coordinate, emote=emote,
                                       source=row, following_body=body, written=False))
            break
    return result


def worksheet_member(archive):
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {item.attrib["Id"]: item.attrib["Target"] for item in relationships}
    sheets = [s for s in workbook.find("m:sheets", NS) if s.attrib["name"] == "ConversationLine"]
    require(len(sheets) == 1, "Missing or duplicate ConversationLine sheet")
    target = targets[sheets[0].attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]]
    return posixpath.normpath(target.lstrip("/") if target.startswith("/") else "xl/" + target)


def patch_xml(original, edits):
    patched = original
    inserted = []
    for item in edits:
        coordinate = item["cell"]
        col, row = re.fullmatch(r"([A-Z]+)([0-9]+)", coordinate).groups()
        matches = list(re.finditer(rb'<row\b[^>]*\br="' + row.encode() + rb'"[^>]*>.*?</row>', patched))
        require(len(matches) == 1, f"Missing/duplicate XML row: {row}")
        row_match = matches[0]
        row_xml = row_match.group()
        cells = list(CELL_RE.finditer(row_xml))
        # Current authored empty targets have no XML node. Fail closed on a new representation.
        require(all(m.group(1).decode() + m.group(2).decode() != coordinate for m in cells),
                f"Unexpected existing XML target: {coordinate}")
        column = column_index_from_string(col)
        next_cell = next((m for m in cells if column_index_from_string(m.group(1).decode()) > column), None)
        offset = row_match.start() + (next_cell.start() if next_cell else row_xml.index(b"</row>"))
        node = f'<c r="{coordinate}" t="inlineStr"><is><t>{item["after"]}</t></is></c>'.encode("utf-8")
        patched = patched[:offset] + node + patched[offset:]
        inserted.append(node)
    stripped = patched
    for node in inserted:
        require(stripped.count(node) == 1, "Inserted node count differs")
        stripped = stripped.replace(node, b"", 1)
    require(stripped == original, "Worksheet bytes changed outside three inserted cells")
    ET.fromstring(patched)
    return patched


def zip_metadata(info):
    return tuple(getattr(info, field) for field in (
        "filename", "date_time", "compress_type", "comment", "extra", "create_system",
        "create_version", "extract_version", "flag_bits", "volume", "internal_attr", "external_attr"))


def repair(apply=False, expected_sha256=None):
    initial = digest(NARRATIVE)
    require(not apply or expected_sha256 is not None, "--apply requires --expected-sha256")
    if expected_sha256 is not None:
        require(initial == expected_sha256.lower(), "Narrative SHA differs; re-plan")
    source_sha = digest(MAP)
    rows = mapped_rows()
    sources, evidence = source_evidence(rows)
    wb = openpyxl.load_workbook(NARRATIVE)
    try:
        before, old_metadata = snapshot(wb), metadata(wb)
        dimensions = {s.title: (s.max_row, s.max_column) for s in wb}
        ws = wb["ConversationLine"]
        headers = {c.value: c.column for c in ws[1] if c.value}
        require({"lineId", "pageId", "speakerKey", "textKey", "emote"} <= headers.keys(), "Schema differs")
        targets, edits = [], []
        for number, _, page, speaker, _, emote in TARGETS:
            key = f"{ROOT_ID}.line.p{number:03}"
            matches = [r for r in range(5, ws.max_row + 1) if ws.cell(r, headers["lineId"]).value == key]
            require(len(matches) == 1, f"Missing/duplicate line: {key}")
            row = matches[0]
            for field, value in (("pageId", ROOT_ID + "." + page), ("speakerKey", speaker),
                                 ("textKey", key.replace(".line.", ".text."))):
                require(ws.cell(row, headers[field]).value == value, f"Unexpected {field}: {key}")
            comment = ws.cell(row, headers["lineId"]).comment
            require(comment is not None and json.loads(comment.text).get("batch") == "A04", f"Wrong ownership: {key}")
            cell = ws.cell(row, headers["emote"])
            require(cell.value in (None, emote), f"Unexpected authored emote: {key}")
            item = dict(sheet=ws.title, cell=cell.coordinate, lineId=key,
                        field="emote", before=cell.value, after=emote)
            targets.append(item)
            if cell.value is None:
                edits.append(item)
        require(len(edits) in (0, 3), "Partial repair state; expected all three blank or repaired")
        allowed = {(i["sheet"], i["cell"]) for i in edits}
        unchanged = {k: v for k, v in before.items() if k not in allowed}
        report = dict(tool="repair_neutral_emote", status="planned", changed=False,
                      idempotent=not edits, writes=bool(apply and edits),
                      before_sha256=initial, after_sha256=initial,
                      source_map_sha256=source_sha, source_docx_sha256=sources,
                      source_evidence=evidence, targets=targets, edits=edits,
                      cells_checked=len(before), preserved_cells=len(unchanged),
                      preserved_cells_signature_sha256=signature_sha256(unchanged),
                      other_explicit_structure_emote_omissions=other_omissions(rows, ws, headers))

        def cas_check():
            require(digest(NARRATIVE) == initial, "Concurrent Narrative change; CAS aborted")
            require(digest(MAP) == source_sha, "Concurrent source-map change; CAS aborted")
            for path, sha in sources.items():
                require(digest(Path(path)) == sha, f"Concurrent DOCX change: {path}")

        cas_check()
        if not apply or not edits:
            report["status"] = "idempotent-noop" if not edits else "planned"
            return report
        out = ROOT / "Library/Locus/tmp/neutral-content" / (datetime.now().strftime("%Y%m%d-%H%M%S-%f") + "-emote-repair")
        out.mkdir(parents=True, exist_ok=False)
        backup, candidate = out / "Narrative.before.xlsx", out / "Narrative.candidate.xlsx"
        shutil.copy2(NARRATIVE, backup)
        require(digest(backup) == initial, "Backup SHA mismatch")
        with ZipFile(backup) as original:
            names = original.namelist()
            require(len(names) == len(set(names)), "Duplicate ZIP members")
            member = worksheet_member(original)
            original_xml = original.read(member)
            patched_xml = patch_xml(original_xml, edits)
            with ZipFile(candidate, "w") as saved:
                saved.comment = original.comment
                for info in original.infolist():
                    saved.writestr(deepcopy(info), patched_xml if info.filename == member else original.read(info.filename))
        with ZipFile(backup) as original, ZipFile(candidate) as saved:
            require(saved.namelist() == original.namelist() and saved.comment == original.comment, "ZIP membership/comment changed")
            part_hashes = {}
            for name in original.namelist():
                require(zip_metadata(saved.getinfo(name)) == zip_metadata(original.getinfo(name)), f"ZIP metadata changed: {name}")
                expected = patched_xml if name == member else original.read(name)
                require(saved.read(name) == expected, f"ZIP payload changed: {name}")
                part_hashes[name] = sha_bytes(saved.read(name))
        saved = openpyxl.load_workbook(candidate)
        try:
            actual = snapshot(saved)
            require(saved.sheetnames == wb.sheetnames and actual.keys() == before.keys(), "Sheet/cell coordinates changed")
            require({s.title: (s.max_row, s.max_column) for s in saved} == dimensions, "Sheet dimensions changed")
            require(metadata(saved) == old_metadata, "Workbook/sheet metadata changed")
            desired = {(i["sheet"], i["cell"]): i["after"] for i in edits}
            for key, old in before.items():
                if key in allowed:
                    require(actual[key][:2] == (desired[key], "s") and actual[key][2:] == old[2:], f"Target readback differs: {key}")
                else:
                    require(actual[key] == old, f"Non-target value/type/style/comment/link changed: {key}")
            require(signature_sha256({k: v for k, v in actual.items() if k not in allowed}) == report["preserved_cells_signature_sha256"], "Preserved signatures differ")
        finally:
            saved.close()
        candidate_sha = digest(candidate)
        report.update(evidence=str(out), backup_sha256=digest(backup), candidate_sha256=candidate_sha,
                      after_sha256=candidate_sha, changed_cells=3, metadata_preserved=True,
                      exact_cell_types_preserved_outside_targets=True, worksheet_member=member,
                      non_target_worksheet_bytes_preserved=True, zip_member_sha256=part_hashes,
                      readback_verified=True, status="verified-candidate")
        report_path = out / "repair-evidence.json"
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        replacement = out / "Narrative.replace.xlsx"
        shutil.copy2(candidate, replacement)
        require(digest(replacement) == candidate_sha and digest(candidate) == candidate_sha and digest(backup) == initial, "Evidence/replacement SHA changed")
        cas_check()
        replacement.replace(NARRATIVE)
        require(digest(NARRATIVE) == candidate_sha, "Post-CAS SHA mismatch")
        report.update(changed=True, status="committed")
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        return report
    finally:
        wb.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--expected-sha256")
    args = parser.parse_args()
    print(json.dumps(repair(args.apply, args.expected_sha256), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
