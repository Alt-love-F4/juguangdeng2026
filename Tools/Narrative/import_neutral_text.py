"""Stage A-neutral source-mapped Text rows without activating any event.
Run from project root with --event 1008|1010.s1|1010.s2|1013.s2|1018.s1.
Generation must subsequently use Assets/ConfigTables/GenerateAll.bat.
"""
from __future__ import annotations
import argparse
from copy import copy
from datetime import datetime
import hashlib
import json
from pathlib import Path
import re
import shutil
import xml.etree.ElementTree as ET
import zipfile

import openpyxl
from openpyxl.comments import Comment

ROOT = Path(__file__).resolve().parents[2]
MAP = ROOT / 'Docs/superpowers/specs/2026-10-16-content-batches/A-neutral/source-map.md'
TEXT = ROOT / 'Assets/ConfigTables/Data/Shared/Text.xlsx'
SOURCES = {
    '1008': ('MA70.10.08-S1.docx', 86, '5d472f1b8fdc51487bb2171928e4ae2332a6d259024580886954a4f414519d58'),
    '1010.s1': ('MA70.10.10-S1.docx', 77, 'e54ba6e06d1213498379775c5f43707818ffe3bfed94bc7f34c671440c940a0f'),
    '1010.s2': ('MA70.10.10-S2.docx', 114, '788e091ab3a06de3cc2ee78d8adc388dbb9a87cd4f9a514bdd760ccd92760120'),
    '1013.s2': ('MA70.10.13-S2.docx', 153, '4cd41271eaaf9c1f3a6892ad96e52729d3bf7553f4271428b4ed7d57cb841caa'),
    '1018.s1': ('MA70.10.18-S1.docx', 97, '069aac8e8a9d08d485b381b7525bbf8c48404939fb12f57e71186df899705044'),
}
SOURCE_DIR = Path('D:/桌面/文案/中立')
PREFIX = re.compile(r'^(?:\{playername\}|店主|柯莱特·布吕内|阿黛尔·米尼耶|安托万·马尚|顾客[ABCD]|记者|？？？)\s*[：:]\s*')
EMOTE = re.compile(r'[（(]\s*#.*?[）)]')


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def source_rows(event: str):
    rows = []
    root = 'neutral.' + ('1008.s1' if event == '1008' else event)
    expected_count = {'1008': 63, '1010.s1': 55, '1010.s2': 73, '1013.s2': 106, '1018.s1': 77}[event]
    in_a01 = False
    for line in MAP.read_text(encoding='utf-8').splitlines():
        if line.startswith('## A') and ' / MA' in line:
            if in_a01:
                break
            in_a01 = root + ' /' in line
        if not in_a01 or not re.match(r'^\| P\d{3} \|', line):
            continue
        cells = [part.strip() for part in line.strip('|').split('|')]
        if len(cells) != 4:
            continue
        number = int(re.search(r'P(\d+)', cells[0]).group(1))
        kind, location = cells[1], cells[2]
        original = line.split('|')[4]
        original = original[1:] if original.startswith(' ') else original
        original = original[:-1] if original.endswith(' ') else original
        location = location.strip('`')
        keys = [item.strip() for item in location.split('/')]
        key = keys[-1] if keys else ''
        if key == '—' or kind not in ('正文', '选项', '选项标签/结构'):
            continue
        if '.text.' not in key:
            continue
        original = original.replace('<br>', '\n')
        rows.append((number, kind, key, original))
    assert len(rows) == expected_count, (event, len(rows), expected_count)
    assert len({r[0] for r in rows}) == len(rows)
    return rows


def verify_docx(event: str, rows):
    filename, count, expected = SOURCES[event]
    path = SOURCE_DIR / filename
    if not path.is_file():
        raise RuntimeError(f'Original source unavailable: {path}')
    assert digest(path) == expected, f'Original changed: {path}'
    with zipfile.ZipFile(path) as archive:
        xml = ET.fromstring(archive.read('word/document.xml'))
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    paragraphs = []
    for p in xml.findall('.//w:body//w:p', ns):
        text = ''.join(n.text or '' if n.tag.endswith('}t') else '\t' if n.tag.endswith('}tab') else '\n' if n.tag.endswith(('}br', '}cr')) else '' for n in p.iter())
        if text.strip():
            paragraphs.append(text)
    assert len(paragraphs) == count, (len(paragraphs), count)
    mapped = []
    active = False
    root = 'neutral.' + ('1008.s1' if event == '1008' else event)
    for line in MAP.read_text(encoding='utf-8').splitlines():
        if line.startswith('## A') and ' / MA' in line:
            if active:
                break
            active = root + ' /' in line
        if active and re.match(r'^\| P\d{3} \|', line):
            parts = line.split('|')
            number = int(re.search(r'P(\d+)', parts[1]).group(1))
            original = parts[4]
            original = original[1:] if original.startswith(' ') else original
            original = original[:-1] if original.endswith(' ') else original
            mapped.append((number, original.replace('<br>', '\n')))
    if [n for n, _ in mapped] != list(range(1, count + 1)):
        raise RuntimeError(f'Incomplete source-map: {event}')
    for number, original in mapped:
        if paragraphs[number - 1] != original:
            raise RuntimeError(f'DOCX/source-map mismatch at {event} P{number:03}')
    return path


def clean(raw: str, kind: str) -> str:
    text = raw.strip()
    if kind == '选项标签/结构':
        return text.removeprefix('选择')
    if kind == '选项':
        match = re.fullmatch(r'[ABC]\.「(.*)」（(?:无需检定|语言\s*≥\s*\d+|智慧\s*≥\s*\d+)）', text)
        assert match, f'Unknown choice label format: {text!r}'
        return match.group(1)
    text = re.sub(r'^(\{playername\}|店主|柯莱特·布吕内|阿黛尔·米尼耶|安托万·马尚|顾客[ABCD]|记者|？？？)(?:[（(]#[^）)]*[）)])?\s*[：:]\s*', '', text)
    return text


def signature(cell):
    return (cell.value, cell.data_type, copy(cell._style), cell.comment.text if cell.comment else None)


def stage(event: str):
    rows = source_rows(event)
    source = verify_docx(event, rows)
    entries = []
    for number, kind, key, original in rows:
        value = clean(original, kind)
        assert value and not PREFIX.match(value) and not EMOTE.search(value), f'Invalid cleaned text at P{number:03}: {original!r} -> {value!r}'
        entries.append({'key': key, 'text': value, 'source': f'A-{event}-P{number:03}', 'kind': kind, 'original': original, 'path': str(source)})
    assert len({entry['key'] for entry in entries}) == len(entries)

    initial_hash = digest(TEXT)
    wb = openpyxl.load_workbook(TEXT)
    ws = wb['Text']
    before = {(sheet.title, cell.coordinate): signature(cell) for sheet in wb for row in sheet for cell in row}
    keys = {}
    for row in ws.iter_rows(min_row=5):
        key = row[1].value
        if key:
            assert key not in keys, f'Duplicate Text key: {key}'
            keys[key] = row[1].row
    root = 'neutral.' + ('1008.s1' if event == '1008' else event)
    prefix = root + '.text.'
    expected = {entry['key'] for entry in entries}
    assert not {key for key in keys if str(key).startswith(prefix)} - expected
    added = []
    for entry in entries:
        if entry['key'] in keys:
            cell = ws.cell(keys[entry['key']], 3)
            if cell.value != entry['text']:
                # Repair only this importer's initial A01 choice labels, retaining source comments.
                assert event == '1008' and entry['kind'] == '选项' and cell.value == re.sub(r'^[ABC]\.\s*', '', entry['original']), entry['key']
                before.pop(('Text', cell.coordinate))
                cell.value = entry['text']
                added.append(cell.row)
            continue
        row = ws.max_row + 1
        for col in range(1, 5):
            ws.cell(row, col)._style = copy(ws.cell(599, col)._style)
        ws.cell(row, 2, entry['key'])
        ws.cell(row, 3, entry['text'])
        ws.cell(row, 4, f"{entry['source']}; {entry['kind']}; text staging only, event disabled")
        ws.cell(row, 3).comment = Comment(f"{entry['path']}\n{entry['source']}\n{entry['original']}", 'A-neutral source')
        added.append(row)
    stamp = datetime.now().strftime('%Y%m%d-%H%M%S-%f')
    out = ROOT / f'Library/Locus/tmp/neutral-content/{stamp}-{event}'
    out.mkdir(parents=True)
    shutil.copy2(TEXT, out / 'Text.before.xlsx')
    candidate = out / 'Text.candidate.xlsx'
    wb.save(candidate)
    saved = openpyxl.load_workbook(candidate)
    for (sheet, coordinate), value in before.items():
        assert signature(saved[sheet][coordinate]) == value, (sheet, coordinate)
    assert digest(TEXT) == initial_hash, 'Concurrent Text.xlsx change; candidate preserved'
    if added:
        candidate.replace(TEXT)
    manifest = {'time': datetime.now().astimezone().isoformat(), 'event': root, 'source_sha256': SOURCES[event][2], 'source_paragraphs': SOURCES[event][1], 'text_rows': len(entries), 'added_rows': added, 'before_sha256': initial_hash, 'after_sha256': digest(TEXT), 'preserved_existing_cells': len(before), 'entries': entries, 'scope': 'Text only. No Event, Date, Action, Result, Choice, or completed fact created.'}
    (out / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({key: value for key, value in manifest.items() if key != 'entries'}, ensure_ascii=False))
    print('Evidence:', out)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--event', choices=SOURCES, required=True)
    stage(parser.parse_args().event)
