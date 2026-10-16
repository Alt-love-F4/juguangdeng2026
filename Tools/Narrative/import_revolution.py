"""Incremental, source-mapped C package authoring. Official GenerateAll.bat owns generated output."""
from pathlib import Path
import re, hashlib, shutil, json, argparse, zipfile
import xml.etree.ElementTree as ET
import openpyxl
from openpyxl.comments import Comment

ROOT = Path(__file__).resolve().parents[2]
PACKAGE = ROOT / 'Docs/superpowers/specs/2026-10-16-content-batches/C-revolution'

def run(chapter):
    roots = {'C13': 'revolution.1013.s2', 'C28': 'revolution.1028.s1', 'C14': 'revolution.1014.s1'}
    root = roots[chapter]
    paths = [ROOT/'Assets/ConfigTables/Data/Narrative/Narrative.xlsx', ROOT/'Assets/ConfigTables/Data/Shared/Text.xlsx']
    hashes = [hashlib.sha256(p.read_bytes()).hexdigest() for p in paths]
    books = [openpyxl.load_workbook(p) for p in paths]
    n, t = books
    if any(r[1] == root for r in n['Event'].values):
        raise RuntimeError('Already configured: ' + root)
    backup = ROOT / 'Library/Locus/tmp/C-revolution' / chapter
    backup.mkdir(parents=True, exist_ok=True)
    for p in paths: shutil.copy2(p, backup / p.name)
    def add(sheet, **values):
        s = t['Text'] if sheet == 'Text' else n[sheet]
        header = [c.value for c in s[1]]
        s.append([values.get(h) for h in header])
        return s.max_row

    line_sheet = n['ConversationLine']
    line_header = [c.value for c in line_sheet[1]]
    if 'speakerNameKey' not in line_header:
        col = line_sheet.max_column + 1
        for row, value in enumerate(['speakerNameKey', 'string', 'c', '逐行显示名 Text key'], 1):
            line_sheet.cell(row, col, value)
        line_header.append('speakerNameKey')
    def text(key, value, source):
        add('Text', key=key, text=value, comment=source)
        t['Text'].cell(t['Text'].max_row, 3).comment = Comment(source, 'C source map')
    source = []
    page = None
    source_map_lines = (PACKAGE/'source-map.md').read_text('utf-8').splitlines()
    for line in source_map_lines:
        m = re.match(r'### '+chapter+r' / (\w+) /', line)
        if m: page = m[1]
        m = re.match(r'\|'+chapter+r'-P(\d+)\|([^|]*)\|(.*)\|$', line)
        if m: source.append((int(m[1]), page, m[2], m[3]))
    day, entry, blockers = {'C13':(8,'Immediate','C-B01,C-B02,C-B05;C-G04'), 'C28':(23,'BeforeNight','C-B05,C-B06;C-G07'), 'C14':(9,'MorningBeforeMap','C-B03,C-B04,C-B05;B-real-completion;C-G02,C-G05,C-G07')}[chapter]
    # blockedBy is an explicit production dependency gate, not a fabricated story condition.
    add('Event', eventId=root, eventKind='world', actorId='world', setupKey=root, onceOnly=True, logTag=chapter,
        presentationClass='Conversation', presentationKey=root+'.conversation', presentationLogTag=chapter,
        eventProperty=f'entry={entry};blocksSave=true;blockedBy={blockers.replace(";", ",")}')
    add('SetupParams', setupKey=root, tplName='Tpl_ArrivalConversation', slotValues='opts=', skipAble=False, autoAdvance=False)
    add('PresentationIndex', presentationKey=root+'.conversation', presentationClass='Conversation', logTag=chapter)
    if chapter != 'C13': add('EventDate', eventId=root, day=day, timeSlot='morning' if chapter=='C14' else 'night', repeatRule='once', priority=0)
    def fact(key):
        if not any(r[1]==key for r in n['ConditionKey'].values):
            text('condition.'+key,key,chapter+' output declaration; no production completion while blocked')
            add('ConditionKey', key=key,nameText='condition.'+key,initialValue=False,onlySetTrue=True)
    fact('story.'+root+'.completed')
    for suffix,typ,target,value in [('day','BAR_BETWEEN','world.calendar',f'{day},{day}'),('incomplete','BOOL_EQ','story.'+root+'.completed','false'),('first_letter','BOOL_EQ','story.oct07.s2.completed','true')]:
        add('Condition',conditionId=root+'.'+suffix,ownerType='event',ownerId=root,groupIndex=1,condType=typ,target=target,value=value)
    # Upstream facts are not authored here: their producers do not exist yet. blockedBy keeps the root unavailable.
    page_rows = {}; counts = {}; choices = []; lines = []; seq = 0
    guards = {'reply1':'choice:'+root+'.c1.1','reply2':'choice:'+root+'.c1.2','reply3':'choice:'+root+'.c1.3',
        'comfort':'choice:'+root+'.c1.1','ask_father':'choice:'+root+'.c1.2','self':'choice:'+root+'.c1.3',
        'history':'dependency:C-B06.history-specialization','flashback':'story.revolution.1013.s2.completed',
        'cafe_plain':'anyChoice:'+root+'.c2.1;'+root+'.c2.2','cafe_strange':'choice:'+root+'.c2.3',
        'stance_neutral':'choice:'+root+'.c4.1','stance_rose':'choice:'+root+'.c4.2',
        'stance_revolution_a':'choice:'+root+'.c4.3','perception6':'choice:'+root+'.c4.3','stance_revolution_b':'choice:'+root+'.c4.3',
        'stance_church':'choice:'+root+'.c4.4','protect':'choice:'+root+'.c5.1','balance':'choice:'+root+'.c5.2','cooperate':'choice:'+root+'.c5.3',
        'reciprocity':'dependency:C-B03.C-B04.C-B05.reciprocity'}
    bg = 'oct07.cafe.day' if chapter=='C14' else 'oct07.cafe.night'
    def ensure(p):
        nonlocal seq
        if p not in page_rows:
            seq += 1
            step = 'PassiveCheck' if p in ['wisdom4','wisdom5','perception6'] else 'Dialogue'
            page_rows[p] = add('ConversationPage',pageId=root+'.'+p,presentationKey=root+'.conversation',seq=seq,pageCondKeys=guards.get(p),bgOverride=bg,stepType=step,checkId=root+'.'+p if step=='PassiveCheck' else None)
        return root+'.'+p
    emotes = dict(zip(['开心','平静','思索','迷惑','尴尬','悲伤','兴奋'],['happy','neutral','thinking','confused','embarrassed','sad','excited']))
    for number, p, mapping, original in source:
        if '选项；Choice=' in mapping:
            m=re.search(r'Choice=`([^`]+)`',mapping); cid=m[1]; group=cid.rsplit('.',1)[0]
            body=re.sub(r'^[①②③④]','',original).split('（勇气检定')[0]
            key=f'ncs.{root}.p{number:03}'
            text(key,body,f'{chapter}-P{number:03}: {original}')
            add('Choice',choiceId=cid,groupKey=group,textKey=key,seq=int(cid.rsplit('.',1)[1]))
            choices.append(cid); continue
        if 'ReadingLine=' in mapping:
            key=f'ncs.{root}.p{number:03}'; text(key,original,f'{chapter}-P{number:03}: {original}')
            add('PresentationIndex',presentationKey=root+'.letter',presentationClass='Reading',logTag=chapter)
            add('ReadingPage',pageId=root+'.letter.page',presentationKey=root+'.letter',seq=1)
            add('ReadingLine',lineId=root+'.letter.p017',pageId=root+'.letter.page',seq=1,mode='narration',textKey=key)
            continue
        if 'line/text后缀=' not in mapping: continue
        # Make source-stage boundaries explicit without adding dialogue text.
        if chapter=='C13' and number>=63: p='after_exit'
        if chapter=='C28' and number>=150: p='after_exit'
        if chapter=='C14' and number>=216: p='after_exit'
        pid=ensure(p)
        body=original; speaker=None
        if mapping.startswith('阿黛尔'): speaker='adele'; body=body.split('：',1)[1]
        elif mapping.startswith('玩家'): speaker='player'; body=body.split('：',1)[1]
        elif mapping.startswith('索雷尔'): speaker='sorel'; body=body.split('：',1)[1]
        parts=re.split(r'（#([^）]+)）',body)
        segments=[]
        if parts[0]: segments.append((None,parts[0]))
        for i in range(1,len(parts),2): segments.append((emotes[parts[i]],parts[i+1]))
        for index,(emote,body) in enumerate(segments):
            suffix=f'p{number:03}'+(chr(97+index) if len(segments)>1 else '')
            key='ncs.'+root+'.'+suffix
            text(key,body,f'{chapter}-P{number:03}: {original}')
            counts[p]=counts.get(p,0)+1
            display_name = None
            if chapter == 'C14' and speaker == 'sorel':
                display_name = 'character.sorel.anonymous2' if number == 26 else 'character.sorel.anonymous3' if number <= 45 else None
            add('ConversationLine',lineId=root+'.'+suffix,pageId=pid,seq=counts[p],speakerKey=speaker,side='right' if speaker in ['adele','sorel'] else 'none',textKey=key,emote=emote,speakerNameKey=display_name)
            lines.append(root+'.'+suffix)
    def options(p,g): n['ConversationPage'].cell(page_rows[p],7,root+'.'+g)
    if chapter=='C13': options('opening','c1')
    if chapter=='C28':
        # Two mutually exclusive attachments reuse the single authored choice group.
        options('family','c1'); options('flashback','c1')
        # Family choice must be conditional, therefore separate choice-only continuation is used below.
        n['ConversationPage'].cell(page_rows['family'],7).value=None
        n['ConversationPage'].cell(page_rows['flashback'],7).value=None
        # Insert at the exact point before branch replies; no duplicated P090/P098.
        insert_seq=n['ConversationPage'].cell(page_rows['comfort'],4).value
        for r in page_rows.values():
            cell=n['ConversationPage'].cell(r,4)
            if cell.value>=insert_seq: cell.value+=1
        add('ConversationPage',pageId=root+'.choice1',presentationKey=root+'.conversation',seq=insert_seq,optionGroupKey=root+'.c1',bgOverride=bg,stepType='Dialogue')
    if chapter=='C14':
        for p,g in [('opening','c1'),('visitor','c2'),('political_lead','c4'),('protection_lead','c5')]: options(p,g)
        for check,attr,bound in [('wisdom4','wisdom',4),('perception6','perception',6),('wisdom5','wisdom',5)]:
            add('AbilityCheck',checkId=root+'.'+check,mode='Inline',comment=chapter+' source Inline pure attribute; no result owner')
            add('AbilityCheckCondition',conditionId=root+'.'+check+'.1',checkId=root+'.'+check,attributeId=attr,threshold=bound,groupIndex=1,seq=1)
        # Letter steps are positioned between authored lines without a second body copy.
        for name,before,kind in [('envelope','before_letter','LetterReveal'),('read_letter','visitor','LetterRead')]:
            at=n['ConversationPage'].cell(page_rows[before],4).value
            for r in range(5,n['ConversationPage'].max_row+1):
                if n['ConversationPage'].cell(r,3).value==root+'.conversation' and n['ConversationPage'].cell(r,4).value>=at: n['ConversationPage'].cell(r,4).value+=1
            add('ConversationPage',pageId=root+'.'+name,presentationKey=root+'.conversation',seq=at,stepType=kind,readingKey=root+'.letter',bgOverride=bg)
        if not any(r[1]=='sorel' for r in n['Object'].values):
            text('character.sorel.name','索雷尔','C14 same object; portrait unavailable C-R04')
            text('character.sorel.anonymous2','？？','C14-P026 anonymous display')
            text('character.sorel.anonymous3','？？？','C14-P028-P045 anonymous display')
            add('Object',objId='sorel',objType='npc',nameText='character.sorel.name',tags='revolution,C-R04-missing-portrait')
    # The unlock reads a real AbilityReader value and shares its snapshot with selection.
    if chapter == 'C14':
        add('Condition', conditionId=root+'.c2.3.courage6', ownerType='choice', ownerId=root+'.c2.3', groupIndex=1,
            condType='ATTRIBUTE_GTE', target='courage', value='6')
    # No fake numeric key, zero result, or partial completion result is emitted for a blocked root.
    for i,p in enumerate(paths):
        tmp=backup/(p.stem+'.new.xlsx'); books[i].save(tmp)
        if hashlib.sha256(p.read_bytes()).hexdigest()!=hashes[i]: raise RuntimeError('Concurrent workbook edit: '+str(p))
        shutil.copyfile(tmp,p)
    report={'chapter':chapter,'root':root,'sourceParagraphs':len(source),'conversationLines':len(lines),'choices':choices,'blockedBy':blockers,'productionReleased':False,'numericResultsBound':False}
    (backup/'configuration.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),'utf-8')
    print(json.dumps(report,ensure_ascii=False))

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('chapter',choices=['C13','C28','C14']);run(p.parse_args().chapter)
