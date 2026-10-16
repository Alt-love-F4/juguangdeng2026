"""Migrate only E page sequences/guards and append authored checkpoint steps.
Uses same cooperative lock/CAS and preserves all unrelated cells.
"""
from copy import copy
from datetime import datetime
import json
import shutil
import openpyxl
from configure_church_long import ROOT, NARRATIVE, digest, signature, require, plan


def main():
    lock = ROOT / 'Library/Locus/tmp/narrative-source-write.lock'
    with lock.open('x', encoding='utf-8') as f:
        f.write('E checkpoint page migration\n')
    try:
        sha = digest(NARRATIVE)
        wb = openpyxl.load_workbook(NARRATIVE)
        old = {(s.title, c.coordinate): signature(c) for s in wb for row in s for c in row}
        ws = wb['ConversationPage']
        heads = {c.value: c.column for c in ws[1]}
        index = {ws.cell(r, heads['pageId']).value: r for r in range(5, ws.max_row + 1)}
        touched = set()
        changes = []
        for item in plan()[0]:
            if item['sheet'] != 'ConversationPage':
                continue
            values = item['values']
            key = values['pageId']
            existed = key in index
            row = index.get(key, ws.max_row + 1)
            if not existed:
                for col in range(1, ws.max_column + 1):
                    ws.cell(row, col)._style = copy(ws.cell(row - 1, col)._style)
            for field, value in values.items():
                cell = ws.cell(row, heads[field])
                if existed and cell.value != value:
                    require(field in ('seq', 'pageCondKeys'), f'Unexpected E mutation: {key}/{field}')
                if cell.value != value:
                    changes.append([key, field, cell.value, value])
                    touched.add((ws.title, cell.coordinate))
                    cell.value = value
        out = ROOT / 'Library/Locus/tmp/E-church' / (datetime.now().strftime('%Y%m%d-%H%M%S') + '-checkpoints')
        out.mkdir(parents=True)
        candidate = out / 'Narrative.candidate.xlsx'
        wb.save(candidate); wb.close()
        saved = openpyxl.load_workbook(candidate)
        for (sheet, coord), before in old.items():
            if (sheet, coord) not in touched:
                require(signature(saved[sheet][coord]) == before, f'Unrelated cell changed: {sheet}!{coord}')
        saved.close()
        require(digest(NARRATIVE) == sha, 'Concurrent mutation; candidate retained')
        shutil.copy2(NARRATIVE, out / 'Narrative.before.xlsx')
        require(digest(out / 'Narrative.before.xlsx') == sha and digest(NARRATIVE) == sha, 'Backup/CAS race')
        candidate.replace(NARRATIVE)
        manifest = {'before': sha, 'after': digest(NARRATIVE), 'changes': changes, 'evidence': str(out)}
        (out / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
        print(json.dumps(manifest, ensure_ascii=False))
    finally:
        lock.unlink()

if __name__ == '__main__':
    main()
