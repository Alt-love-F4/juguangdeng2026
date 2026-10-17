"""Add an isolated, source-ordered preview of all Rose 1011 display text."""
from copy import copy
from datetime import datetime
import hashlib
import shutil
from pathlib import Path

import openpyxl

import import_rose_text as source

ROOT = source.ROOT
NARRATIVE = ROOT / 'Assets/ConfigTables/Data/Narrative/Narrative.xlsx'
TEXT = ROOT / 'Assets/ConfigTables/Data/Shared/Text.xlsx'
initial_hashes = {path: hashlib.sha256(path.read_bytes()).hexdigest() for path in (NARRATIVE, TEXT)}
book = openpyxl.load_workbook(NARRATIVE)
text_book = openpyxl.load_workbook(TEXT)
entries = source.expected_rows('1011')
originals = {row[0]: row[4] for row in source.source_rows('1011')}


def upsert(workbook, sheet_name, values):
    ws = workbook[sheet_name]
    headers = [cell.value for cell in ws[1]]
    keyfield = headers[1]
    matches = [row for row in ws.iter_rows(min_row=5) if row[1].value == values[keyfield]]
    if len(matches) > 1:
        raise AssertionError(f'Duplicate row {sheet_name}/{values[keyfield]}')
    if matches:
        row_number = matches[0][1].row
        for key, value in values.items():
            assert ws.cell(row_number, headers.index(key) + 1).value == value, \
                f'Existing different row {sheet_name}/{values[keyfield]}/{key}'
        return
    row_number = ws.max_row + 1
    for column in range(1, ws.max_column + 1):
        ws.cell(row_number, column)._style = copy(ws.cell(5, column)._style)
    for key, value in values.items():
        ws.cell(row_number, headers.index(key) + 1, value)


preview_key = 'rose.1011.s1.fulltext.preview'
page_id = preview_key + '.page'
upsert(book, 'PresentationIndex', dict(
    presentationKey=preview_key, presentationClass='Conversation',
    logTag='Rose 1011 source-ordered text preview; isolated, no gameplay results'))
upsert(book, 'ConversationPage', dict(
    pageId=page_id, presentationKey=preview_key, seq=1, pageCondKeys=None, pageBarConds=None,
    optionGroupKey=None, bgOverride='cafe.day', bgmOverride=None, stepType=None,
    readingKey=None, checkId=None))

mixed_speakers = {
    52: {'a': 'sylvain', 'b': None, 'c': 'sylvain', 'd': 'adele', 'e': 'sylvain', 'f': None, 'g': 'sylvain'},
    85: {'a': 'player', 'b': 'player', 'c': 'player'},
    87: {'a': 'sylvain', 'b': 'sylvain'},
    110: {'a': 'sylvain', 'b': 'sylvain', 'c': 'sylvain'},
    116: {'a': 'sparrow', 'b': 'sparrow'},
    117: {'a': 'sylvain', 'b': 'sylvain', 'c': 'sylvain', 'd': 'sylvain'},
    122: {'a': 'sylvain', 'b': 'sylvain'},
    125: {'a': 'sylvain', 'b': 'sylvain'},
}

def base_speaker(raw):
    raw = raw.strip()
    if raw.startswith('{playername}'):
        return 'player'
    for prefix, speaker in [('阿黛尔', 'adele'), ('麻雀', 'sparrow'), ('西尔万', 'sylvain')]:
        if raw.startswith(prefix):
            return speaker
    return None

label_key = 'rose.1011.s1.preview.choice_label'
upsert(text_book, 'Text', dict(key=label_key, text='选项', comment='Isolated Rose 1011 full-text preview label'))
for seq, entry in enumerate(entries, 1):
    number = int(entry['source'].split('P')[-1])
    suffix = entry['key'].rsplit('.', 1)[-1]
    speaker = mixed_speakers.get(number, {}).get(suffix, base_speaker(originals[number]))
    is_choice = entry['kind'] == '选项'
    if is_choice:
        speaker = 'player'
    upsert(book, 'ConversationLine', dict(
        lineId=f'rose.1011.s1.preview.p{number:03}.{suffix}', pageId=page_id, seq=seq,
        speakerKey=speaker, side='none', textKey=entry['key'], voiceKey=None, emote=None,
        sfxKey=None, accentFxKey=None, accentTargetKey=None,
        speakerNameKey=label_key if is_choice else None))

stamp = datetime.now().strftime('%Y%m%d-%H%M%S')
out = ROOT / 'Library/Locus/tmp/rose-content' / f'{stamp}-1011-fulltext-preview'
out.mkdir(parents=True)
for path, name in ((NARRATIVE, 'Narrative'), (TEXT, 'Text')):
    shutil.copy2(path, out / f'{name}.before.xlsx')
(book).save(out / 'Narrative.candidate.xlsx')
text_book.save(out / 'Text.candidate.xlsx')
if any(hashlib.sha256(path.read_bytes()).hexdigest() != digest for path, digest in initial_hashes.items()):
    raise RuntimeError('Concurrent source workbook change; candidates preserved and sources not replaced')
(out / 'Narrative.candidate.xlsx').replace(NARRATIVE)
(out / 'Text.candidate.xlsx').replace(TEXT)
print(f'Added isolated Rose 1011 full-text preview: {len(entries)} ordered dialogue/choice lines.')
print(f'Preview key: conversation:{preview_key}; it adds no event, result, completion fact, or production behavior.')
print(f'Backup: {out}')
