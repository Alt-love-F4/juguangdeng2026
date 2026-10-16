"""Prepare ONLY blocked church.1020.s1 configuration from the current D source map.

Default: validate a read-only plan (no files saved). --apply: append with backups,
round-trip preservation checks and cooperative SHA-256 CAS. Run --apply only when
instructed by the shared-workbook owner. Official generation belongs to the parent.
No Action, EventDate, Condition, ConditionKey, Result or gameplay effects.
Shared tails use the public selectedCheck guard and terminal FadeOut steps.
Formal entry and settlement remain blocked by detention and map-range decisions.
Existing legacy dependency guards migrate only from their exact known values.
"""
from __future__ import annotations

import argparse
from collections import Counter
from copy import copy
from datetime import datetime
import json
import re
import shutil

import openpyxl

from import_church_text import ROOT, MAP, TEXT, SOURCES, digest, expected_rows

NARRATIVE = ROOT / 'Assets/ConfigTables/Data/Narrative/Narrative.xlsx'
EVENT = 'church.1020.s1'
PROPERTY = 'entry=MapAction;blocksSave=true;blockedBy=D-C04-detention,D-B04-map-range'
HEADERS = {
    'Event': 'eventId eventKind actorId targetId titleText contentKey setupType setupKey onceOnly logTag presentationClass presentationKey presentationLogTag eventProperty',
    'PresentationIndex': 'presentationKey presentationClass logTag',
    'SetupParams': 'setupKey tplName slotValues skipAble autoAdvance checkId',
    'ConversationPage': 'pageId presentationKey seq pageCondKeys pageBarConds optionGroupKey bgOverride bgmOverride stepType readingKey checkId',
    'ConversationLine': 'lineId pageId seq speakerKey side textKey voiceKey emote sfxKey accentFxKey accentTargetKey speakerNameKey stageLeftKey stageRightKey stageCenterKey',
    'Choice': 'choiceId groupKey textKey resultOwnerId seq checkId',
    'AbilityCheck': 'checkId successEventId failureEventId successResultOwnerId failureResultOwnerId comment mode',
    'AbilityCheckCondition': 'conditionId checkId attributeId threshold groupIndex seq',
    'Object': 'objId objType nameText iconPath portraitPath portraitIndoorPath avtarPath descText tags',
    'Text': 'key text comment',
}
EXPLANATIONS = {'language': 44, 'courage': 45, 'wisdom': 46, 'perception': 47}
PENDING = {
    'routing': 'selectedCheck reads only the selected explanation snapshot; not: shares it. Missing/ambiguous selection is an error. Complete root content can be previewed without settlement.',
    'settlement': 'No Result/ConditionKey: cleared, detention, burned/seized letters and root completion must join the eventual authoritative atomic settlement after terminal fade.',
    'entry': 'No Action/EventDate/Condition: parent must resolve real letter source, day15 eligibility, cafe mandatory-action policy, detention calendar and map-range blockers before formal release.',
    'art': 'Generic inspector and inspector group have blank art. Cafe uses oct07.cafe.day; church.detention.interior is registered as an explicit placeholder at P149.',
    'effects': 'P145/P153 use terminal FadeOut. P005 intrusion audio, P011 group silhouettes, P042 suggested tension sound and P143 burning-letter image remain unbound; no guessed audio/effects.',
}


def require(ok, message):
    if not ok:
        raise RuntimeError(message)


def signature(cell):
    return (cell.value, cell.data_type, tuple(cell._style) if cell.has_style else (0,) * 9,
            (cell.comment.text, cell.comment.author) if cell.comment else None, copy(cell.hyperlink))


def guards(short):
    if short == 'cooperate':
        return 'anyChoice:' + ';'.join(EVENT + f'.p{n:03}' for n in (20, 21))
    if short.startswith('defy_'):
        return 'choice:' + EVENT + '.p022|' + ('not:' if short.endswith('_fail') else '') + 'check:' + EVENT + '.defy'
    if short in ('cleared', 'detained'):
        return ('not:' if short == 'detained' else '') + 'selectedCheck:' + ';'.join(
            EVENT + f'.p{n:03}=' + EVENT + '.explain.' + attr for attr, n in EXPLANATIONS.items())
    if '_' in short:
        attribute, outcome = short.split('_')
        if attribute in EXPLANATIONS:
            result = 'choice:' + EVENT + f'.p{EXPLANATIONS[attribute]:03}'
            if outcome != 'intro':
                result += '|' + ('not:' if outcome == 'fail' else '') + 'check:' + EVENT + '.explain.' + attribute
            return result
    require(short in ('raid', 'letter_found'), 'Unknown page: ' + short)
    return None


def plan():
    entries = expected_rows('1020.s1')  # Original DOCX hash and all 153 exact map paragraphs.
    body = [e for e in entries if e['kind'].startswith('正文')]
    choices = [e for e in entries if e['kind'].startswith('选项')]
    require(len(body) == 117 and len(choices) == 7, 'Source coverage changed')
    rows = []

    def add(sheet, **values):
        rows.append(dict(sheet=sheet, values=values))

    add('Event', eventId=EVENT, eventKind='world', actorId='world', setupKey=EVENT,
        onceOnly=True, logTag='D1020', presentationClass='Conversation',
        presentationKey=EVENT + '.conversation', presentationLogTag='D1020', eventProperty=PROPERTY)
    add('PresentationIndex', presentationKey=EVENT + '.conversation', presentationClass='Conversation', logTag='D1020')
    add('SetupParams', setupKey=EVENT, tplName='Tpl_ArrivalConversation', slotValues='opts=', skipAble=False, autoAdvance=False)
    # Only source-provided identity labels; these are not narrative additions.
    for key, label in [('church.inspector.name', '监察官'), ('church.inspectors.name', '众监察员'),
                       (EVENT + '.inspector.uncertain.name', '疑似领头的人')]:
        add('Text', key=key, text=label, comment='D1020 source speaker identity/alias; no narrative addition')
    for obj in ('church.inspector', 'church.inspectors'):
        add('Object', objId=obj, objType='npc', nameText=obj + '.name')
    for short, attribute, threshold in [('defy', 'courage', 4)] + [('explain.' + a, a, 6) for a in EXPLANATIONS]:
        check = EVENT + '.' + short
        add('AbilityCheck', checkId=check, mode='Inline', comment='D1020 selected path only; cached snapshot; no branch event or result owner')
        add('AbilityCheckCondition', conditionId=check + '.1', checkId=check,
            attributeId=attribute, threshold=threshold, groupIndex=1, seq=1)
    # Dict insertion order preserves current source-map page order, never hardcodes body text.
    pages = list(dict.fromkeys(e['page'] for e in body))
    require(len(pages) == 18, 'Source pages changed')
    page_seq = 0
    for page in pages:
        short = page.removeprefix(EVENT + '.')
        content = [e for e in body if e['page'] == page]
        # P149 explicitly enters detention; give that source anchor its own background.
        segments = [content] if short != 'detained' else [content[:2], content[2:]]
        for segment_index, segment in enumerate(segments):
            page_seq += 1
            actual_page = page if segment_index == 0 else page + '.interior'
            group = 'cooperate' if short == 'raid' else 'explain' if short == 'letter_found' else None
            add('ConversationPage', pageId=actual_page, presentationKey=EVENT + '.conversation', seq=page_seq,
                pageCondKeys=guards(short), optionGroupKey=EVENT + '.' + group + '.choices' if group else None,
                bgOverride='church.detention.interior' if segment_index else 'oct07.cafe.day', stepType='Dialogue')
            for seq, entry in enumerate(segment, 1):
                original = entry['original']
                speaker = ('adele' if original.startswith(('阿黛尔·米尼耶：', '阿黛尔：')) else
                           'church.inspector' if original.startswith(('监察官：', '疑似领头的人：')) else
                           'church.inspectors' if original.startswith('众监察员：') else
                           'player' if re.match(r'^\{player(?:name|nanem|nmae)\}[：；]', original) else None)
                add('ConversationLine', lineId=entry['key'].removesuffix('.text'), pageId=actual_page,
                    seq=seq, speakerKey=speaker, side='none', textKey=entry['key'],
                    speakerNameKey=EVENT + '.inspector.uncertain.name' if original.startswith('疑似领头的人：') else None)
    # Both mutually exclusive tails converge on the single final fade required by the loader.
    add('ConversationPage', pageId=EVENT + '.fade', presentationKey=EVENT + '.conversation',
        seq=page_seq + 1, stepType='FadeOut')
    for entry in choices:
        num = int(entry['source'].rsplit('P', 1)[1])
        group, sequence = ('cooperate', (20, 21, 22).index(num) + 1) if num < 44 else ('explain', num - 43)
        add('Choice', choiceId=entry['key'].removesuffix('.text'), groupKey=EVENT + '.' + group + '.choices',
            textKey=entry['key'], seq=sequence)
    require(Counter(r['sheet'] for r in rows)['ConversationLine'] == 117, 'Line coverage')
    require(len({r['values']['lineId'] for r in rows if r['sheet'] == 'ConversationLine'}) == 117, 'Duplicate source lines')
    return rows, entries


def prepare(books, rows, entries, migrations=None):
    text_rows = {}
    for row in books[TEXT]['Text'].iter_rows(min_row=5, values_only=True):
        if str(row[1]).startswith(EVENT + '.p'):
            require(row[1] not in text_rows, 'Duplicate existing source Text: ' + str(row[1]))
            text_rows[row[1]] = row[2]
    require(text_rows == {e['key']: e['text'] for e in entries}, 'Import source Text first: existing Text/source mismatch')
    existing_objects = {r[1] for r in books[NARRATIVE]['Object'].iter_rows(min_row=5, values_only=True)}
    require({'player', 'adele'} <= existing_objects, 'Missing reusable speakers')
    added, locations = Counter(), []
    for item in rows:
        sheet, values = item['sheet'], item['values']
        path = TEXT if sheet == 'Text' else NARRATIVE
        ws = books[path][sheet]
        require([c.value for c in ws[1]] == ['##var'] + HEADERS[sheet].split(), 'Schema changed: ' + sheet)
        heads = {c.value: c.column for c in ws[1]}
        key = next(iter(values.values()))
        matches = [r for r in range(5, ws.max_row + 1) if ws.cell(r, 2).value == key]
        require(len(matches) <= 1, 'Duplicate key: ' + str(key))
        if matches:
            row = matches[0]
            for field in HEADERS[sheet].split():
                cell = ws.cell(row, heads[field])
                desired = values.get(field)
                if cell.value != desired:
                    legacy = None
                    short = str(key).removeprefix(EVENT + '.')
                    if sheet == 'ConversationPage' and short in ('cleared', 'detained', 'detained.interior'):
                        if field == 'pageCondKeys':
                            legacy = 'dependency:' + EVENT + '.selected_explanation_' + ('pass' if short == 'cleared' else 'fail')
                        elif field == 'seq' and short != 'cleared':
                            # Interim branch-fade candidate (never a valid runtime delivery).
                            legacy = 19 if short == 'detained' else 20
                    require(legacy is not None and cell.value == legacy,
                            f'Conflicting existing row: {sheet}/{key}/{field}')
                    if migrations is not None:
                        migrations.append(dict(sheet=sheet, cell=cell.coordinate, key=key, field=field,
                                               before=cell.value, after=desired))
                    cell.value = desired
        else:
            row = ws.max_row + 1
            for col in range(1, ws.max_column + 1):
                ws.cell(row, col)._style = copy(ws.cell(row - 1, col)._style)
            for field, value in values.items():
                ws.cell(row, heads[field], value)
            added[sheet] += 1
        locations.append(dict(sheet=sheet, key=key, row=row))
    for sheet in ('Action', 'EventDate', 'Condition', 'ConditionKey', 'Result', 'Ending'):
        require(not any(EVENT in str(c.value) for row in books[NARRATIVE][sheet].iter_rows(min_row=5) for c in row), 'Unexpected formal integration: ' + sheet)
    return dict(added), locations


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--apply', action='store_true', help='Owner-authorized CAS append; generation is separate')
    args = parser.parse_args()
    initial = {p: digest(p) for p in (NARRATIVE, TEXT, MAP)}
    rows, entries = plan()
    books = {p: openpyxl.load_workbook(p) for p in (NARRATIVE, TEXT)}
    try:
        before = {p: {(s.title, c.coordinate): signature(c) for s in wb for row in s for c in row} for p, wb in books.items()}
        migrations = []
        added, locations = prepare(books, rows, entries, migrations)
        report = dict(event=EVENT, eventProperty=PROPERTY, lines=117, choices=7, checks=5,
                      pages=sum(r['sheet'] == 'ConversationPage' for r in rows), added=added,
                      pending=PENDING, migrations=migrations, generation_run=False, writes=False)
        for p, sha in initial.items():
            require(digest(p) == sha, 'Concurrent change during validation: ' + str(p))
        if not args.apply or not (added or migrations):
            print(json.dumps(report, ensure_ascii=False, indent=2))
            return
        out = ROOT / 'Library/Locus/tmp/church-content' / (datetime.now().strftime('%Y%m%d-%H%M%S-%f') + '-1020-config')
        out.mkdir(parents=True)
        candidates = {}
        for path, wb in books.items():
            backup = out / (path.stem + '.before.xlsx')
            shutil.copy2(path, backup)
            require(digest(backup) == initial[path], 'Backup race: ' + str(path))
            candidate = out / (path.stem + '.candidate.xlsx')
            wb.save(candidate)
            candidates[path] = candidate
        saved = {p: openpyxl.load_workbook(c) for p, c in candidates.items()}
        try:
            for path in books:
                changed = {(m['sheet'], m['cell']): m for m in migrations} if path == NARRATIVE else {}
                for (sheet, coordinate), value in before[path].items():
                    actual = signature(saved[path][sheet][coordinate])
                    if (sheet, coordinate) in changed:
                        m = changed[(sheet, coordinate)]
                        require(value[0] == m['before'] and actual[0] == m['after'] and actual[1:] == value[1:],
                                f'Migration mismatch: {path}/{sheet}/{coordinate}')
                    else:
                        require(actual == value, f'Existing cell changed: {path}/{sheet}/{coordinate}')
            check_added, check_locations = prepare(saved, rows, entries)
            require(not check_added and check_locations == locations, 'Candidate round-trip mismatch')
        finally:
            for wb in saved.values():
                wb.close()
        report.update(locations=locations, entries=entries, before_sha256={str(p): h for p, h in initial.items()},
                      source_sha256=SOURCES['1020.s1'][2], source_paragraphs=153,
                      preserved={str(p): len(v) for p, v in before.items()}, replaced=[])
        journal = out / 'manifest.json'
        journal.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        for path, sha in initial.items():
            require(digest(path) == sha, 'Concurrent change; candidates retained: ' + str(path))
        # Cooperative CAS, not an OS-level multi-file transaction. Backups and
        # per-file journal entries permit recovery if a replace fails midway.
        for path, candidate in candidates.items():
            require(digest(path) == initial[path], 'Concurrent change before replace: ' + str(path))
            candidate.replace(path)
            report['replaced'].append(str(path))
            report['writes'] = True
            journal.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        report['after_sha256'] = {str(p): digest(p) for p in books}
        journal.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        print(json.dumps(dict(added=added, manifest=str(journal), generation_run=False), ensure_ascii=False))
    finally:
        for wb in books.values():
            wb.close()


if __name__ == '__main__':
    main()
