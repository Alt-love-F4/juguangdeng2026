"""Apply only the user-approved neutral wallet assessment to the CURRENT source XLSX.
Preserves other in-progress story/presentation edits. Run GenerateAll.bat afterwards.
"""
from pathlib import Path
import argparse
import sys
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(Path(__file__).parent / "Content"))
import neutral


def apply(path):
    data = neutral.build(ROOT)
    wb = openpyxl.load_workbook(path)
    root = "neutral.route.1129.s1"
    ordinary = "neutral.route.1130.s2"

    def records(name):
        sheet = wb[name]
        headers = [c.value for c in sheet[1]]
        return sheet, headers, [dict(zip(headers, values)) for values in sheet.iter_rows(min_row=5, values_only=True) if values[1] is not None]

    def save(name, headers, rows):
        sheet = wb[name]
        for i, row in enumerate(rows, 5):
            for j, key in enumerate(headers, 1):
                sheet.cell(i, j).value = row.get(key)
        for i in range(len(rows) + 5, sheet.max_row + 1):
            for cell in sheet[i]:
                cell.value = None

    _, h, rows = records("Event")
    event = next(r for r in rows if r["eventId"] == root)
    event["eventProperty"] = ";".join(s for s in event["eventProperty"].split(";") if not s.startswith("blockedBy="))
    save("Event", h, rows)

    _, h, rows = records("ConversationPage")
    low_id = root + ".money_below_target"
    target_pages = [r for r in rows if r["presentationKey"] == root and r["pageId"] != low_id]
    target_pages.sort(key=lambda r: r["seq"])
    enough = f"bar:cafe.money:ge:{neutral.BUSINESS_TYCOON_MONEY}"
    for seq, row in enumerate(target_pages, 2):
        row["seq"] = seq
        # Do not gate the final unconditional End.
        if row["stepType"] != "End":
            guards = [s for s in str(row.get("pageCondKeys") or "").split("|") if s and s != enough]
            row["pageCondKeys"] = "|".join(guards + [enough])
    rows = [r for r in rows if r["pageId"] != low_id]
    rows.append(next(r for r in data["ConversationPage"] if r["pageId"] == low_id))
    save("ConversationPage", h, rows)

    _, h, rows = records("Result")
    result_ids = {root + ".result.committee", root + ".result.money_below_target"}
    rows = [r for r in rows if r["resultId"] not in result_ids]
    rows += [r for r in data["Result"] if r["resultId"] in result_ids]
    save("Result", h, rows)

    _, h, rows = records("Condition")
    def affected(r):
        return ((r["ownerType"] == "event" and r["ownerId"] == root)
                or (r["ownerType"] == "result" and r["ownerId"] in result_ids)
                or (r["ownerType"] == "event" and r["ownerId"] == ordinary and r["groupIndex"] == 3))
    rows = [r for r in rows if not affected(r)]
    incoming = [r for r in data["Condition"] if affected(r)]
    # Stable names for patch-owned conditions avoid collisions with retained source rows.
    for i, row in enumerate(incoming):
        row = dict(row, conditionId=f"neutral.wallet.assessment.{i:03}")
        rows.append(row)
    save("Condition", h, rows)
    wb.save(path)
    wb.close()
    print(f"Applied wallet >= {neutral.BUSINESS_TYCOON_MONEY} assessment to {path}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=ROOT / "Assets/ConfigTables/Data/Narrative/Narrative.xlsx")
    apply(parser.parse_args().workbook)
