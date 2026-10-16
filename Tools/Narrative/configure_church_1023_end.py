"""Append ONLY blocked church.1023.s2 content; no release producer or Ending.
Default is a read-only plan; --apply uses incremental backup and cooperative CAS.
No generation, Editor operations, Text writes or asset changes.
"""
from __future__ import annotations

import argparse
from collections import Counter
from copy import copy
from datetime import datetime
import json
from pathlib import Path
import shutil

import openpyxl

from import_church_text import ROOT, MAP, TEXT, digest, expected_rows, SOURCES

P = ROOT / 'Assets/ConfigTables/Data/Narrative/Narrative.xlsx'
EVENT = 'church.1023.s2'
PROPERTY = 'entry=Immediate;blocksSave=true;blockedBy=D-B01-release-failure-unavailable,D-C05-ending'
HEADERS = {
    'Event': 'eventId eventKind actorId targetId titleText contentKey setupType setupKey onceOnly logTag presentationClass presentationKey presentationLogTag eventProperty',
    'PresentationIndex': 'presentationKey presentationClass logTag',
    'SetupParams': 'setupKey tplName slotValues skipAble autoAdvance checkId',
    'ConversationPage': 'pageId presentationKey seq pageCondKeys pageBarConds optionGroupKey bgOverride bgmOverride stepType readingKey checkId',
    'ConversationLine': 'lineId pageId seq speakerKey side textKey voiceKey emote sfxKey accentFxKey accentTargetKey speakerNameKey',
}


def signature(cell):
    return (cell.value, cell.data_type, tuple(cell._style) if cell.has_style else (0,) * 9,
            (cell.comment.text, cell.comment.author) if cell.comment else None,
            copy(cell.hyperlink))


def plan():
    entries = expected_rows('1023.s2')  # All 34 DOCX paragraphs + hash + exact map text.
    assert len(entries) == 25 and all(e['kind'].startswith('正文') for e in entries)
    text = openpyxl.load_workbook(TEXT, read_only=True)
    try:
        actual = {}
        for row in text['Text'].iter_rows(min_row=5, values_only=True):
            if str(row[1]).startswith(EVENT + '.'):
                assert row[1] not in actual, row[1]
                actual[row[1]] = row[2]
        assert actual == {e['key']: e['text'] for e in entries}, 'Existing Text/source mismatch'
    finally:
        text.close()
    rows = []
    def add(sheet, **values):
        rows.append(dict(sheet=sheet, values=values))
    add('Event', eventId=EVENT, eventKind='world', actorId='world', setupKey=EVENT,
        onceOnly=True, logTag='D1023S2', presentationClass='Conversation',
        presentationKey=EVENT + '.conversation', presentationLogTag='D1023S2', eventProperty=PROPERTY)
    add('PresentationIndex', presentationKey=EVENT + '.conversation', presentationClass='Conversation', logTag='D1023S2')
    add('SetupParams', setupKey=EVENT, tplName='Tpl_ArrivalConversation', slotValues='opts=', skipAble=False, autoAdvance=False)
    for seq, short in enumerate(('prison', 'execution'), 1):
        page = EVENT + '.' + short
        add('ConversationPage', pageId=page, presentationKey=EVENT + '.conversation', seq=seq,
            bgOverride='church.detention.interior' if short == 'prison' else 'church.execution.black', stepType='Dialogue')
        content = [e for e in entries if e['page'] == page]
        assert len(content) == (6 if short == 'prison' else 19)
        for line_seq, entry in enumerate(content, 1):
            add('ConversationLine', lineId=entry['key'].removesuffix('.text'), pageId=page,
                seq=line_seq, side='none', textKey=entry['key'])
    assert len(rows) == 30
    return rows, entries


def prepare(wb, rows):
    locations, added = [], Counter()
    for item in rows:
        sheet, values = item['sheet'], item['values']
        ws = wb[sheet]
        assert [c.value for c in ws[1]] == ['##var'] + HEADERS[sheet].split(), sheet
        heads = {c.value: c.column for c in ws[1]}
        key = next(iter(values.values()))
        matches = [r for r in range(5, ws.max_row + 1) if ws.cell(r, 2).value == key]
        assert len(matches) <= 1, (sheet, key)
        if matches:
            r = matches[0]
            for field in HEADERS[sheet].split():
                assert ws.cell(r, heads[field]).value == values.get(field), (sheet, key, field)
        else:
            r = ws.max_row + 1
            for col in range(1, ws.max_column + 1):
                ws.cell(r, col)._style = copy(ws.cell(r - 1, col)._style)
            for field, value in values.items():
                ws.cell(r, heads[field], value)
            added[sheet] += 1
        locations.append(dict(sheet=sheet, key=key, row=r))
    # This isolated preparation must never carry outcomes or formal eligibility.
    for sheet in ('EventDate', 'Ending', 'Result', 'ConditionKey', 'Condition', 'Action', 'Choice'):
        assert not any(EVENT in str(c.value) for row in wb[sheet].iter_rows(min_row=5) for c in row), sheet
    return dict(added), locations


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--apply', action='store_true')
    args = parser.parse_args()
    initial = {p: digest(p) for p in (P, TEXT, MAP)}
    rows, entries = plan()
    wb = openpyxl.load_workbook(P)
    try:
        before = {(s.title, c.coordinate): signature(c) for s in wb for row in s for c in row}
        added, locations = prepare(wb, rows)
        if not args.apply or not added:
            print(json.dumps(dict(writes=False, added=added, locations=locations), ensure_ascii=False))
            return
        out = ROOT / 'Library/Locus/tmp/church-content' / (datetime.now().strftime('%Y%m%d-%H%M%S-%f') + '-1023-end-config')
        out.mkdir(parents=True)
        assert digest(P) == initial[P], 'Concurrent Narrative change before backup'
        backup = out / 'Narrative.before.xlsx'
        shutil.copy2(P, backup)
        assert digest(backup) == initial[P], 'Backup race'
        candidate = out / 'Narrative.candidate.xlsx'
        wb.save(candidate)
        saved = openpyxl.load_workbook(candidate)
        try:
            for (sheet, coordinate), value in before.items():
                assert signature(saved[sheet][coordinate]) == value, (sheet, coordinate)
            check_added, check_locations = prepare(saved, rows)
            assert not check_added and check_locations == locations
        finally:
            saved.close()
        manifest = dict(time=datetime.now().astimezone().isoformat(), event=EVENT,
            scope='Blocked isolated content only; authoritative upstream release.failed not produced; no formal eligibility or settlement.',
            eventProperty=PROPERTY, added=added, locations=locations, entries=entries,
            source_paragraphs=34, source_sha256=SOURCES['1023.s2'][2], source_map_sha256=initial[MAP],
            preserved_existing_cells=len(before), before_sha256={str(p): h for p, h in initial.items()},
            pending={'registry': {'church.detention.interior': 'neutral prison placeholder; parent owns asset registration',
                                  'church.execution.black': 'black placeholder; parent owns asset registration'},
                     'P011-P012': 'fade to black not configured', 'P015': 'door/steps/indistinct speech audio not configured',
                     'P033': 'single red frame then black not configured; no completion/Ending until supported'},
            omitted_sheets=['EventDate', 'Ending', 'Result', 'ConditionKey', 'Condition', 'Action', 'Choice'],
            generation_run=False, editor_operations=False, replaced=False)
        journal = out / 'manifest.json'
        journal.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
        for path, sha in initial.items():
            assert digest(path) == sha, f'Concurrent change: {path}; candidate retained at {out}'
        candidate.replace(P)
        manifest.update(replaced=True, after_sha256=digest(P), text_unchanged=digest(TEXT) == initial[TEXT])
        journal.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
        print(json.dumps(dict(added=added, manifest=str(journal), preserved_existing_cells=len(before), after_sha256=manifest['after_sha256']), ensure_ascii=False))
    finally:
        wb.close()


if __name__ == '__main__':
    main()
