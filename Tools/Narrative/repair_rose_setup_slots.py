"""Migrate only B-owned empty opts slots after validator inspection; no content changes."""
from datetime import datetime
import shutil
import openpyxl
import import_rose_text as source
from rose_write_guard import source_write_guard


def main():
    path=source.ROOT/'Assets/ConfigTables/Data/Narrative/Narrative.xlsx'
    with source_write_guard('repair B optional opts slots'):
        initial=source.digest(path)
        wb=openpyxl.load_workbook(path)
        try:
            ws=wb['SetupParams']; h={c.value:c.column for c in ws[1]}
            changed=[]
            for row in range(5,ws.max_row+1):
                key=ws.cell(row,h['setupKey']).value
                if key not in ['rose.1009.s1','rose.1011.s1','rose.1012.s1','rose.1016.s1']: continue
                assert ws.cell(row,h['tplName']).value=='Tpl_ArrivalConversation'
                cell=ws.cell(row,h['slotValues'])
                assert cell.value in (None,'opts='), f'Unexpected authored slot: {key}'
                if cell.value is None: cell.value='opts='; changed.append(key)
            if not changed: print('No migration needed'); return
            out=source.ROOT/'Library/Locus/tmp/rose-content'/(datetime.now().strftime('%Y%m%d-%H%M%S-%f')+'-opts-migration')
            out.mkdir(parents=True)
            assert source.digest(path)==initial
            shutil.copy2(path,out/'Narrative.before.xlsx')
            candidate=out/'Narrative.candidate.xlsx';wb.save(candidate)
            assert source.digest(path)==initial
            candidate.replace(path)
            print('Explicit optional opts= assigned:',changed,'backup:',out)
        finally: wb.close()


if __name__=='__main__': main()
