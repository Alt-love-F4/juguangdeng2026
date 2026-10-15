"""Prepare or append E-church Text rows; never creates a playable event.
Default is read-only audit. --apply appends to the latest Text.xlsx using a
short lock and hash recheck. Run GenerateAll.bat separately afterwards.
"""
from __future__ import annotations
import argparse
from collections import Counter
from copy import copy
from datetime import datetime
import hashlib
import json
from pathlib import Path
import re
import shutil
import xml.etree.ElementTree as ET
import zipfile

import openpyxl
from openpyxl.comments import Comment

ROOT = Path(__file__).resolve().parents[2]
MAP = ROOT / 'Docs/superpowers/specs/2026-10-16-content-batches/E-church/source-map.md'
TEXT = ROOT / 'Assets/ConfigTables/Data/Shared/Text.xlsx'
SOURCE = Path('D:/桌面/文案/王庭教会')
SOURCES = {
    'MA70.11.04-S1.docx': ('E04', 473, '85060b970849e5e76c616504ba6621867603b678e1f919a917e8fd8cb0f6b955'),
    '【施工中】MA70.11.05-S1.docx': ('E05A', 1, 'e4b24b8a56befdce94e316a91f5bc04dcc109bcf7fa285fd6d44741025f1cb29'),
    '【施工中】MA70.11.05-S2.docx': ('E05B', 1, '5f6025cac385e3cad5a9c97c30871efbb09f8d3d74a10f7fbfc0f9e31e27032b'),
}

def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def paragraphs(path):
    with zipfile.ZipFile(path) as archive:
        xml = ET.fromstring(archive.read('word/document.xml'))
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    return [''.join(n.text or '' for n in p.findall('.//w:t', ns))
            for p in xml.findall('./w:body/w:p', ns)
            if ''.join(n.text or '' for n in p.findall('.//w:t', ns)).strip()]


def expected_rows():
    sections = {prefix: [] for prefix, _, _ in SOURCES.values()}
    page = None
    for line_number, line in enumerate(MAP.read_text(encoding='utf-8').splitlines(), 1):
        if line.startswith('拟 page/步骤：'):
            page = re.search(r'`([^`]+)`', line).group(1)
        if not re.match(r'^\| E(?:04|05A|05B)-P\d{3}（\d+）', line):
            continue
        cells = [s.strip() for s in re.split(r'(?<!\\)\|', line)[1:-1]]
        prefix, number = re.search(r'(E04|E05A|E05B)-P(\d+)', cells[0]).groups()
        raw = cells[2].replace(r'\|', '|')
        sections[prefix].append({'number': int(number), 'source': f'{prefix}-P{number}',
                                 'meta': cells[1], 'original': raw, 'page': page,
                                 'source_map_line': line_number})
    for filename, (prefix, count, expected_hash) in SOURCES.items():
        path = SOURCE / filename
        assert digest(path) == expected_hash, f'Changed original: {path}'
        rows = sections[prefix]
        assert [r['number'] for r in rows] == list(range(1, count + 1))
        assert paragraphs(path) == [r['original'] for r in rows], f'DOCX mismatch: {filename}'
    counts = Counter(r['meta'].split('；')[0] for r in sections['E04'])
    assert counts == {'标题': 1, '控制': 26, '演出': 33, '正文': 405, '选项': 8}, counts
    entries = []
    for row in sections['E04']:
        if not row['meta'].startswith(('正文；', '选项；')):
            continue
        text = row['original']
        choice = row['meta'].startswith('选项；')
        if choice:
            text = re.sub(r'^[①②③]', '', text)
        else:
            speaker = re.search(r'；speaker=([^；]+)', row['meta'])
            if speaker:
                text = re.sub(r'^[^：:]+[：:]\s*', '', text, count=1)
        text = re.sub(r'（#[^）]*）', '', text)
        # Only explicit author/reward metadata from the source map is removed.
        text = re.sub(r'【(?:数值未定|金钱\+(?:1|xx)|目前不太清楚这个国家的消费水平来着)】', '', text)
        text = text.replace('{playernam}', '{playername}')
        key = f'church.1104.s1.p{row["number"]:03}'
        assert key in row['meta'] and text and '【' not in text and '（#' not in text
        entries.append(dict(row, key=key, text=text, kind='选项' if choice else '正文'))
    assert len(entries) == 413 and len({e['key'] for e in entries}) == 413
    assert sum('xx' in e['original'] for e in entries) == 15
    assert sum('xx' in e['text'] for e in entries) == 14
    return entries, dict(counts)


def signature(cell):
    return cell.value, cell.data_type, copy(cell._style), cell.comment.text if cell.comment else None


def stage(apply=False):
    entries, counts = expected_rows()
    if not apply:
        print(json.dumps({'source_counts': counts, 'verified_docx_paragraphs': 475,
                          'text_entries': len(entries), 'unresolved_xx_entries': 14,
                          'scope': 'Source audit only; no files written.'}, ensure_ascii=False))
        return
    evidence = ROOT / 'Library/Locus/tmp/E-church'
    evidence.mkdir(parents=True, exist_ok=True)
    # Cooperative writers may use this same lock; the CAS check also detects
    # non-cooperating writers before replacement. Do not steal stale locks.
    lock = ROOT / 'Library/Locus/tmp/narrative-source-write.lock'
    with lock.open('x', encoding='utf-8') as handle:
        handle.write('E-church Text append\n')
    try:
        initial_hash = digest(TEXT)
        wb = openpyxl.load_workbook(TEXT)
        ws = wb['Text']
        assert [ws.cell(1, c).value for c in range(1, 5)] == ['##var', 'key', 'text', 'comment']
        before = {(s.title, c.coordinate): signature(c) for s in wb for row in s for c in row}
        keys = {}
        for row in ws.iter_rows(min_row=5):
            key = row[1].value
            if key:
                assert key not in keys, f'Duplicate key: {key}'
                keys[key] = row[1].row
        assert not {k for k in keys if str(k).startswith('church.1104.s1.')} - {e['key'] for e in entries}
        added, mapping = [], []
        for entry in entries:
            key = entry['key']
            if key in keys:
                index = keys[key]
                assert ws.cell(index, 3).value == entry['text'], f'Existing E text differs: {key}'
            else:
                index = ws.max_row + 1
                for col in range(1, 5):
                    ws.cell(index, col)._style = copy(ws.cell(599, col)._style)
                ws.cell(index, 2, key)
                ws.cell(index, 3, entry['text'])
                ws.cell(index, 4, f'{entry["source"]}; {entry["page"]}; Text staging only; root disabled')
                ws.cell(index, 3).comment = Comment(f'{SOURCE / "MA70.11.04-S1.docx"}\n{entry["source"]}\n{entry["original"]}', 'E-church source')
                added.append(index)
            mapping.append({'key': key, 'row': index, 'source': entry['source'],
                            'source_map_line': entry['source_map_line'], 'page': entry['page']})
        out = evidence / datetime.now().strftime('%Y%m%d-%H%M%S-%f')
        out.mkdir()
        shutil.copy2(TEXT, out / 'Text.before.xlsx')
        candidate = out / 'Text.candidate.xlsx'
        wb.save(candidate)
        saved = openpyxl.load_workbook(candidate)
        for (sheet, coordinate), old in before.items():
            assert signature(saved[sheet][coordinate]) == old, f'Existing cell changed: {sheet}!{coordinate}'
        for entry, location in zip(entries, mapping):
            assert saved['Text'].cell(location['row'], 3).value == entry['text']
        assert digest(TEXT) == initial_hash, f'Concurrent Text change; candidate kept at {candidate}; rerun with latest workbook'
        if added:
            candidate.replace(TEXT)
        manifest = {'time': datetime.now().astimezone().isoformat(), 'added_rows': added,
                    'source_counts': counts, 'preserved_existing_cells': len(before),
                    'before_sha256': initial_hash, 'after_sha256': digest(TEXT), 'mapping': mapping,
                    'scope': 'Text only; no Event/Date/Choice/Result/complete fact or 11.05 text created.'}
        (out / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
        print(json.dumps({k: v for k, v in manifest.items() if k not in ('mapping', 'added_rows')}, ensure_ascii=False))
        print('Rows:', min(added) if added else 'existing', max(added) if added else 'existing', 'count=', len(added))
        print('Evidence:', out)
    finally:
        lock.unlink()


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--apply', action='store_true')
    stage(parser.parse_args().apply)
