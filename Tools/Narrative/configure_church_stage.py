"""Append optional stage-slot columns and apply source-anchored E stage commands.
Empty keeps a slot; '-' clears; Object key places. Official generation is separate.
"""
from copy import copy
from datetime import datetime
import json
from pathlib import Path
import shutil
import openpyxl
from configure_church_long import ROOT, NARRATIVE, digest, signature, require


def main():
    base = ROOT / 'Library/Locus/tmp'
    lock = base / 'narrative-source-write.lock'
    with lock.open('x', encoding='utf-8') as f:
        f.write('E stage columns and source commands\n')
    try:
        sha = digest(NARRATIVE)
        wb = openpyxl.load_workbook(NARRATIVE)
        ws = wb['ConversationLine']
        old = {(s.title, c.coordinate): signature(c) for s in wb for row in s for c in row}
        headers = {c.value: c.column for c in ws[1]}
        for field in ('stageLeftKey', 'stageRightKey', 'stageCenterKey'):
            if field not in headers:
                col = ws.max_column + 1
                headers[field] = col
                for r, value in enumerate((field, 'string', 'Empty=keep; -=clear; Object key=place', None), 1):
                    ws.cell(r, col)._style = copy(ws.cell(r, 13)._style)
                    ws.cell(r, col, value)
        commands = {
            'church.1104.s1.p015': {'stageLeftKey': 'alienor', 'stageRightKey': 'adele'},
            'church.1104.s1.p071': {'stageLeftKey': 'alienor', 'stageRightKey': 'adele'},
            # Independent night preview reconstructs the two established characters.
            'church.1104.s1.p414': {'stageLeftKey': 'alienor', 'stageRightKey': 'adele'},
            'church.1104.s1.p426': {'stageRightKey': '-'},
        }
        touched = set()
        found = set()
        for row in range(5, ws.max_row + 1):
            key = ws.cell(row, headers['lineId']).value
            if key not in commands:
                continue
            found.add(key)
            for field, value in commands[key].items():
                c = ws.cell(row, headers[field])
                require(c.value in (None, value), f'Concurrent stage conflict: {key}/{field}')
                c.value = value
                touched.add((ws.title, c.coordinate))
        require(found == set(commands), 'Source stage anchor missing')
        out = base / 'E-church' / (datetime.now().strftime('%Y%m%d-%H%M%S') + '-stage')
        out.mkdir(parents=True)
        candidate = out / 'Narrative.candidate.xlsx'
        wb.save(candidate)
        saved = openpyxl.load_workbook(candidate)
        for (sheet, coord), before in old.items():
            if (sheet, coord) not in touched:
                require(signature(saved[sheet][coord]) == before, f'Unrelated cell changed: {sheet}!{coord}')
        saved.close(); wb.close()
        require(digest(NARRATIVE) == sha, 'Concurrent workbook mutation; candidate retained')
        shutil.copy2(NARRATIVE, out / 'Narrative.before.xlsx')
        require(digest(out / 'Narrative.before.xlsx') == sha and digest(NARRATIVE) == sha, 'Backup/CAS race')
        candidate.replace(NARRATIVE)
        manifest = {'before': sha, 'after': digest(NARRATIVE), 'commands': commands, 'evidence': str(out)}
        (out / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
        print(json.dumps(manifest, ensure_ascii=False))
    finally:
        lock.unlink()

if __name__ == '__main__':
    main()
