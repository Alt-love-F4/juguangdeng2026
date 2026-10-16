"""Prepare the blocked E long chapter in the official Narrative/Text sheets.

Usage (run later, under the shared-table owner's coordination):
  python Tools/Narrative/configure_church_long.py          # read-only plan
  python Tools/Narrative/configure_church_long.py --apply  # CAS append + evidence
Then run Assets/ConfigTables/GenerateAll.bat separately.

No Date, Result, completion flag, fake business or 1105 content is created.
The 42 steps comprise 39 text pages, two Wait checkpoints and a terminal fade.
Stage guards use the shared checkpointStage protocol; the root retains blockedBy
until production contracts and full acceptance are resolved. Isolated stage
previews use NarrativeChapterPreview.PlayChurchLongConstruction; they cannot
supply production business receipts or authoritative faction semantics.
"""
from __future__ import annotations

import argparse
from collections import Counter
from copy import copy
from datetime import datetime
import json
from pathlib import Path
import re
import shutil

import openpyxl
from openpyxl.comments import Comment

from import_church_long_text import ROOT, MAP, TEXT, SOURCE, SOURCES, digest, expected_rows

NARRATIVE = ROOT / 'Assets/ConfigTables/Data/Narrative/Narrative.xlsx'
EVENT = 'church.1104.s1'
PRESENTATION = EVENT + '.conversation'
BLOCKER = 'E-B01,E-B02,E-B03,E-B04,E-C01,E-C04,E-resources-unverified'
# These are the currently inspected row-1 contracts, not a new schema.
HEADERS = {
    'Event': 'eventId eventKind actorId targetId titleText contentKey setupType setupKey onceOnly logTag presentationClass presentationKey presentationLogTag eventProperty',
    'PresentationIndex': 'presentationKey presentationClass logTag',
    'SetupParams': 'setupKey tplName slotValues skipAble autoAdvance checkId',
    'ConversationPage': 'pageId presentationKey seq pageCondKeys pageBarConds optionGroupKey bgOverride bgmOverride stepType readingKey checkId',
    'ConversationLine': 'lineId pageId seq speakerKey side textKey voiceKey emote sfxKey accentFxKey accentTargetKey speakerNameKey stageLeftKey stageRightKey stageCenterKey',
    'Choice': 'choiceId groupKey textKey resultOwnerId seq checkId',
    'AbilityCheck': 'checkId successEventId failureEventId successResultOwnerId failureResultOwnerId comment mode',
    'AbilityCheckCondition': 'conditionId checkId attributeId threshold groupIndex seq',
    'Object': 'objId objType nameText iconPath portraitPath portraitIndoorPath avtarPath descText tags',
    'Text': 'key text comment',
}
UNSUPPORTED = {
    'E-B01': 'P002: real 1103 completion/shelter producers and exact keys remain unresolved; no Date or fabricated facts.',
    'E-B02': 'P043/P061/P064 and both market passages: Q fee, C coat and F purchasing-power values/relationships unresolved; xx retained.',
    'E-B03': 'P440/P453: real revolution classification producer, timing, threshold and unknown semantics unresolved; dependency:E-B03.revolution throws formally.',
    'E-B04': 'P149-150/P350-351: D1014 recollection reachability or alternate source text unresolved; no invented prerequisite.',
    'E-B05': 'Both 1105 manuscripts contain only titles; retained solely in provenance, no configuration.',
    'E-C01': 'Shared same-root WaitBusiness/WaitNight, token confirmation, AfterBusiness/BeforeNight and checkpoint save/cancel/retry implemented. Production E entry and final settlement remain blocked pending authored contracts and full integration acceptance.',
    'E-C02': 'choice/anyChoice/not/check gates use controller local snapshots; three Inline checks and choices validated. pageBarConds remains empty.',
    'E-C03': 'speakerNameKey aliases and stage slot snapshots connected; P015/P071/P414 entry and P426 adele-only exit validated. Production art and arbitrary-page context reconstruction remain unverified.',
    'E-C04': 'Existing isolated construction preview supports day30, perception/wisdom, one stage and explicit night faction fixture. Arbitrary-page context and complete formal save/AP/wallet isolation acceptance remain outstanding.',
    'E-R001-022': 'Terminal black fade and character stage anchors connected; production art, expression differences, anchored sounds and object overlays remain unverified. No guessed audio/effects.',
    'FINAL': 'Intentionally no money, completed, thibaut_reported or other Result. Parent must resolve blockers and add one final atomic commit after terminal fade, never at stage end.',
}
CHOICES = {41: 'fee.free', 42: 'fee.one', 43: 'fee.quoted',
           75: 'route.far', 76: 'route.near',
           429: 'night.wait', 430: 'night.ask', 431: 'night.hurry'}
CHECKS = {'morning.perception': ('perception', 3),
          'old.perception': ('perception', 3), 'street.wisdom': ('wisdom', 4)}


def require(ok, message):
    if not ok:
        raise RuntimeError(message)


def full_source():
    """Keep all 475 original rows, routing cells, XML indices and group gates."""
    groups, rows, current = [], [], None
    for number, line in enumerate(MAP.read_text(encoding='utf-8').splitlines(), 1):
        if line.startswith('拟 page/步骤：'):
            current = {'page': re.search(r'`([^`]+)`', line).group(1),
                       'routing': line, 'source_map_line': number, 'rows': []}
            groups.append(current)
        match = re.match(r'^\| (E(?:04|05A|05B)-P(\d{3}))（(\d+)）', line)
        if not match:
            continue
        cells = [c.strip().replace(r'\|', '|') for c in re.split(r'(?<!\\)\|', line)[1:-1]]
        row = dict(source=match[1], number=int(match[2]), xml_paragraph=int(match[3]),
                   meta=cells[1], original=cells[2], routing=cells[3],
                   source_map_line=number, source_map_raw=line)
        rows.append(row)
        if row['source'].startswith('E04-'):
            require(current is not None, 'Source row without page')
            current['rows'].append(row)
    require(len(groups) == 42 and len(rows) == 475, 'Source group/paragraph coverage changed')
    require([r['number'] for r in rows if r['source'].startswith('E04-')] == list(range(1, 474)), 'Source order changed')
    return groups, rows


def stage_of(short):
    if short.startswith(('morning.', 'fee.')):
        return 'morning'
    if short.startswith(('tour.', 'far.', 'near.')):
        return 'tour'
    if short.startswith('night.'):
        return 'night'
    raise RuntimeError('Not a content page: ' + short)


def page_guards(short):
    guards = ['checkpointStage:' + str(('morning', 'tour', 'night').index(stage_of(short)))]
    if short.startswith('fee.'):
        guards.append('choice:' + EVENT + '.' + short)
    if short.startswith(('far.', 'near.')):
        guards.append('choice:' + EVENT + '.route.' + short.split('.')[0])
    if short.endswith('.fail'):
        check = 'morning.perception' if short.startswith('morning.') else 'street.wisdom'
        guards.append('not:check:' + EVENT + '.' + check)
    if short == 'night.question':
        guards.append('anyChoice:' + ';'.join(EVENT + '.night.' + x for x in ('wait', 'ask', 'hurry')))
    if short in ('night.nonrevolution', 'night.revolution'):
        guards.append(('not:' if short.endswith('.nonrevolution') else '') + 'dependency:E-B03.revolution')
    return '|'.join(guards)


def plan():
    entries, counts = expected_rows()  # Verifies all three original DOCX hashes/texts.
    groups, sources = full_source()
    display = {e['number']: e for e in entries}
    rows = []

    def add(sheet, **values):
        rows.append({'sheet': sheet, 'values': values})

    add('Event', eventId=EVENT, eventKind='world', actorId='world', setupKey=EVENT,
        onceOnly=True, logTag='E1104', presentationClass='Conversation',
        presentationKey=PRESENTATION, presentationLogTag='E1104',
        eventProperty='entry=MorningBeforeMap;blocksSave=true;blockedBy=' + BLOCKER)
    add('PresentationIndex', presentationKey=PRESENTATION, presentationClass='Conversation', logTag='E1104')
    add('SetupParams', setupKey=EVENT, tplName='Tpl_ArrivalConversation', slotValues='opts=', skipAble=False, autoAdvance=False)
    names = {'alienor': '阿莉耶诺尔·康多尔', EVENT + '.alias.alienor': '埃莉诺',
             EVENT + '.alias.landlord.uncertain': '房东？'}
    npc_names = dict(mother='女人', vendor='摊主', civilian_a='平民A', civilian_b='平民B',
                     elder_lady='老妇人', buyer='男子', tenant='男人', landlord='房东')
    names.update({EVENT + '.npc.' + k: v for k, v in npc_names.items()})
    for key, value in names.items():
        add('Text', key=key + '.name', text=value, comment='E source identity/display alias; no narrative addition')
    add('Object', objId='alienor', objType='npc', nameText='alienor.name',
        portraitPath='Characters/alienor/neutral', portraitIndoorPath='Characters/alienor/indoor')
    for short in npc_names:
        key = EVENT + '.npc.' + short
        add('Object', objId=key, objType='npc', nameText=key + '.name')
    for entry in entries:
        add('Text', key=entry['key'], text=entry['text'],
            comment=f'{entry["source"]}; {entry["page"]}; blocked E preparation; xx unresolved')
    for short, (attribute, threshold) in CHECKS.items():
        check = EVENT + '.' + short
        add('AbilityCheck', checkId=check, mode='Inline', comment='E source single cached snapshot; no result owner')
        add('AbilityCheckCondition', conditionId=check + '.1', checkId=check,
            attributeId=attribute, threshold=threshold, groupIndex=1, seq=1)
    page_seq = 0
    for group in groups:
        page = group['page']
        short = page.removeprefix(EVENT + '.')
        content = [r for r in group['rows'] if r['meta'].startswith('正文；')]
        if not content:
            require(short == 'source' or short.startswith('boundary.'), 'Unexpected empty content group')
            continue
        if short in ('tour.depart', 'night.return'):
            page_seq += 1
            wait = 'business' if short == 'tour.depart' else 'night'
            stage = 'morning' if wait == 'business' else 'tour'
            add('ConversationPage', pageId=EVENT + '.checkpoint.' + wait,
                presentationKey=PRESENTATION, seq=page_seq,
                pageCondKeys='checkpointStage:' + str(('morning', 'tour', 'night').index(stage)),
                stepType='WaitBusiness' if wait == 'business' else 'WaitNight')
        page_seq += 1
        check = ('morning.perception' if short == 'morning.perception.pass' else
                 'old.perception' if short.endswith('.old.perception') else
                 'street.wisdom' if short == 'far.street.wisdom.pass' else None)
        option = {'morning.invite': 'fee', 'tour.depart': 'route', 'night.private': 'night'}.get(short)
        background = ('oct07.cafe.night' if stage_of(short) == 'night' else
                      'oct07.cafe.day' if short.startswith(('fee.', 'morning.cafe', 'morning.perception', 'morning.invite', 'tour.')) else
                      'oct07.street.day' if short in ('near.street', 'far.street.intro') else
                      EVENT + '.background.' + short)
        add('ConversationPage', pageId=page, presentationKey=PRESENTATION, seq=page_seq,
            pageCondKeys=page_guards(short), optionGroupKey=EVENT + '.' + option + '.choices' if option else None,
            bgOverride=background, stepType='PassiveCheck' if check else 'Dialogue',
            checkId=EVENT + '.' + check if check else None)
        for seq, source in enumerate(content, 1):
            meta = source['meta']
            speaker = re.search(r'；speaker=([^；]+)', meta)
            side = re.search(r'；side=([^；]+)', meta)
            emote = re.search(r'；表情=([^；]+)', meta)
            speaker = speaker[1] if speaker else None
            alias = (EVENT + '.alias.alienor.name' if speaker == 'alienor' else
                     EVENT + '.alias.landlord.uncertain.name' if source['number'] in (196, 296) else None)
            entry = display[source['number']]
            add('ConversationLine', lineId=entry['key'], pageId=page, seq=seq,
                speakerKey=speaker, side=side[1] if side else 'none', textKey=entry['key'],
                emote=emote[1] if emote and emote[1] != '空→保持' else None, speakerNameKey=alias,
                stageLeftKey='alienor' if source['number'] in (15, 71, 414) else None,
                stageRightKey='adele' if source['number'] in (15, 71, 414) else '-' if source['number'] == 426 else None)
    # Both source terminal blackouts (P452/P473) share the same post-text step.
    add('ConversationPage', pageId=EVENT + '.night.fade', presentationKey=PRESENTATION,
        seq=page_seq + 1, pageCondKeys='checkpointStage:2', stepType='FadeOut')
    for number, short in CHOICES.items():
        group = short.split('.')[0]
        seq = [n for n, s in CHOICES.items() if s.startswith(group + '.')].index(number) + 1
        add('Choice', choiceId=EVENT + '.' + short, groupKey=EVENT + '.' + group + '.choices',
            textKey=display[number]['key'], seq=seq)
    totals = Counter(r['sheet'] for r in rows)
    require(totals['ConversationLine'] == 405 and totals['Choice'] == 8 and totals['ConversationPage'] == 42, 'Prepared coverage changed')
    return rows, {'event': EVENT, 'source_counts': counts, 'counts': dict(totals),
                  'groups': groups, 'sources': sources, 'unsupported': UNSUPPORTED,
                  'preview_enabled': False, 'source_map_sha256': digest(MAP),
                  'docx_sha256': {name: digest(SOURCE / name) for name in SOURCES}}


def signature(cell):
    return (cell.value, cell.data_type, tuple(cell._style) if cell._style is not None else (0,) * 9,
            (cell.comment.text, cell.comment.author) if cell.comment else None,
            copy(cell.hyperlink))


def prepare(books, rows, manifest):
    locations, added = [], Counter()
    groups = {g['page']: g for g in manifest['groups']}
    sources = {EVENT + f'.p{r["number"]:03}': r for r in manifest['sources'] if r['source'].startswith('E04-')}
    indexes = {}
    for item in rows:
        sheet, values = item['sheet'], item['values']
        path = TEXT if sheet == 'Text' else NARRATIVE
        ws = books[path][sheet]
        headers = ['##var'] + HEADERS[sheet].split()
        require([c.value for c in ws[1]] == headers, f'Schema drift: {sheet}')
        heads = {v: n for n, v in enumerate(headers, 1)}
        key_field = HEADERS[sheet].split()[0]
        key = values[key_field]
        if sheet not in indexes:
            index = {}
            for row in ws.iter_rows(min_row=5):
                old_key = row[1].value
                if old_key is not None:
                    require(old_key not in index, f'Duplicate existing key: {sheet}/{old_key}')
                    index[old_key] = row[1].row
            indexes[sheet] = index
        index = indexes[sheet]
        if key in index:
            number = index[key]
            # Text already staged by the authoritative importer has different comments.
            fields = ['key', 'text'] if sheet == 'Text' else HEADERS[sheet].split()
            # A shared character registered elsewhere must be reused without rewriting.
            if sheet == 'Object' and key == 'alienor':
                require(ws.cell(number, heads['objType']).value == 'npc', 'alienor identity conflict')
            else:
                for field in fields:
                    require(ws.cell(number, heads[field]).value == values.get(field), f'Existing row differs: {sheet}/{key}/{field}')
        else:
            number = ws.max_row + 1
            for column in range(1, ws.max_column + 1):
                ws.cell(number, column)._style = copy(ws.cell(number - 1, column)._style)
            for field, value in values.items():
                if value is not None:
                    ws.cell(number, heads[field], value)
            provenance = groups.get(key) or sources.get(key) or {'scope': 'E blocked configuration preparation'}
            ws.cell(number, 2).comment = Comment(json.dumps(provenance, ensure_ascii=False, indent=2), 'E-church source routing')
            index[key] = number
            added[sheet] += 1
        locations.append({'sheet': sheet, 'key': key, 'row': number})
    # Never inherit externally added E outcomes/date candidates without an explicit extension.
    for sheet, field in [('Result', 'eventId'), ('EventDate', 'eventId')]:
        ws = books[NARRATIVE][sheet]
        headers = {c.value: c.column for c in ws[1]}
        require(not any(ws.cell(r, headers[field]).value == EVENT for r in range(5, ws.max_row + 1)), f'E must remain result/date-free: {sheet}')
    return dict(added), locations


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--apply', action='store_true', help='Append under cooperative lock; verify latest workbook hashes before replacement')
    args = parser.parse_args()
    rows, manifest = plan()
    paths = (NARRATIVE, TEXT)
    if not args.apply:
        books = {p: openpyxl.load_workbook(p) for p in paths}
        try:
            added, locations = prepare(books, rows, manifest)  # Memory only.
            print(json.dumps({**manifest, 'would_append': added, 'locations': locations, 'writes': False}, ensure_ascii=False, indent=2))
        finally:
            for wb in books.values():
                wb.close()
        return
    base = ROOT / 'Library/Locus/tmp'
    base.mkdir(parents=True, exist_ok=True)
    lock = base / 'narrative-source-write.lock'
    with lock.open('x', encoding='utf-8') as handle:
        handle.write('E blocked long chapter append\n')
    books = {}
    try:
        hashes = {p: digest(p) for p in paths}
        books = {p: openpyxl.load_workbook(p) for p in paths}
        before = {p: {(s.title, c.coordinate): signature(c) for s in w for row in s for c in row} for p, w in books.items()}
        added, locations = prepare(books, rows, manifest)
        if not added:
            print(json.dumps({'idempotent': True, 'changed': False, 'unsupported': UNSUPPORTED}, ensure_ascii=False, indent=2))
            return
        out = base / 'E-church' / (datetime.now().strftime('%Y%m%d-%H%M%S-%f') + '-config')
        out.mkdir(parents=True)
        candidates = {}
        for p, wb in books.items():
            require(digest(p) == hashes[p], f'Concurrent workbook change before backup: {p}')
            shutil.copy2(p, out / (p.stem + '.before.xlsx'))
            require(digest(out / (p.stem + '.before.xlsx')) == hashes[p], f'Backup race: {p}')
            candidate = out / (p.stem + '.candidate.xlsx')
            wb.save(candidate)
            saved = openpyxl.load_workbook(candidate)
            try:
                for (sheet, coordinate), old in before[p].items():
                    require(signature(saved[sheet][coordinate]) == old, f'Existing cell changed: {sheet}!{coordinate}')
                for item in rows:
                    sheet, values = item['sheet'], item['values']
                    if (sheet == 'Text') != (p == TEXT):
                        continue
                    location = next(x for x in locations if x['sheet'] == sheet and x['key'] == next(iter(values.values())))
                    for column in range(1, books[p][sheet].max_column + 1):
                        coordinate = books[p][sheet].cell(location['row'], column).coordinate
                        require(signature(saved[sheet][coordinate]) == signature(books[p][sheet][coordinate]),
                                f'Candidate append changed: {sheet}!{coordinate}')
            finally:
                saved.close()
            candidates[p] = candidate
        manifest.update(time=datetime.now().astimezone().isoformat(), added=added, locations=locations,
                        before_sha256={str(p): h for p, h in hashes.items()}, replaced=[])
        journal = out / 'manifest.json'
        def record():
            journal.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
        record()
        # Two filesystem replacements are not one atomic transaction. Journal each
        # success, never roll back over another writer, and rerun after a partial
        # apply. Per-file CAS plus the shared lock is cooperative, not an OS CAS.
        for p in paths:
            require(digest(p) == hashes[p], f'Concurrent workbook change: {p}; candidates retained at {out}')
        for p in (TEXT, NARRATIVE):  # Dependencies before references.
            require(digest(p) == hashes[p], f'Concurrent workbook change: {p}; see {journal}')
            if not any(count for sheet, count in added.items() if (sheet == 'Text') == (p == TEXT)):
                continue
            candidates[p].replace(p)
            manifest['replaced'].append(str(p))
            manifest.setdefault('after_sha256', {})[str(p)] = digest(p)
            record()
        print(json.dumps({'added': added, 'evidence': str(out), 'preview_enabled': False}, ensure_ascii=False, indent=2))
    finally:
        for wb in books.values():
            wb.close()
        lock.unlink()


if __name__ == '__main__':
    main()
